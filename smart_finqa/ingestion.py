from __future__ import annotations

import re
import shutil
import subprocess
import warnings
from dataclasses import dataclass
from enum import Enum
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook
from pypdf import PdfReader
from pypdf.errors import PyPdfError

from .config import OCRConfig
from .core import MetadataRecognitionError, parse_report_period
from .extraction import extract_metric_snapshot_from_text
from .ocr import OCR_PREPROCESSING_VERSION, OCREngine, OCRPage, OCRPageError, OCRToken, RapidOCREngine

REPORT_METADATA_VERSION = "report-metadata-v1"
REPORT_SELECTION_VERSION = "report-selection-v1"
FINANCIAL_STATEMENT_PAGE_PATTERN = re.compile(
    r"(?m)^\s*(?:(?:合并|母公司)?(?:资产负债表|利润表|损益表|现金流量表)|"
    r"主要会计数据(?:和|及)财务指标|主要财务指标)\s*$"
)
LARGE_RASTER_IMAGE_MIN_PIXELS = 500_000
LARGE_RASTER_IMAGE_MIN_SIDE = 400


@dataclass(frozen=True, slots=True)
class PageText:
    page_no: int
    text: str
    source_type: str = "native_text"
    ocr_tokens: tuple[OCRToken, ...] = ()
    ocr_engine: str | None = None
    ocr_engine_version: str | None = None
    page_image_sha256: str | None = None
    text_layer_text: str | None = None
    has_large_raster_image: bool = False

    def __post_init__(self) -> None:
        if isinstance(self.page_no, bool) or not isinstance(self.page_no, int) or self.page_no < 1:
            raise ValueError("page_no must be a positive integer")
        if not isinstance(self.text, str):
            raise TypeError("text must be a string")
        if not isinstance(self.has_large_raster_image, bool):
            raise TypeError("has_large_raster_image must be a boolean")
        if self.source_type not in {"native_text", "pdftotext", "ocr", "mixed"}:
            raise ValueError("source_type must be native_text, pdftotext, ocr, or mixed")
        if any(token.page_no != self.page_no for token in self.ocr_tokens):
            raise ValueError("OCR tokens must belong to the PageText physical page")
        if self.source_type in {"ocr", "mixed"}:
            if not self.ocr_engine or not self.ocr_engine_version or not self.page_image_sha256:
                raise ValueError("OCR-backed PageText requires engine and page image provenance")
            if any(token.page_image_sha256 != self.page_image_sha256 for token in self.ocr_tokens):
                raise ValueError("OCR tokens must match the PageText page image hash")
        elif self.ocr_tokens or self.ocr_engine or self.ocr_engine_version or self.page_image_sha256:
            raise ValueError("non-OCR PageText cannot carry OCR provenance")
        if self.source_type == "mixed":
            if self.text_layer_text is None or not self.text_layer_text.strip():
                raise ValueError("mixed PageText requires non-empty text-layer evidence")
            if self.text_layer_text.strip() not in self.text:
                raise ValueError("mixed PageText must preserve its text-layer evidence")
        elif self.text_layer_text is not None:
            raise ValueError("only mixed PageText can carry separate text-layer evidence")


class ReportRevisionKind(str, Enum):
    ORIGINAL = "original"
    PRE_REVISION = "pre_revision"
    REVISED = "revised"


@dataclass(frozen=True, slots=True)
class ReportFileMetadata:
    path: Path
    stock_code: str
    stock_abbr: str
    report_period: str
    report_year: int
    is_summary: bool
    is_english: bool
    revision_kind: ReportRevisionKind
    published_date: str | None
    copy_index: int
    page_count: int | None
    head_text_length: int
    read_error: str | None = None
    company_in_master: bool = True


@dataclass(frozen=True, slots=True)
class ReportSelectionDecision:
    path: Path
    status: str
    reason: str
    report: ReportFileMetadata


def report_metadata_to_dict(report: ReportFileMetadata) -> dict[str, Any]:
    return {
        "stock_code": report.stock_code,
        "stock_abbr": report.stock_abbr,
        "report_period": report.report_period,
        "report_year": report.report_year,
        "is_summary": report.is_summary,
        "is_english": report.is_english,
        "revision_kind": report.revision_kind.value,
        "published_date": report.published_date,
        "copy_index": report.copy_index,
        "page_count": report.page_count,
        "head_text_length": report.head_text_length,
        "read_error": report.read_error,
        "company_in_master": report.company_in_master,
    }


def report_metadata_from_dict(path: Path, values: Mapping[str, Any]) -> ReportFileMetadata:
    try:
        return ReportFileMetadata(
            path=path,
            stock_code=str(values["stock_code"]),
            stock_abbr=str(values["stock_abbr"]),
            report_period=str(values["report_period"]),
            report_year=int(values["report_year"]),
            is_summary=_required_bool(values, "is_summary"),
            is_english=_required_bool(values, "is_english"),
            revision_kind=ReportRevisionKind(str(values["revision_kind"])),
            published_date=str(values["published_date"]) if values.get("published_date") else None,
            copy_index=int(values["copy_index"]),
            page_count=int(values["page_count"]) if values.get("page_count") is not None else None,
            head_text_length=int(values["head_text_length"]),
            read_error=str(values["read_error"]) if values.get("read_error") else None,
            company_in_master=_required_bool(values, "company_in_master"),
        )
    except (KeyError, TypeError, ValueError) as exc:
        raise ValueError(f"invalid cached report metadata for {path.name}") from exc


def _required_bool(values: Mapping[str, Any], key: str) -> bool:
    value = values[key]
    if not isinstance(value, bool):
        raise TypeError(f"{key} must be a boolean")
    return value


@dataclass(slots=True)
class ReportRecord:
    source_path: Path
    stock_code: str
    stock_abbr: str
    report_period: str
    report_year: int
    text: str
    pages: tuple[PageText, ...]
    snapshot: dict[str, float | None]


@dataclass(slots=True)
class CompanyIndex:
    code_to_abbr: dict[str, str]
    abbr_to_code: dict[str, str]

    @classmethod
    def from_xlsx(cls, path: Path) -> "CompanyIndex":
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        code_to_abbr: dict[str, str] = {}
        abbr_to_code: dict[str, str] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            code = row[1]
            abbr = row[2]
            if code is None or not abbr:
                continue
            code_str = str(int(code)).zfill(6)
            abbr_str = str(abbr).strip()
            code_to_abbr[code_str] = abbr_str
            abbr_to_code[abbr_str] = code_str
            abbr_to_code[_normalize_company_name(abbr_str)] = code_str
        return cls(code_to_abbr=code_to_abbr, abbr_to_code=abbr_to_code)


_REPORT_TITLE_PATTERN = re.compile(
    r"(?P<year>\d{4})\s*年\s*"
    r"(?P<period>一季度|半年度|三季度|年度|年报|第三季度|第一季度)\s*"
    r"(?:报告)?\s*(?P<summary>摘要)?"
)
_PRE_REVISION_PATTERN = re.compile(r"更新前|更正前|修订前")
_REVISED_PATTERN = re.compile(r"更新后|更正后|修订后|修订版|修订稿|更新版|更正版")
_ENGLISH_VERSION_PATTERN = re.compile(r"英文版|English\s+Version", re.IGNORECASE)
_PUBLISHED_DATE_PATTERN = re.compile(r"(?:^|_)(?P<date>20\d{6})(?:_|\.)")
_COPY_INDEX_PATTERN = re.compile(r"[（(](?P<index>\d+)[)）]\s*$")


def classify_report_file_name(
    path: Path,
    company_index: CompanyIndex,
    *,
    text_head: str = "",
    page_count: int | None = None,
    head_text_length: int | None = None,
    read_error: str | None = None,
) -> ReportFileMetadata:
    """Classify one report from its public metadata without selecting it for ingestion."""
    source = f"{path.name}\n{text_head}"
    file_name_company = _company_name_from_file_name(path.name)
    file_name_in_master = file_name_company is None or file_name_company in company_index.abbr_to_code
    try:
        stock_code = _infer_stock_code(path.name, text_head, company_index)
    except MetadataRecognitionError:
        if file_name_in_master or file_name_company is None:
            raise
        stock_code = ""
    stock_abbr = (
        _infer_stock_abbr(path.name, text_head, stock_code, company_index) if stock_code else file_name_company or ""
    )
    report_period, report_year = parse_report_period(path.name, text_head)
    title_match = _REPORT_TITLE_PATTERN.search(source)
    if title_match is None:
        raise MetadataRecognitionError("report_type", path.name)
    revision_kind = ReportRevisionKind.ORIGINAL
    if _PRE_REVISION_PATTERN.search(source):
        revision_kind = ReportRevisionKind.PRE_REVISION
    elif _REVISED_PATTERN.search(source):
        revision_kind = ReportRevisionKind.REVISED
    published_match = _PUBLISHED_DATE_PATTERN.search(path.name)
    copy_match = _COPY_INDEX_PATTERN.search(path.stem)
    return ReportFileMetadata(
        path=path,
        stock_code=stock_code,
        stock_abbr=stock_abbr,
        report_period=report_period,
        report_year=report_year,
        is_summary=bool(title_match.group("summary")),
        is_english=bool(_ENGLISH_VERSION_PATTERN.search(source)),
        revision_kind=revision_kind,
        published_date=published_match.group("date") if published_match else None,
        copy_index=int(copy_match.group("index")) if copy_match else 0,
        page_count=page_count,
        head_text_length=len(text_head) if head_text_length is None else head_text_length,
        read_error=read_error,
        company_in_master=stock_code in company_index.code_to_abbr,
    )


def inspect_report_file_metadata(
    path: Path,
    company_index: CompanyIndex,
    *,
    page_extractor: PdfPageExtractor | None = None,
) -> ReportFileMetadata:
    """Read at most the first two pages and return auditable report metadata."""
    if page_extractor is not None:
        pages, page_count = page_extractor.extract_head(path, max_pages=2)
        text_head = "\n".join(page.text for page in pages)[:10000]
        return classify_report_file_name(
            path,
            company_index,
            text_head=text_head,
            page_count=page_count,
            head_text_length=len(text_head),
        )
    text_parts: list[str] = []
    read_errors: list[str] = []
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            reader = PdfReader(str(path))
            if getattr(reader, "is_encrypted", False):
                try:
                    reader.decrypt("")
                except Exception as exc:
                    read_errors.append(f"decrypt:{type(exc).__name__}")
            page_count = len(reader.pages)
            for page_index in range(min(2, page_count)):
                try:
                    text_parts.append(reader.pages[page_index].extract_text() or "")
                except Exception as exc:
                    read_errors.append(f"page_{page_index + 1}:{type(exc).__name__}")
    except Exception as exc:
        raise MetadataRecognitionError("pdf_read", path.name) from exc
    text_head = "\n".join(text_parts)[:10000]
    return classify_report_file_name(
        path,
        company_index,
        text_head=text_head,
        page_count=page_count,
        head_text_length=len(text_head),
        read_error=";".join(read_errors) or None,
    )


def select_authoritative_reports(
    reports: Iterable[ReportFileMetadata],
) -> tuple[ReportSelectionDecision, ...]:
    """Choose one authoritative full Chinese report for every company and period."""
    ordered = tuple(reports)
    if len({report.path for report in ordered}) != len(ordered):
        raise ValueError("report paths must be unique")
    decisions: dict[Path, ReportSelectionDecision] = {}
    candidates: dict[tuple[str, str], list[ReportFileMetadata]] = {}
    for report in ordered:
        if not report.company_in_master:
            decisions[report.path] = _selection_decision(report, "needs_review", "outside_company_master")
        elif report.is_summary:
            decisions[report.path] = _selection_decision(report, "skipped_summary", "summary_not_authoritative")
        elif report.is_english:
            decisions[report.path] = _selection_decision(report, "superseded", "english_translation")
        elif report.revision_kind is ReportRevisionKind.PRE_REVISION:
            decisions[report.path] = _selection_decision(report, "superseded", "pre_revision")
        else:
            candidates.setdefault((report.stock_code, report.report_period), []).append(report)

    for group in candidates.values():
        scores = {report.path: _authority_score(report) for report in group}
        highest = max(scores.values())
        winners = [report for report in group if scores[report.path] == highest]
        if len(winners) != 1:
            for report in winners:
                decisions[report.path] = _selection_decision(
                    report,
                    "needs_review",
                    "ambiguous_authoritative_version",
                )
            for report in group:
                if report.path not in decisions:
                    decisions[report.path] = _selection_decision(report, "superseded", "higher_ranked_version_exists")
            continue
        selected = winners[0]
        decisions[selected.path] = _selection_decision(selected, "selected", "authoritative_version")
        for report in group:
            if report.path != selected.path:
                decisions[report.path] = _selection_decision(report, "superseded", "higher_ranked_version_selected")
    return tuple(decisions[report.path] for report in ordered)


def _selection_decision(report: ReportFileMetadata, status: str, reason: str) -> ReportSelectionDecision:
    return ReportSelectionDecision(path=report.path, status=status, reason=reason, report=report)


def _authority_score(report: ReportFileMetadata) -> tuple[int, str, int]:
    revision_rank = 1 if report.revision_kind is ReportRevisionKind.REVISED else 0
    return revision_rank, report.published_date or "", -report.copy_index


def _normalize_company_name(value: str) -> str:
    return re.sub(r"\s+", "", value.strip())


def _company_name_from_file_name(file_name: str) -> str | None:
    if "：" not in file_name and ":" not in file_name:
        return None
    candidate = _normalize_company_name(re.split(r"[:：]", file_name, maxsplit=1)[0])
    return candidate or None


class PdfPageExtractor:
    """Read a PDF once and OCR only the configured physical pages."""

    def __init__(self, config: OCRConfig | None = None, *, ocr_engine: OCREngine | None = None) -> None:
        self.config = config or OCRConfig()
        self._ocr_engine = ocr_engine

    @property
    def profile(self) -> dict[str, Any]:
        return {
            "preprocessing_version": OCR_PREPROCESSING_VERSION,
            "engine": self.config.engine,
            "policy": self.config.policy,
            "dpi": self.config.dpi,
            "min_page_text_chars": self.config.min_page_text_chars,
            "min_confidence": self.config.min_confidence,
            "max_page_pixels": self.config.max_page_pixels,
        }

    def extract(self, pdf_path: Path, *, ocr_page_numbers: set[int] | None = None) -> tuple[PageText, ...]:
        native_pages = _read_native_pdf_pages(pdf_path)
        should_try_alternate = self.config.policy != "always" and any(
            _non_whitespace_length(page.text) < self.config.min_page_text_chars for page in native_pages
        )
        alternate_pages = _pdftotext_pages_fallback(pdf_path) if should_try_alternate else ()
        alternate_by_page = {page.page_no: page for page in alternate_pages}
        pages = tuple(_prefer_page_text(page, alternate_by_page.get(page.page_no)) for page in native_pages)
        return self._apply_ocr(pdf_path, pages, ocr_page_numbers=ocr_page_numbers)

    def extract_head(self, pdf_path: Path, *, max_pages: int = 2) -> tuple[tuple[PageText, ...], int]:
        if isinstance(max_pages, bool) or not isinstance(max_pages, int) or max_pages < 1:
            raise ValueError("max_pages must be a positive integer")
        pages, page_count = _read_native_pdf_head(pdf_path, max_pages=max_pages)
        return self._apply_ocr(pdf_path, pages), page_count

    def _apply_ocr(
        self,
        pdf_path: Path,
        pages: tuple[PageText, ...],
        *,
        ocr_page_numbers: set[int] | None = None,
    ) -> tuple[PageText, ...]:
        if not self.config.enabled:
            return pages

        physical_pages = {page.page_no for page in pages}
        allowed_pages = physical_pages if ocr_page_numbers is None else ocr_page_numbers & physical_pages
        targets = [page for page in pages if page.page_no in allowed_pages and self._should_ocr(page)]
        if not targets:
            return pages

        engine = self._engine()
        ocr_by_page = {}
        for page in targets:
            ocr_page = engine.recognize_page(pdf_path, page.page_no, dpi=self.config.dpi)
            if ocr_page.page_no != page.page_no:
                raise OCRPageError(
                    f"OCR_PAGE_FAILED: engine returned page {ocr_page.page_no} for requested page {page.page_no}"
                )
            ocr_by_page[page.page_no] = ocr_page
        return tuple(
            self._merge_ocr_page(page, ocr_by_page[page.page_no], engine) if page.page_no in ocr_by_page else page
            for page in pages
        )

    def _should_ocr(self, page: PageText) -> bool:
        if self.config.policy == "always":
            return True
        if _non_whitespace_length(page.text) < self.config.min_page_text_chars:
            return True
        return self.config.policy == "financial_pages_and_low_text" and bool(
            page.has_large_raster_image or FINANCIAL_STATEMENT_PAGE_PATTERN.search(page.text)
        )

    @staticmethod
    def _merge_ocr_page(page: PageText, ocr_page: OCRPage, engine: OCREngine) -> PageText:
        text_layer = page.text.strip()
        if not text_layer:
            return PageText(
                page_no=page.page_no,
                text=ocr_page.text,
                source_type="ocr",
                ocr_tokens=ocr_page.tokens,
                ocr_engine=engine.name,
                ocr_engine_version=engine.version,
                page_image_sha256=ocr_page.page_image_sha256,
                has_large_raster_image=page.has_large_raster_image,
            )
        return PageText(
            page_no=page.page_no,
            text=_merge_text_evidence(page.text, ocr_page.text),
            source_type="mixed",
            ocr_tokens=ocr_page.tokens,
            ocr_engine=engine.name,
            ocr_engine_version=engine.version,
            page_image_sha256=ocr_page.page_image_sha256,
            text_layer_text=page.text,
            has_large_raster_image=page.has_large_raster_image,
        )

    def _engine(self) -> OCREngine:
        if self._ocr_engine is None:
            if self.config.engine == "rapidocr":
                self._ocr_engine = RapidOCREngine(self.config)
            else:  # Configuration validation prevents unknown engines.
                raise ValueError(f"Unsupported OCR engine: {self.config.engine!r}")
        return self._ocr_engine


def extract_pdf_pages(
    pdf_path: Path,
    *,
    ocr_config: OCRConfig | None = None,
    ocr_engine: OCREngine | None = None,
    ocr_page_numbers: set[int] | None = None,
) -> tuple[PageText, ...]:
    return PdfPageExtractor(ocr_config, ocr_engine=ocr_engine).extract(
        pdf_path,
        ocr_page_numbers=ocr_page_numbers,
    )


def _merge_text_evidence(text_layer: str, ocr_text: str) -> str:
    left = text_layer.strip()
    right = ocr_text.strip()
    if not right or right in left:
        return left
    if left in right:
        return right
    return f"{left}\n{right}"


def _read_native_pdf_pages(pdf_path: Path) -> tuple[PageText, ...]:
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            reader = PdfReader(str(pdf_path))
            if getattr(reader, "is_encrypted", False):
                try:
                    if reader.decrypt("") == 0:
                        raise ValueError("PDF requires a password")
                except (PyPdfError, TypeError, ValueError) as exc:
                    raise ValueError(f"cannot decrypt PDF: {pdf_path}") from exc
            return tuple(
                PageText(
                    page_no=index,
                    text=page.extract_text() or "",
                    has_large_raster_image=_has_large_raster_image(page),
                )
                for index, page in enumerate(reader.pages, 1)
            )
    except (OSError, PyPdfError, KeyError, TypeError, ValueError) as exc:
        raise ValueError(f"failed to read PDF text layer: {pdf_path}") from exc


def _read_native_pdf_head(pdf_path: Path, *, max_pages: int) -> tuple[tuple[PageText, ...], int]:
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            reader = PdfReader(str(pdf_path))
            if getattr(reader, "is_encrypted", False):
                try:
                    if reader.decrypt("") == 0:
                        raise ValueError("PDF requires a password")
                except (PyPdfError, TypeError, ValueError) as exc:
                    raise ValueError(f"cannot decrypt PDF: {pdf_path}") from exc
            page_count = len(reader.pages)
            pages = tuple(
                PageText(
                    page_no=index + 1,
                    text=reader.pages[index].extract_text() or "",
                    has_large_raster_image=_has_large_raster_image(reader.pages[index]),
                )
                for index in range(min(max_pages, page_count))
            )
            return pages, page_count
    except (OSError, PyPdfError, KeyError, TypeError, ValueError) as exc:
        raise ValueError(f"failed to read PDF metadata pages: {pdf_path}") from exc


def _prefer_page_text(native: PageText, alternate: PageText | None) -> PageText:
    if alternate is None or _non_whitespace_length(native.text) >= _non_whitespace_length(alternate.text):
        return native
    return PageText(
        page_no=native.page_no,
        text=alternate.text,
        source_type="pdftotext",
        has_large_raster_image=native.has_large_raster_image,
    )


def _has_large_raster_image(page: Any) -> bool:
    resources = _resolve_pdf_object(page.get("/Resources"))
    if not hasattr(resources, "get"):
        return False
    return _resources_have_large_raster_image(resources, visited=set())


def _resources_have_large_raster_image(resources: Any, *, visited: set[int]) -> bool:
    xobjects = _resolve_pdf_object(resources.get("/XObject"))
    if not hasattr(xobjects, "values"):
        return False
    for candidate in xobjects.values():
        obj = _resolve_pdf_object(candidate)
        identity = id(obj)
        if identity in visited or not hasattr(obj, "get"):
            continue
        visited.add(identity)
        subtype = str(obj.get("/Subtype") or "")
        if subtype == "/Image":
            width = int(obj.get("/Width") or 0)
            height = int(obj.get("/Height") or 0)
            if min(width, height) >= LARGE_RASTER_IMAGE_MIN_SIDE and width * height >= LARGE_RASTER_IMAGE_MIN_PIXELS:
                return True
        if subtype == "/Form":
            nested_resources = _resolve_pdf_object(obj.get("/Resources"))
            if hasattr(nested_resources, "get") and _resources_have_large_raster_image(
                nested_resources, visited=visited
            ):
                return True
    return False


def _resolve_pdf_object(value: Any) -> Any:
    get_object = getattr(value, "get_object", None)
    return get_object() if callable(get_object) else value


def _non_whitespace_length(text: str) -> int:
    return len(re.sub(r"\s+", "", text))


def extract_pdf_text(pdf_path: Path) -> str:
    """Extract flat text while preserving the existing public interface."""
    return _join_page_text(extract_pdf_pages(pdf_path))


def _join_page_text(pages: tuple[PageText, ...]) -> str:
    return "\n".join(page.text for page in pages)


def _pages_from_text(text: str) -> tuple[PageText, ...]:
    if not text:
        return ()
    chunks = text.split("\f")
    if chunks and not chunks[-1].strip():
        chunks.pop()
    return tuple(PageText(page_no=index, text=chunk) for index, chunk in enumerate(chunks, 1))


def _pdftotext_pages_fallback(pdf_path: Path) -> tuple[PageText, ...]:
    if not shutil.which("pdftotext"):
        return ()
    try:
        proc = subprocess.run(
            ["pdftotext", str(pdf_path), "-"],
            check=False,
            capture_output=True,
            text=True,
            errors="ignore",
            timeout=60,
        )
        if proc.stdout:
            return tuple(
                PageText(page_no=page.page_no, text=page.text, source_type="pdftotext")
                for page in _pages_from_text(proc.stdout)
            )
    except (OSError, subprocess.SubprocessError):
        return ()
    return ()


def _pdftotext_fallback(pdf_path: Path) -> str:
    return _join_page_text(_pdftotext_pages_fallback(pdf_path))


def _infer_stock_code(file_name: str, text: str, company_index: CompanyIndex) -> str:
    code_match = re.search(r"(?<!\d)(\d{6})(?!\d)", file_name)
    if code_match and code_match.group(1) != "000000":
        return code_match.group(1)
    if "：" in file_name or ":" in file_name:
        candidate = re.split(r"[:：]", file_name, maxsplit=1)[0]
        indexed_code = company_index.abbr_to_code.get(candidate)
        if indexed_code is None:
            indexed_code = company_index.abbr_to_code.get(_normalize_company_name(candidate))
        if indexed_code and indexed_code != "000000":
            return indexed_code
    text_match = re.search(r"(?:公司|证券)代码\s*[:：]?\s*(?<!\d)(\d{6})(?!\d)", text)
    if text_match and text_match.group(1) != "000000":
        return text_match.group(1)
    for abbr, code in company_index.abbr_to_code.items():
        if abbr in text and code != "000000":
            return code
    raise MetadataRecognitionError("stock_code", file_name)


def _infer_stock_abbr(file_name: str, text: str, stock_code: str, company_index: CompanyIndex) -> str:
    indexed_abbr = company_index.code_to_abbr.get(stock_code)
    if indexed_abbr and _is_known_company_name(indexed_abbr):
        return indexed_abbr
    if "：" in file_name or ":" in file_name:
        candidate = re.split(r"[:：]", file_name, maxsplit=1)[0]
        if _is_known_company_name(candidate):
            return _normalize_company_name(candidate)
    match = re.search(r"(?:公司|证券)简称\s*[:：]?\s*([^\s]+)", text)
    if match and _is_known_company_name(match.group(1)):
        return match.group(1).strip()
    raise MetadataRecognitionError("company", file_name)


def _is_known_company_name(value: str) -> bool:
    return bool(value.strip()) and value.strip().lower() not in {"unknown", "未知", "未知公司"}


def extract_report_record(
    pdf_path: Path,
    company_index: CompanyIndex,
    metadata: ReportFileMetadata | None = None,
    *,
    page_extractor: PdfPageExtractor | None = None,
) -> ReportRecord:
    pages = page_extractor.extract(pdf_path) if page_extractor is not None else extract_pdf_pages(pdf_path)
    text = _join_page_text(pages)
    text_head = text[:5000]
    metadata = metadata or classify_report_file_name(
        pdf_path,
        company_index,
        text_head=text_head,
        page_count=len(pages) or None,
        head_text_length=len(text_head),
    )
    snapshot = extract_metric_snapshot_from_text(text)
    return ReportRecord(
        source_path=pdf_path,
        stock_code=metadata.stock_code,
        stock_abbr=metadata.stock_abbr,
        report_period=metadata.report_period,
        report_year=metadata.report_year,
        text=text,
        pages=pages,
        snapshot=snapshot,
    )
