from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass(slots=True)
class FieldSpec:
    field_name: str
    cn_name: str
    raw_type: str
    description: str


SHEET_TO_TABLE = {
    "核心业绩指标表": "core_performance_indicators_sheet",
    "资产负债表": "balance_sheet",
    "现金流量表": "cash_flow_sheet",
    "利润表": "income_sheet",
}
FIELD_NAME_PATTERN = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


def map_to_sqlite_type(raw_type: str | None) -> str:
    text = (raw_type or "").strip().lower()
    if "int" in text:
        return "INTEGER"
    if "decimal" in text or "float" in text or "double" in text:
        return "REAL"
    return "TEXT"


def load_schema_from_xlsx(path: Path) -> dict[str, list[FieldSpec]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    schema: dict[str, list[FieldSpec]] = {}
    for sheet_name, table_name in SHEET_TO_TABLE.items():
        ws = wb[sheet_name]
        fields: list[FieldSpec] = []
        seen_fields: set[str] = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            field_name = str(row[0]).strip()
            if FIELD_NAME_PATTERN.fullmatch(field_name) is None:
                raise ValueError(f"Invalid schema field name in {sheet_name}: {field_name!r}")
            if field_name in seen_fields:
                raise ValueError(f"Duplicate schema field name in {sheet_name}: {field_name!r}")
            seen_fields.add(field_name)
            fields.append(
                FieldSpec(
                    field_name=field_name,
                    cn_name=str(row[1] or "").strip(),
                    raw_type=str(row[2] or "").strip(),
                    description=str(row[3] or "").strip(),
                )
            )
        schema[table_name] = fields
    return schema
