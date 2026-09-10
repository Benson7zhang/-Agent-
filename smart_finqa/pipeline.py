from __future__ import annotations

import json
import os
import re
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

os.environ.setdefault("MPLCONFIGDIR", "/tmp/matplotlib")

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook

from .calculations import apply_post_compute
from .config import AppConfig, DatabaseConfig
from .core import MetadataRecognitionError, normalize_numeric, report_period_sort_key
from .database import FinanceDatabase
from .facts import (
    AuthoritativeSource,
    ExtractedFactCandidate,
    FinancialFact,
    SourceAuthorityStatus,
    StatementScope,
    UntrustedFactError,
    ValidationIssue,
    expected_period_type_for_table,
)
from .ingestion import (
    REPORT_METADATA_VERSION,
    REPORT_SELECTION_VERSION,
    CompanyIndex,
    PdfPageExtractor,
    ReportFileMetadata,
    ReportRecord,
    extract_report_record,
    inspect_report_file_metadata,
    report_metadata_from_dict,
    report_metadata_to_dict,
    select_authoritative_reports,
)
from .ingestion_state import (
    IngestionStateStore,
    InvalidIngestionStateError,
    dataset_source_uri,
    ingestion_context_fingerprint,
    sha256_file,
)
from .kb import EvidenceFilter, EvidenceStatus, SimpleKnowledgeBase
from .llm import LLMClient
from .logger import get_logger
from .monitor import ProgressTracker, ResourceMonitor
from .planner import TaskPlanner
from .quality import (
    METRIC_LABELS,
    METRIC_UNITS,
    format_reason_with_references,
    format_single_metric_answer,
    format_topn_analysis_answer,
    format_trend_analysis_answer,
)
from .safe_query import CompiledQuery, QueryAuditRecord, QueryValidationError, SafeQueryExecutor, SessionState
from .schema import FieldSpec, load_schema_from_xlsx
from .sql_planner import SQLPlanner
from .task2 import TASK2_METRICS, CompanyResolver, QuestionAnalyzer

plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["font.sans-serif"] = ["PingFang SC", "SimHei", "Arial Unicode MS", "DejaVu Sans"]

TABLE_SECTION_KEYWORDS = {
    "income_sheet": ["合并利润表", "利润表"],
    "balance_sheet": ["合并资产负债表", "资产负债表"],
    "cash_flow_sheet": ["合并现金流量表", "现金流量表"],
    "core_performance_indicators_sheet": ["主要会计数据和财务指标", "核心财务指标"],
}
TABLE_SECTION_ROW_HINTS = {
    "income_sheet": ["营业总收入", "净利润", "销售费用", "研发费用"],
    "balance_sheet": ["货币资金", "应收账款", "存货", "总资产", "短期借款"],
    "cash_flow_sheet": ["经营活动产生的现金流量净额", "销售商品、提供劳务收到的现金", "投资活动产生的现金流量净额"],
    "core_performance_indicators_sheet": [
        "营业收入",
        "归属于上市公司股东的净利润",
        "基本每股收益",
        "加权平均净资产收益率",
    ],
}

BASE_FIELDS = {"serial_number", "stock_code", "stock_abbr", "report_period", "report_year"}
INGESTION_SIGNATURE_FIELDS = ("size", "mtime_ns", "content_sha256", "context_fingerprint")
INGESTION_STATE_VERSION = 4
EXTRACTOR_VERSION = "legacy-regex-v2"
RUN_LOG_SCHEMA_VERSION = 2
VALIDATION_REPORT_SCHEMA_VERSION = 2
SAMPLED_INGESTION_STATUSES = frozenset({"parsed"})
RESEARCH_COMPANY_METADATA_KEYS = (
    "company",
    "company_name",
    "stock_name",
    "stock_abbr",
    "stockName",
    "公司",
    "公司名称",
    "股票简称",
)
RESEARCH_INDUSTRY_METADATA_KEYS = (
    "industry",
    "industry_name",
    "industryName",
    "indvInduName",
    "行业",
    "行业名称",
)
RESEARCH_PUBLICATION_METADATA_KEYS = (
    "published_at",
    "publish_time",
    "publish_date",
    "publishDate",
    "发布日期",
    "报告日期",
    "发布时间",
)


def _period_sort_key(period: str) -> tuple[int, int]:
    return report_period_sort_key(period)


def _to_bool(text: str | bool | None) -> bool:
    if isinstance(text, bool):
        return text
    return str(text or "").strip().lower() in {"1", "true", "yes", "y", "on"}


@dataclass(slots=True)
class PipelinePaths:
    base_dir: Path
    sample_dir: Path
    schema_xlsx: Path
    company_xlsx: Path
    reports_dir: Path
    task2_xlsx: Path
    task3_xlsx: Path
    research_dir: Path
    output_dir: Path
    result_dir: Path
    db_path: Path

    @classmethod
    def from_base_dir(cls, base_dir: Path, full_data: bool = False) -> "PipelinePaths":
        sample_dir = cls._resolve_dataset_dir(base_dir=base_dir, full_data=full_data)
        output_dir = base_dir / "outputs"
        result_dir = base_dir / "result"
        return cls(
            base_dir=base_dir,
            sample_dir=sample_dir,
            schema_xlsx=sample_dir / "附件3：数据库-表名及字段说明.xlsx",
            company_xlsx=cls._resolve_company_xlsx(sample_dir),
            reports_dir=sample_dir / "附件2：财务报告",
            task2_xlsx=sample_dir / "附件4：问题汇总.xlsx",
            task3_xlsx=sample_dir / "附件6：问题汇总.xlsx",
            research_dir=sample_dir / "附件5：研报数据",
            output_dir=output_dir,
            result_dir=result_dir,
            db_path=output_dir / "finance.db",
        )

    @staticmethod
    def _resolve_dataset_dir(base_dir: Path, full_data: bool) -> Path:
        sample_default = base_dir / "样例数据"
        schema_name = "附件3：数据库-表名及字段说明.xlsx"
        if not full_data:
            return sample_default if sample_default.exists() else base_dir

        direct_candidates = [
            base_dir / "正式数据",
            base_dir / "数据" / "测试数据",
            base_dir / "数据" / "全量数据",
            base_dir / "测试数据",
            base_dir / "全量数据",
            sample_default,
            base_dir,
        ]
        for candidate in direct_candidates:
            if (candidate / schema_name).exists():
                return candidate

        hits: list[Path] = []
        for schema_path in base_dir.rglob(schema_name):
            hits.append(schema_path.parent)

        for label in ("正式数据", "测试数据", "全量数据", "样例数据"):
            label_hits = [path for path in hits if label in str(path)]
            if label_hits:
                label_hits.sort(key=lambda p: (len(p.parts), len(str(p))))
                return label_hits[0]

        if hits:
            hits.sort(key=lambda p: (len(p.parts), len(str(p))))
            return hits[0]
        return sample_default if sample_default.exists() else base_dir

    @staticmethod
    def _resolve_company_xlsx(sample_dir: Path) -> Path:
        matches = sorted(sample_dir.glob("附件1：*上市公司基本信息*.xlsx"))
        if matches:
            return matches[0]
        return sample_dir / "附件1：上市公司基本信息.xlsx"


class SmartFinancePipeline:
    def __init__(
        self, paths: PipelinePaths, app_config: AppConfig | None = None, llm_client: LLMClient | None = None
    ) -> None:
        self.paths = paths
        self.app_config = app_config or AppConfig.from_env()
        self.pdf_page_extractor = PdfPageExtractor(self.app_config.ocr)
        ocr_profile = self.pdf_page_extractor.profile
        self.extractor_version = EXTRACTOR_VERSION
        if self.app_config.ocr.enabled:
            self.extractor_version += (
                f"+ocr-{ocr_profile['preprocessing_version']}-{ocr_profile['engine']}-{ocr_profile['policy']}-"
                f"{ocr_profile['dpi']}-{ocr_profile['min_page_text_chars']}-{ocr_profile['min_confidence']}-"
                f"{ocr_profile['max_page_pixels']}"
            )
        self.paths.output_dir.mkdir(parents=True, exist_ok=True)
        self.paths.result_dir.mkdir(parents=True, exist_ok=True)

        self.logger = get_logger()
        self.resource_monitor = ResourceMonitor(memory_limit_percent=85.0)

        self.schema = load_schema_from_xlsx(self.paths.schema_xlsx)
        self.company_index = CompanyIndex.from_xlsx(self.paths.company_xlsx)
        self.company_resolver = CompanyResolver.from_xlsx(self.paths.company_xlsx)

        db_cfg = self.app_config.db
        if db_cfg.backend == "sqlite" and not db_cfg.sqlite_path:
            db_cfg = DatabaseConfig(backend="sqlite", sqlite_path=str(self.paths.db_path))
        self.db = FinanceDatabase(
            self.paths.db_path,
            self.schema,
            db_config=db_cfg,
            connect_immediately=False,
            enable_cache=self.app_config.enable_cache,
            cache_size=self.app_config.cache_size,
        )
        self.db.create_tables()

        self.llm_client = llm_client or LLMClient(self.app_config.llm)
        self.planner = TaskPlanner(self.llm_client if self.llm_client.enabled else None)
        self.sql_planner = SQLPlanner()
        self.safe_query_executor = SafeQueryExecutor(
            self.db,
            self.sql_planner.registry,
            timeout_seconds=self.app_config.query_timeout_seconds,
        )
        self.task2_analyzer = QuestionAnalyzer(company_resolver=self.company_resolver)
        self.kb = SimpleKnowledgeBase(
            self.llm_client if self.llm_client.enabled else None,
            use_embeddings=self.app_config.kb_use_embeddings,
        )
        self._file_signature_cache: dict[Path, dict[str, int | str]] = {}
        self._ingestion_context_hashes: tuple[str, str] | None = None
        self._serial_counter = {table: 1 for table in self.schema}
        self.run_log = self._new_run_log(self.app_config.mode, status="NOT_STARTED")

    def _new_run_log(self, selected_mode: str, *, status: str = "RUNNING") -> dict[str, Any]:
        return {
            "schema_version": RUN_LOG_SCHEMA_VERSION,
            "run_id": str(uuid4()),
            "started_at": self._utc_timestamp() if status == "RUNNING" else None,
            "finished_at": None,
            "status": status,
            "ingestion": [],
            "ingestion_event_stats": {},
            "ingestion_overflow": 0,
            "task2": [],
            "task3": [],
            "validation": {},
            "config": {
                "mode": selected_mode,
                "full_data": self.app_config.full_data,
                "db_backend": self.db.backend,
                "llm_enabled": self.llm_client.enabled,
                "ingest_workers": self.app_config.ingest_workers,
                "incremental_ingest": self.app_config.incremental_ingest,
                "kb_max_documents": self.app_config.kb_max_documents,
                "kb_max_chunks_per_paper": self.app_config.kb_max_chunks_per_paper,
                "kb_use_embeddings": self.app_config.kb_use_embeddings,
                "query_timeout_seconds": self.app_config.query_timeout_seconds,
                "ocr": self.pdf_page_extractor.profile,
            },
        }

    @staticmethod
    def _utc_timestamp() -> str:
        return datetime.now(timezone.utc).isoformat(timespec="milliseconds")

    def run(self, mode: str | None = None) -> dict[str, str]:
        selected_mode = (mode or self.app_config.mode or "all").strip().lower()
        if selected_mode not in {"all", "ingest", "task2", "task3"}:
            raise ValueError(f"Unsupported pipeline mode: {selected_mode}")
        self.run_log = self._new_run_log(selected_mode)
        outputs = {
            "db_path": str(self.paths.db_path),
            "result_2": "",
            "result_3": "",
            "run_log": "",
            "validation_report": "",
        }

        log_path = self.paths.output_dir / "run_log.json"
        try:
            if selected_mode in {"all", "ingest"}:
                self.ingest_reports()
                validation = self.validate_database()
                self.run_log["validation"] = {
                    **validation,
                    "schema_version": VALIDATION_REPORT_SCHEMA_VERSION,
                    "run_id": self.run_log["run_id"],
                    "generated_at": self._utc_timestamp(),
                }
                validation_path = self.paths.output_dir / "validation_report.json"
                validation_path.write_text(
                    json.dumps(self.run_log["validation"], ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                outputs["validation_report"] = str(validation_path)

            if selected_mode in {"all", "task2", "task3"}:
                self._ensure_database_ready()

            if selected_mode in {"all", "task3"}:
                self.build_knowledge_base()

            if selected_mode in {"all", "task2"}:
                outputs["result_2"] = str(self.answer_task2())
            if selected_mode in {"all", "task3"}:
                outputs["result_3"] = str(self.answer_task3())
            self.run_log["status"] = "SUCCEEDED"
            return outputs
        except Exception as exc:
            self.run_log["status"] = "FAILED"
            self.run_log["failure"] = {"type": type(exc).__name__, "message": str(exc)}
            raise
        finally:
            self.run_log["finished_at"] = self._utc_timestamp()
            log_path.write_text(json.dumps(self.run_log, ensure_ascii=False, indent=2), encoding="utf-8")
            outputs["run_log"] = str(log_path)

    # ---------- Task 1 ----------
    def ingest_reports(self) -> None:
        self._file_signature_cache.clear()
        self._ingestion_context_hashes = (
            sha256_file(self.paths.schema_xlsx),
            sha256_file(self.paths.company_xlsx),
        )
        report_files = sorted(self.paths.reports_dir.rglob("*.pdf"))
        total_files = len(report_files)
        self.logger.info(f"Found {total_files} PDF files to process")
        projection_rows: dict[tuple[str, str, str], dict[str, Any]] = {}
        extracted_facts: list[FinancialFact] = []
        state_path = self.paths.output_dir / "ingestion_state.json"
        state = self._load_ingestion_state(state_path)

        metadata_batch_size = max(10, min(100, max(1, total_files // 10)))
        workers = max(1, min(self.app_config.ingest_workers, os.cpu_count() or 4))
        parse_batch_size = max(1, min(16, workers * 2))
        inspected_reports: list[ReportFileMetadata] = []
        inspect_candidates: list[Path] = []
        metadata_failures = 0
        reused_metadata = 0
        for pdf_path in report_files:
            signature = self._file_signature(pdf_path)
            cached = state.get(str(pdf_path))
            if (
                self.app_config.incremental_ingest
                and cached
                and self._matches_ingestion_signature(cached, signature)
                and cached.get("metadata_version") == REPORT_METADATA_VERSION
                and cached.get("selection_version") == REPORT_SELECTION_VERSION
            ):
                try:
                    inspected_reports.append(report_metadata_from_dict(pdf_path, cached))
                except ValueError:
                    inspect_candidates.append(pdf_path)
                else:
                    reused_metadata += 1
                    continue
            else:
                inspect_candidates.append(pdf_path)

        inspection_progress = ProgressTracker(len(inspect_candidates), "Inspecting report metadata")
        for batch_start in range(0, len(inspect_candidates), metadata_batch_size):
            batch = inspect_candidates[batch_start : batch_start + metadata_batch_size]
            with ThreadPoolExecutor(max_workers=workers) as pool:
                future_map = {
                    pool.submit(
                        inspect_report_file_metadata,
                        path,
                        self.company_index,
                        page_extractor=self.pdf_page_extractor,
                    ): path
                    for path in batch
                }
                for future in as_completed(future_map):
                    pdf_path = future_map[future]
                    try:
                        inspected_reports.append(future.result())
                    except MetadataRecognitionError as exc:
                        metadata_failures += 1
                        item = {
                            **self._file_signature(pdf_path),
                            "status": "needs_review",
                            "error_code": "metadata_recognition_failed",
                            "field": exc.field,
                            "error": str(exc),
                            "metadata_version": REPORT_METADATA_VERSION,
                            "selection_version": REPORT_SELECTION_VERSION,
                        }
                        state[str(pdf_path)] = item
                        self._append_ingestion_log({"pdf": str(pdf_path), **item})
                        self.logger.warning(f"Report requires metadata review: {pdf_path.name} ({exc.field})")
                    inspection_progress.update()
        inspection_progress.finish()

        decisions = select_authoritative_reports(inspected_reports)
        selection_counts: dict[str, int] = {}
        selected_reports: list[ReportFileMetadata] = []
        for decision in decisions:
            selection_counts[decision.status] = selection_counts.get(decision.status, 0) + 1
            report = decision.report
            item = {
                **self._file_signature(report.path),
                **report_metadata_to_dict(report),
                "status": decision.status,
                "reason": decision.reason,
                "metadata_version": REPORT_METADATA_VERSION,
                "selection_version": REPORT_SELECTION_VERSION,
            }
            if decision.status == "needs_review":
                item["error_code"] = decision.reason
            if report.read_error:
                item["read_error"] = report.read_error
            existing = state.get(str(report.path))
            preserve_parsed_state = (
                decision.status == "selected"
                and existing is not None
                and existing.get("status") == "parsed"
                and self._matches_ingestion_signature(existing, item)
                and existing.get("extractor_version") == self.extractor_version
                and existing.get("selection_version") == REPORT_SELECTION_VERSION
            )
            if not preserve_parsed_state:
                state[str(report.path)] = item
            if decision.status == "selected":
                selected_reports.append(report)
            else:
                self._append_ingestion_log({"pdf": str(report.path), **item})

        selected_source_uris = {
            report.path: dataset_source_uri(self.paths.sample_dir, report.path) for report in selected_reports
        }
        selected_sources = {
            report.path: AuthoritativeSource(
                source_file=selected_source_uris[report.path],
                source_content_sha256=str(self._file_signature(report.path)["content_sha256"]),
                stock_code=report.stock_code,
                period=report.report_period,
                selection_version=REPORT_SELECTION_VERSION,
            )
            for report in selected_reports
        }

        # Preserve completed metadata work if full-document extraction is interrupted.
        self._save_ingestion_state(state_path, state)

        selected_by_path = {report.path: report for report in selected_reports}

        parse_candidates: list[ReportFileMetadata] = []
        skipped_unchanged = 0
        for report in selected_reports:
            signature = self._file_signature(report.path)
            prev = state.get(str(report.path))
            if (
                self.app_config.incremental_ingest
                and prev
                and self._matches_ingestion_signature(prev, signature)
                and prev.get("status") == "parsed"
                and prev.get("extractor_version") == self.extractor_version
                and prev.get("selection_version") == REPORT_SELECTION_VERSION
                and self._has_complete_period_record(
                    str(prev.get("stock_code", "")),
                    str(prev.get("report_period", "")),
                    prev.get("available_tables") if "available_tables" in prev else None,
                )
            ):
                skipped_unchanged += 1
                continue
            parse_candidates.append(report)

        self.logger.info(f"Skipped {skipped_unchanged} unchanged files, processing {len(parse_candidates)} files")

        if not parse_candidates:
            self.logger.info("No files to process")
            self._verify_authoritative_inputs(selected_sources)
            self.db.reconcile_batch_authoritative_sources(
                list(selected_sources.values()),
                selection_version=REPORT_SELECTION_VERSION,
            )
            self._save_ingestion_state(state_path, state)
            resource_summary = self.resource_monitor.get_summary()
            self.run_log["ingestion_summary"] = {
                "total_files": total_files,
                "metadata_failures": metadata_failures,
                "reused_metadata": reused_metadata,
                "selection": selection_counts,
                "parse_batch_size": parse_batch_size,
                "parsed_files": 0,
                "skipped_unchanged": skipped_unchanged,
                "candidate_facts_upserted": 0,
                "projection_seed_rows_upserted": 0,
                "projection_seed_rows_by_table": {},
                "projection_rows_with_validated_metrics": 0,
                "projection_rows_with_validated_metrics_by_table": {},
                "validated_metric_cells_projected": 0,
                "validated_metric_cells_projected_by_table": {},
                "fact_status": self._fact_status_counts(),
                "resource_usage": resource_summary,
            }
            return

        # Keep full PDF text and page objects bounded to one batch. Facts and
        # projection rows remain in memory until the authoritative write below.
        progress = ProgressTracker(len(parse_candidates), "Processing PDFs")
        successfully_parsed = 0
        pending_commit_paths: list[Path] = []

        for batch_start in range(0, len(parse_candidates), parse_batch_size):
            batch = parse_candidates[batch_start : batch_start + parse_batch_size]
            batch_records: list[tuple[Path, ReportRecord]] = []

            with ThreadPoolExecutor(max_workers=workers) as pool:
                future_map = {
                    pool.submit(
                        extract_report_record,
                        report.path,
                        self.company_index,
                        report,
                        page_extractor=self.pdf_page_extractor,
                    ): report.path
                    for report in batch
                }
                for fut in as_completed(future_map):
                    pdf_path = future_map[fut]
                    try:
                        record = fut.result()
                        batch_records.append((pdf_path, record))
                    except MetadataRecognitionError as exc:
                        report = selected_by_path[pdf_path]
                        state[str(pdf_path)] = {
                            **self._file_signature(pdf_path),
                            **report_metadata_to_dict(report),
                            "status": "needs_review",
                            "error_code": "metadata_recognition_failed",
                            "field": exc.field,
                            "error": str(exc),
                            "metadata_version": REPORT_METADATA_VERSION,
                            "selection_version": REPORT_SELECTION_VERSION,
                        }
                        self._append_ingestion_log(
                            {
                                "pdf": str(pdf_path),
                                "status": "needs_review",
                                "error_code": "metadata_recognition_failed",
                                "field": exc.field,
                                "error": str(exc),
                            }
                        )
                        self.logger.warning(f"Report requires metadata review: {pdf_path.name} ({exc.field})")
                        progress.update()
                    except Exception as exc:
                        report = selected_by_path[pdf_path]
                        state[str(pdf_path)] = {
                            **self._file_signature(pdf_path),
                            **report_metadata_to_dict(report),
                            "status": "error",
                            "error": str(exc),
                            "metadata_version": REPORT_METADATA_VERSION,
                            "selection_version": REPORT_SELECTION_VERSION,
                        }
                        self._append_ingestion_log({"pdf": str(pdf_path), "error": str(exc), "status": "error"})
                        self.logger.error(f"Failed to parse {pdf_path.name}: {exc}")
                        progress.update()

            # Futures retain their results; release them before consuming the
            # batch so cleared records can be collected before the next batch.
            future_map.clear()
            del future_map, fut, pool

            while batch_records:
                pdf_path, record = batch_records.pop()
                report = selected_by_path[pdf_path]
                if self._consume_parsed_record(
                    pdf_path,
                    record,
                    report,
                    selected_source_uris[report.path],
                    state,
                    projection_rows,
                    extracted_facts,
                ):
                    successfully_parsed += 1
                    pending_commit_paths.append(pdf_path)
                progress.update()
                del record

            del batch_records
            self._save_ingestion_state(state_path, state)

            if self.resource_monitor.should_gc():
                gc_stats = self.resource_monitor.force_gc()
                self.logger.info(
                    f"Memory cleanup: freed {gc_stats['freed_mb']:.1f}MB ({gc_stats['freed_percent']:.1f}%)"
                )

        progress.finish()

        self._verify_pending_commit_inputs(
            pending_commit_paths,
            selected_sources,
            state,
            extracted_facts,
        )
        self.db.reconcile_batch_authoritative_sources(
            list(selected_sources.values()),
            selection_version=REPORT_SELECTION_VERSION,
        )

        # Persist candidates first, then rebuild wide rows from reviewed facts only.
        grouped_rows: dict[str, list[dict[str, Any]]] = {}
        for (table, _stock, _period), row in projection_rows.items():
            grouped_rows.setdefault(table, []).append(row)

        self.logger.info(f"Inserting {sum(len(rows) for rows in grouped_rows.values())} rows to database")
        self.db.upsert_financial_facts(extracted_facts)
        for table, rows in grouped_rows.items():
            for row in rows:
                self.db.refresh_fact_projection(table, row)
            self.logger.info(f"Refreshed {len(rows)} validated-fact projections in {table}")
        seed_rows_by_table = {table: len(rows) for table, rows in grouped_rows.items()}
        projected_rows_by_table, projected_cells_by_table = self._projection_validated_metric_counts(grouped_rows)

        for pdf_path in pending_commit_paths:
            item = state[str(pdf_path)]
            item["status"] = "parsed"
            self._append_ingestion_log(
                {
                    "pdf": str(pdf_path),
                    "stock_code": item["stock_code"],
                    "stock_abbr": item["stock_abbr"],
                    "report_period": item["report_period"],
                    "status": "parsed",
                }
            )

        # Mark files parsed only after all authoritative writes complete.
        self._save_ingestion_state(state_path, state)

        # Final resource check
        resource_summary = self.resource_monitor.get_summary()
        self.logger.info(f"Resource usage: {resource_summary}")

        self.run_log["ingestion_summary"] = {
            "total_files": len(report_files),
            "metadata_failures": metadata_failures,
            "reused_metadata": reused_metadata,
            "selection": selection_counts,
            "parse_batch_size": parse_batch_size,
            "parsed_files": successfully_parsed,
            "skipped_unchanged": skipped_unchanged,
            "candidate_facts_upserted": len(extracted_facts),
            "projection_seed_rows_upserted": sum(seed_rows_by_table.values()),
            "projection_seed_rows_by_table": seed_rows_by_table,
            "projection_rows_with_validated_metrics": sum(projected_rows_by_table.values()),
            "projection_rows_with_validated_metrics_by_table": projected_rows_by_table,
            "validated_metric_cells_projected": sum(projected_cells_by_table.values()),
            "validated_metric_cells_projected_by_table": projected_cells_by_table,
            "fact_status": self._fact_status_counts(),
            "resource_usage": resource_summary,
        }

    def _verify_pending_commit_inputs(
        self,
        pending_commit_paths: list[Path],
        selected_sources: dict[Path, AuthoritativeSource],
        state: dict[str, dict[str, Any]],
        extracted_facts: list[FinancialFact],
    ) -> None:
        self._verify_authoritative_inputs(selected_sources)

        expected_digest_by_source: dict[str, str] = {}
        source_path_by_uri: dict[str, Path] = {}
        for pdf_path in pending_commit_paths:
            item = state.get(str(pdf_path))
            if item is None or item.get("status") != "parsed_pending_commit":
                raise RuntimeError(f"source is not pending authoritative commit: {pdf_path}")
            expected_digest = item.get("content_sha256")
            if not isinstance(expected_digest, str):
                raise RuntimeError(f"pending source state has no content SHA-256: {pdf_path}")
            selected_source = selected_sources.get(pdf_path)
            if selected_source is None:
                raise RuntimeError(f"pending source has no authoritative selection: {pdf_path}")
            if selected_source.source_content_sha256 != expected_digest:
                raise RuntimeError(f"pending source state hash differs from authoritative selection: {pdf_path}")
            expected_digest_by_source[selected_source.source_file] = expected_digest
            source_path_by_uri[selected_source.source_file] = pdf_path

        for fact in extracted_facts:
            expected_digest = expected_digest_by_source.get(fact.source_file)
            if expected_digest is None:
                raise RuntimeError(f"fact references a source outside the pending commit: {fact.source_file}")
            if fact.source_content_sha256 != expected_digest:
                pdf_path = source_path_by_uri[fact.source_file]
                raise RuntimeError(f"fact source hash differs from pending source state: {pdf_path}")

    def _verify_authoritative_inputs(self, selected_sources: dict[Path, AuthoritativeSource]) -> None:
        if self._ingestion_context_hashes is None:
            raise RuntimeError("ingestion context hashes are not initialized")
        expected_schema_digest, expected_company_digest = self._ingestion_context_hashes
        if sha256_file(self.paths.schema_xlsx) != expected_schema_digest:
            raise RuntimeError(f"schema workbook changed before authoritative commit: {self.paths.schema_xlsx}")
        if sha256_file(self.paths.company_xlsx) != expected_company_digest:
            raise RuntimeError(f"company workbook changed before authoritative commit: {self.paths.company_xlsx}")

        for pdf_path, selected_source in selected_sources.items():
            current_digest = sha256_file(pdf_path)
            if current_digest != selected_source.source_content_sha256:
                raise RuntimeError(f"source file changed before authoritative commit: {pdf_path}")

    def _consume_parsed_record(
        self,
        pdf_path: Path,
        record: ReportRecord,
        report: ReportFileMetadata,
        source_file: str,
        state: dict[str, dict[str, Any]],
        projection_rows: dict[tuple[str, str, str], dict[str, Any]],
        extracted_facts: list[FinancialFact],
    ) -> bool:
        try:
            source_signature = self._file_signature(pdf_path)
            if sha256_file(pdf_path) != source_signature["content_sha256"]:
                raise RuntimeError(f"source file changed during ingestion: {pdf_path}")
            if len(record.text.strip()) < 80:
                item = {
                    **self._file_signature(pdf_path),
                    **report_metadata_to_dict(report),
                    "status": "needs_review",
                    "error_code": "document_text_unavailable",
                    "reason": "document_text_unavailable",
                    "metadata_version": REPORT_METADATA_VERSION,
                    "selection_version": REPORT_SELECTION_VERSION,
                }
                state[str(pdf_path)] = item
                self._append_ingestion_log({"pdf": str(pdf_path), **item})
                return False

            table_rows = self._build_rows_from_record(record)
            record_facts = self._build_facts_from_rows(
                record,
                table_rows,
                source_file,
                str(source_signature["content_sha256"]),
            )
            record_projection_rows: dict[tuple[str, str, str], dict[str, Any]] = {}
            for table, row in table_rows.items():
                row = self._sanitize_row_values(table, row)
                key = (table, str(row.get("stock_code")), str(row.get("report_period")))
                record_projection_rows[key] = {field: row.get(field) for field in BASE_FIELDS if field in row}
                ok, issues = self._validate_row(table, row)
                if not ok:
                    self._append_ingestion_log(
                        {
                            "pdf": str(pdf_path),
                            "table": table,
                            "stock_code": record.stock_code,
                            "report_period": record.report_period,
                            "status": "skip_invalid",
                            "issues": issues,
                        }
                    )

            extracted_facts.extend(record_facts)
            projection_rows.update(record_projection_rows)
            state[str(pdf_path)] = {
                **self._file_signature(pdf_path),
                "status": "parsed_pending_commit",
                **report_metadata_to_dict(report),
                "metadata_version": REPORT_METADATA_VERSION,
                "selection_version": REPORT_SELECTION_VERSION,
                "extractor_version": self.extractor_version,
                "available_tables": self._available_tables_in_text(record.text),
            }
            return True
        except Exception as exc:
            self.logger.exception(f"Failed to build facts and rows for {pdf_path.name}")
            state[str(pdf_path)] = {
                **self._file_signature(pdf_path),
                **report_metadata_to_dict(report),
                "status": "error",
                "error": f"Row building failed: {exc}",
                "metadata_version": REPORT_METADATA_VERSION,
                "selection_version": REPORT_SELECTION_VERSION,
            }
            return False

    def _build_rows_from_record(self, record: ReportRecord) -> dict[str, dict[str, Any]]:
        rows: dict[str, dict[str, Any]] = {}
        for table, fields in self.schema.items():
            row: dict[str, Any] = {f.field_name: None for f in fields}
            row["serial_number"] = self._serial_counter[table]
            self._serial_counter[table] += 1
            row["stock_code"] = record.stock_code
            row["stock_abbr"] = record.stock_abbr
            row["report_period"] = record.report_period
            row["report_year"] = record.report_year

            section_text = self._locate_table_section(record.text, table)
            section_lines = [ln.strip() for ln in section_text.splitlines() if ln.strip()]
            for field in fields:
                if field.field_name in BASE_FIELDS:
                    continue
                value = self._extract_field_value(section_text, section_lines, field)
                row[field.field_name] = value

            self._overlay_snapshot_values(table, row, record.snapshot)
            rows[table] = row
        return rows

    def _build_facts_from_rows(
        self,
        record: ReportRecord,
        table_rows: dict[str, dict[str, Any]],
        source_file: str,
        source_content_sha256: str,
    ) -> list[FinancialFact]:
        facts: list[FinancialFact] = []
        for table, row in table_rows.items():
            fields = {field.field_name: field for field in self.schema[table]}
            for metric, value in row.items():
                if metric in BASE_FIELDS or value is None or metric not in fields:
                    continue
                field = fields[metric]
                page_no = self._find_source_page(record, field)
                unit = self._target_unit(field)
                issues = (
                    ValidationIssue(
                        code="legacy_heuristic_extraction",
                        message="该数值来自尚未完成表格列结构恢复的兼容抽取器",
                        field="normalized_value",
                    ),
                )
                candidate = ExtractedFactCandidate(
                    metric=metric,
                    raw_value=value,
                    source_unit=unit,
                    page_no=page_no,
                    table_name=table,
                    row_label=field.cn_name or field.description or metric,
                    column_label=None,
                    confidence=0.5 if page_no is not None else 0.25,
                )
                facts.append(
                    FinancialFact.from_candidate(
                        candidate,
                        company_id=record.stock_code,
                        stock_code=record.stock_code,
                        period=record.report_period,
                        statement_scope=StatementScope.CONSOLIDATED,
                        period_type=expected_period_type_for_table(table),
                        target_unit=unit,
                        currency="CNY",
                        source_file=source_file,
                        source_content_sha256=source_content_sha256,
                        extractor_version=self.extractor_version,
                        validation_issues=issues,
                    )
                )
        return facts

    def _find_source_page(self, record: ReportRecord, field: FieldSpec) -> int | None:
        aliases = self._field_aliases(field)
        for page in record.pages:
            if any(alias in page.text for alias in aliases):
                return page.page_no
        return None

    @staticmethod
    def _target_unit(field: FieldSpec) -> str:
        if (
            "%" in field.cn_name
            or "ratio" in field.field_name
            or "yoy" in field.field_name
            or "qoq" in field.field_name
        ):
            return "%"
        if "万元" in field.cn_name:
            return "万元"
        return "元"

    def _locate_table_section(self, text: str, table: str) -> str:
        candidates = TABLE_SECTION_KEYWORDS.get(table, [])
        positions: list[tuple[int, str]] = []
        for marker in candidates:
            positions.extend((match.start(), marker) for match in re.finditer(re.escape(marker), text))
        if not positions:
            return text

        start = max(
            (idx for idx, _marker in positions),
            key=lambda idx: (self._score_table_section_candidate(text, table, idx), idx),
        )
        next_cut = len(text)
        for other_table, markers in TABLE_SECTION_KEYWORDS.items():
            if other_table == table:
                continue
            for marker in markers:
                idx = text.find(marker, start + 1)
                if idx > start:
                    next_cut = min(next_cut, idx)
        return text[start:next_cut]

    def _score_table_section_candidate(self, text: str, table: str, start: int) -> int:
        window = text[start : start + 8000]
        preview = window[:300]
        hints = TABLE_SECTION_ROW_HINTS.get(table, [])
        score = 0
        score += sum(120 for hint in hints if hint in window)
        score += min(len(re.findall(r"-?\d{1,3}(?:,\d{3})*(?:\.\d+)?|-?\d+\.\d+|-?\d+", window[:4000])), 80)
        score += 40 if "编制单位" in preview else 0
        score += 40 if "单位" in preview else 0
        score += 40 if "项目" in preview else 0
        score -= 160 if "和母公司" in preview else 0
        score -= 120 if "审计报告" in preview else 0
        score -= 80 if "关键审计事项" in preview else 0
        return score

    def _extract_field_value(self, text: str, lines: list[str], field: FieldSpec) -> float | int | None:
        aliases = self._field_aliases(field)
        candidates: list[tuple[float, str]] = []
        for alias in aliases:
            # Fast path: line-level scan avoids repeatedly scanning whole section.
            matched = False
            for idx, line in enumerate(lines):
                if alias not in line:
                    continue
                matched = True
                snippet = line
                start_at = snippet.find(alias) + len(alias)
                nums = re.findall(r"-?\d{1,3}(?:,\d{3})*(?:\.\d+)?|-?\d+\.\d+|-?\d+", snippet[start_at:])
                if not nums and idx + 1 < len(lines):
                    extended = " ".join(lines[idx : min(idx + 3, len(lines))])
                    extended_start = extended.find(alias) + len(alias)
                    extended_nums = re.findall(
                        r"-?\d{1,3}(?:,\d{3})*(?:\.\d+)?|-?\d+\.\d+|-?\d+",
                        extended[extended_start:],
                    )
                    if extended_nums:
                        snippet = extended
                        nums = extended_nums
                if nums:
                    try:
                        candidates.append((float(nums[0].replace(",", "")), snippet[:120]))
                    except ValueError:
                        pass
            if matched:
                continue

            # Fallback to section-level regex when line segmentation misses merged rows.
            for match in re.finditer(re.escape(alias), text):
                snippet = text[match.start() : match.start() + 220]
                nums = re.findall(r"-?\d{1,3}(?:,\d{3})*(?:\.\d+)?|-?\d+\.\d+|-?\d+", snippet)
                for token in nums[:4]:
                    try:
                        val = float(token.replace(",", ""))
                    except ValueError:
                        continue
                    candidates.append((val, snippet[:120]))

        if not candidates:
            return None
        chosen = self._pick_best_candidate(field, candidates)
        return self._normalize_field_value(field, chosen)

    def _field_aliases(self, field: FieldSpec) -> list[str]:
        cn = field.cn_name or ""
        aliases = {cn}
        stripped = re.sub(r"（.*?）|\(.*?\)", "", cn).strip()
        aliases.add(stripped)
        aliases.add(cn.replace("同比增长", "同比").replace("季度环比增长", "环比"))
        aliases.add(stripped.replace("同比增长", "同比").replace("季度环比增长", "环比"))
        aliases.add(cn.replace("营业总支出", "营业总成本"))
        aliases.add(cn.replace("净利润", "归属于上市公司股东的净利润"))
        for token in (cn, stripped):
            for sep in ("-", "—", "–"):
                if sep in token:
                    _left, right = token.split(sep, 1)
                    aliases.add(right.strip())
        if "经营性现金流" in stripped:
            aliases.add("经营活动产生的现金流量净额")
        if "融资性现金流" in stripped:
            aliases.add("融资活动产生的现金流量净额")
        if "投资性现金流" in stripped:
            aliases.add("投资活动产生的现金流量净额")
        aliases = {x for x in aliases if x}
        return sorted(aliases, key=len, reverse=True)

    def _pick_best_candidate(self, field: FieldSpec, candidates: list[tuple[float, str]]) -> float:
        # Use LLM disambiguation when available and multiple strong candidates exist.
        if self.llm_client.enabled and len(candidates) > 1:
            prompt = {
                "field_name": field.field_name,
                "cn_name": field.cn_name,
                "candidates": [{"value": c[0], "context": c[1]} for c in candidates[:6]],
            }
            ans = self.llm_client.complete_json(
                '你是财报字段数值裁决器。返回JSON: {"value": 数值}。',
                json.dumps(prompt, ensure_ascii=False),
            )
            if not isinstance(ans, dict) or "value" not in ans:
                raise RuntimeError("Invalid LLM extraction response: missing value")
            return float(ans["value"])

        values = [c[0] for c in candidates]
        if "yoy" in field.field_name or "qoq" in field.field_name or "%" in field.cn_name:
            valid = [v for v in values if -1000 <= v <= 1000]
            return valid[0] if valid else values[0]
        if "ratio" in field.field_name:
            valid = [v for v in values if 0 <= v <= 100]
            return valid[0] if valid else values[0]
        # For amounts pick the largest absolute among nearby candidates.
        return sorted(values, key=lambda v: abs(v), reverse=True)[0]

    def _normalize_field_value(self, field: FieldSpec, value: float) -> float | int | None:
        if value is None:
            return None
        field_type = field.raw_type.lower()
        if "int" in field_type:
            return int(round(value))
        cn = field.cn_name
        if "万元" in cn:
            return normalize_numeric(value, source_unit="元", target_unit="万元")
        if "元" in cn:
            return normalize_numeric(value, source_unit="元", target_unit="元")
        if "%" in cn:
            return round(float(value), 4)
        return round(float(value), 4)

    def _overlay_snapshot_values(self, table: str, row: dict[str, Any], snapshot: dict[str, float | None]) -> None:
        if table == "income_sheet":
            if snapshot.get("total_profit_yuan") is not None:
                row["total_profit"] = normalize_numeric(snapshot["total_profit_yuan"], "元", "万元")
            if snapshot.get("net_profit_yuan") is not None:
                row["net_profit"] = normalize_numeric(snapshot["net_profit_yuan"], "元", "万元")
            if snapshot.get("total_operating_revenue_yuan") is not None:
                row["total_operating_revenue"] = normalize_numeric(
                    snapshot["total_operating_revenue_yuan"], "元", "万元"
                )
            if snapshot.get("operating_revenue_yoy") is not None:
                row["operating_revenue_yoy_growth"] = snapshot["operating_revenue_yoy"]
            if snapshot.get("net_profit_yoy") is not None:
                row["net_profit_yoy_growth"] = snapshot["net_profit_yoy"]
        elif table == "core_performance_indicators_sheet":
            if snapshot.get("eps") is not None:
                row["eps"] = snapshot["eps"]
            if snapshot.get("total_operating_revenue_yuan") is not None:
                row["total_operating_revenue"] = normalize_numeric(
                    snapshot["total_operating_revenue_yuan"], "元", "万元"
                )
            if snapshot.get("net_profit_yuan") is not None:
                row["net_profit_10k_yuan"] = normalize_numeric(snapshot["net_profit_yuan"], "元", "万元")
            if snapshot.get("net_profit_yoy") is not None:
                row["net_profit_yoy_growth"] = snapshot["net_profit_yoy"]
            if snapshot.get("operating_revenue_yoy") is not None:
                row["operating_revenue_yoy_growth"] = snapshot["operating_revenue_yoy"]
        elif table == "cash_flow_sheet":
            if snapshot.get("operating_cf_net_yuan") is not None:
                row["operating_cf_net_amount"] = normalize_numeric(snapshot["operating_cf_net_yuan"], "元", "万元")
                row["net_cash_flow"] = snapshot["operating_cf_net_yuan"]
            if snapshot.get("operating_cf_yoy") is not None:
                row["net_cash_flow_yoy_growth"] = snapshot["operating_cf_yoy"]
        elif table == "balance_sheet":
            if snapshot.get("asset_liability_ratio") is not None:
                row["asset_liability_ratio"] = snapshot["asset_liability_ratio"]

    def _sanitize_row_values(self, table: str, row: dict[str, Any]) -> dict[str, Any]:
        cleaned = dict(row)
        for key, value in row.items():
            if value is None:
                continue
            if ("yoy" in key or "qoq" in key) and (
                not isinstance(value, (int, float)) or value < -1000 or value > 1000
            ):
                cleaned[key] = None
            if key == "asset_liability_ratio" and (not isinstance(value, (int, float)) or value < 0 or value > 100):
                cleaned[key] = None
        return cleaned

    @staticmethod
    def _available_tables_in_text(text: str) -> list[str]:
        return [table for table, markers in TABLE_SECTION_KEYWORDS.items() if any(marker in text for marker in markers)]

    def _projection_validated_metric_counts(
        self,
        grouped_rows: dict[str, list[dict[str, Any]]],
    ) -> tuple[dict[str, int], dict[str, int]]:
        rows_by_table: dict[str, int] = {}
        cells_by_table: dict[str, int] = {}
        for table, seed_rows in grouped_rows.items():
            touched_keys = {
                (str(row.get("stock_code") or ""), str(row.get("report_period") or "")) for row in seed_rows
            }
            metric_fields = [field.field_name for field in self.schema[table] if field.field_name not in BASE_FIELDS]
            rows_with_metrics = 0
            projected_cells = 0
            if metric_fields and touched_keys:
                columns = ", ".join(("stock_code", "report_period", *metric_fields))
                for row in self.db.query(f"SELECT {columns} FROM {table}", use_cache=False):
                    key = (str(row["stock_code"]), str(row["report_period"]))
                    if key not in touched_keys:
                        continue
                    non_null_cells = sum(row.get(metric) is not None for metric in metric_fields)
                    if non_null_cells:
                        rows_with_metrics += 1
                        projected_cells += non_null_cells
            rows_by_table[table] = rows_with_metrics
            cells_by_table[table] = projected_cells
        return rows_by_table, cells_by_table

    def _validate_row(self, table: str, row: dict[str, Any]) -> tuple[bool, list[str]]:
        issues: list[str] = []
        business_fields = [k for k in row if k not in BASE_FIELDS]
        non_null = sum(1 for k in business_fields if row.get(k) is not None)
        ratio = non_null / max(len(business_fields), 1)
        if ratio < 0.12:
            issues.append(f"non_null_ratio_too_low:{ratio:.3f}")

        for key, value in row.items():
            if value is None:
                continue
            if ("yoy" in key or "qoq" in key) and (
                not isinstance(value, (int, float)) or value < -1000 or value > 1000
            ):
                issues.append(f"invalid_growth:{key}")
            if key == "asset_liability_ratio" and (not isinstance(value, (int, float)) or value < 0 or value > 100):
                issues.append("invalid_asset_liability_ratio")

        return (len(issues) == 0, issues)

    def validate_database(self) -> dict[str, Any]:
        fact_status = self._fact_status_counts()
        consistency_inputs = self._financial_consistency_inputs()
        financial_consistency = self._financial_consistency_issues(consistency_inputs)
        validation_coverage = self._validation_coverage(consistency_inputs)
        validation_execution = {
            "eligible_records": sum(int(item["eligible_records"]) for item in validation_coverage.values()),
            "rules_with_eligible_records": sum(
                int(int(item["eligible_records"]) > 0) for item in validation_coverage.values()
            ),
        }
        result: dict[str, Any] = {
            "unique": {},
            "cross_table": [],
            "financial_consistency": financial_consistency,
            "fact_status": fact_status,
            "all_fact_status": self._fact_status_counts(current_only=False),
            "query_readiness": self._query_readiness(fact_status),
            "fact_traceability": self._fact_traceability(),
            "validation_coverage": validation_coverage,
            "validation_execution": validation_execution,
        }
        for table in self.schema:
            dup = self.db.query(
                f"SELECT stock_code, report_period, COUNT(*) as cnt FROM {table} "
                "GROUP BY stock_code, report_period HAVING cnt > 1"
            )
            result["unique"][table] = len(dup) == 0

        presence_rows = self.db.query(
            "SELECT stock_code, report_period, 'income_sheet' AS table_name FROM income_sheet "
            "UNION ALL "
            "SELECT stock_code, report_period, 'balance_sheet' AS table_name FROM balance_sheet "
            "UNION ALL "
            "SELECT stock_code, report_period, 'cash_flow_sheet' AS table_name FROM cash_flow_sheet "
            "UNION ALL "
            "SELECT stock_code, report_period, 'core_performance_indicators_sheet' AS table_name FROM core_performance_indicators_sheet"
        )
        presence: dict[tuple[str, str], set[str]] = {}
        for row in presence_rows:
            key = (str(row["stock_code"]), str(row["report_period"]))
            presence.setdefault(key, set()).add(str(row["table_name"]))

        required_by_period = self._required_tables_by_period()
        for key in sorted(set(presence) | set(required_by_period)):
            stock_code, report_period = key
            required = required_by_period.get(key)
            if required is None:
                required = set(self.schema.keys())
            missing = sorted(required - presence.get(key, set()))
            if missing:
                result["cross_table"].append(
                    {
                        "stock_code": stock_code,
                        "report_period": report_period,
                        "missing_tables": missing,
                    }
                )
        if financial_consistency or result["cross_table"] or not all(result["unique"].values()):
            result["status"] = "FAILED"
        elif not fact_status:
            result["status"] = "UNAVAILABLE"
        elif fact_status.get("NEEDS_REVIEW", 0) or fact_status.get("REJECTED", 0):
            result["status"] = "NEEDS_REVIEW"
        elif validation_execution["eligible_records"] == 0:
            result["status"] = "VALIDATION_UNAVAILABLE"
        else:
            result["status"] = "PASS"
        return result

    def _fact_status_counts(self, *, current_only: bool = True) -> dict[str, int]:
        if current_only:
            rows = self.db.query(
                "SELECT ff.validation_status, COUNT(*) AS cnt FROM financial_fact ff "
                "INNER JOIN financial_source fs ON fs.source_key = ff.source_key "
                "AND fs.stock_code = ff.stock_code AND fs.period = ff.period "
                f"WHERE fs.authority_status = {self.db.parameter_marker} GROUP BY ff.validation_status",
                (SourceAuthorityStatus.CURRENT.value,),
            )
        else:
            rows = self.db.query(
                "SELECT validation_status, COUNT(*) AS cnt FROM financial_fact GROUP BY validation_status"
            )
        return {str(row["validation_status"]): int(row["cnt"]) for row in rows}

    @staticmethod
    def _query_readiness(fact_status: dict[str, int]) -> dict[str, Any]:
        validated_fact_count = fact_status.get("VALIDATED", 0)
        if validated_fact_count <= 0:
            return {
                "status": "BLOCKED",
                "validated_fact_count": 0,
                "reason": "no VALIDATED financial facts; complete human review before controlled queries",
            }
        return {
            "status": "READY",
            "validated_fact_count": validated_fact_count,
            "reason": "VALIDATED financial facts are available for controlled queries",
        }

    def _fact_traceability(self) -> dict[str, int | float]:
        rows = self.db.query(
            "SELECT COUNT(*) AS total, "
            "SUM(CASE WHEN TRIM(ff.source_file) <> '' THEN 1 ELSE 0 END) AS with_source_file, "
            "SUM(CASE WHEN ff.page_no IS NOT NULL THEN 1 ELSE 0 END) AS with_page_no, "
            "SUM(CASE WHEN ff.row_label IS NOT NULL AND TRIM(ff.row_label) <> '' THEN 1 ELSE 0 END) AS with_row_label, "
            "SUM(CASE WHEN ff.column_label IS NOT NULL AND TRIM(ff.column_label) <> '' THEN 1 ELSE 0 END) "
            "AS with_column_label, "
            "SUM(CASE WHEN TRIM(ff.source_file) <> '' AND ff.page_no IS NOT NULL "
            "AND ff.row_label IS NOT NULL AND TRIM(ff.row_label) <> '' "
            "AND ff.column_label IS NOT NULL AND TRIM(ff.column_label) <> '' THEN 1 ELSE 0 END) AS fully_traceable "
            "FROM financial_fact ff INNER JOIN financial_source fs ON fs.source_key = ff.source_key "
            "AND fs.stock_code = ff.stock_code AND fs.period = ff.period "
            f"WHERE fs.authority_status = {self.db.parameter_marker}",
            (SourceAuthorityStatus.CURRENT.value,),
        )
        row = rows[0] if rows else {}
        counts = {
            key: int(row.get(key) or 0)
            for key in (
                "total",
                "with_source_file",
                "with_page_no",
                "with_row_label",
                "with_column_label",
                "fully_traceable",
            )
        }
        total = counts["total"]

        def coverage(count: int) -> float:
            return round(count / total, 6) if total else 0.0

        return {
            **counts,
            "source_file_coverage": coverage(counts["with_source_file"]),
            "page_no_coverage": coverage(counts["with_page_no"]),
            "row_label_coverage": coverage(counts["with_row_label"]),
            "column_label_coverage": coverage(counts["with_column_label"]),
            "fully_traceable_coverage": coverage(counts["fully_traceable"]),
        }

    def _validation_coverage(
        self,
        consistency_inputs: dict[str, list[dict[str, Any]]] | None = None,
    ) -> dict[str, dict[str, Any]]:
        balance_fields = {field.field_name for field in self.schema.get("balance_sheet", [])}
        income_fields = {field.field_name for field in self.schema.get("income_sheet", [])}
        cash_fields = {field.field_name for field in self.schema.get("cash_flow_sheet", [])}
        balance_required = {"asset_total_assets", "liability_total_liabilities", "equity_total_equity"}
        yoy_required = {"total_operating_revenue", "operating_revenue_yoy_growth", "report_year"}
        balance_schema_ready = balance_required <= balance_fields
        yoy_schema_ready = yoy_required <= income_fields
        inputs = consistency_inputs if consistency_inputs is not None else self._financial_consistency_inputs()
        balance_eligible = len(inputs["balance_equation"])
        yoy_eligible = len(inputs["revenue_yoy_recalculation"])
        return {
            "balance_equation": {
                "status": "available" if balance_schema_ready and balance_eligible else "unavailable",
                "schema_ready": balance_schema_ready,
                "eligible_records": balance_eligible,
                "reason": (
                    f"{balance_eligible} record(s) have complete VALIDATED inputs"
                    if balance_eligible
                    else (
                        "no records have complete VALIDATED inputs"
                        if balance_schema_ready
                        else f"schema lacks required fields: {', '.join(sorted(balance_required - balance_fields))}"
                    )
                ),
            },
            "revenue_yoy_recalculation": {
                "status": "available" if yoy_schema_ready and yoy_eligible else "unavailable",
                "schema_ready": yoy_schema_ready,
                "eligible_records": yoy_eligible,
                "reason": (
                    f"{yoy_eligible} record(s) have complete VALIDATED current/prior-period inputs"
                    if yoy_eligible
                    else (
                        "no records have complete VALIDATED current/prior-period inputs"
                        if yoy_schema_ready
                        else f"schema lacks required fields: {', '.join(sorted(yoy_required - income_fields))}"
                    )
                ),
            },
            "cash_net_change": {
                "status": "unavailable",
                "schema_ready": {"opening_cash", "closing_cash", "net_cash_change"} <= cash_fields,
                "eligible_records": 0,
                "reason": (
                    "schema lacks opening_cash, closing_cash, and net_cash_change fields"
                    if not {"opening_cash", "closing_cash", "net_cash_change"} <= cash_fields
                    else "validator not implemented"
                ),
            },
        }

    def _financial_consistency_inputs(self) -> dict[str, list[dict[str, Any]]]:
        balance_required = ("asset_total_assets", "liability_total_liabilities", "equity_total_equity")
        income_required = ("total_operating_revenue", "operating_revenue_yoy_growth")
        balance_fields = {field.field_name for field in self.schema.get("balance_sheet", [])}
        income_fields = {field.field_name for field in self.schema.get("income_sheet", [])}

        balance_rows: list[dict[str, Any]] = []
        if set(balance_required) <= balance_fields:
            balance_rows = [
                row
                for row in self._validated_projection_rows("balance_sheet", balance_required)
                if all(isinstance(row.get(field), (int, float)) for field in balance_required)
            ]

        yoy_rows: list[dict[str, Any]] = []
        if {*income_required, "report_year"} <= income_fields:
            income_rows = self._validated_projection_rows("income_sheet", income_required)
            grouped: dict[tuple[str, str], dict[int, dict[str, Any]]] = {}
            for row in income_rows:
                current_value = row.get("total_operating_revenue")
                if not isinstance(current_value, (int, float)):
                    continue
                report_period = str(row.get("report_period") or "")
                report_year = row.get("report_year")
                if not isinstance(report_year, int):
                    continue
                grouped.setdefault((str(row["stock_code"]), report_period[-2:]), {})[report_year] = row
            for by_year in grouped.values():
                for year, current in by_year.items():
                    previous = by_year.get(year - 1)
                    reported = current.get("operating_revenue_yoy_growth")
                    if (
                        previous is None
                        or not isinstance(reported, (int, float))
                        or previous.get("total_operating_revenue") in (None, 0)
                    ):
                        continue
                    yoy_rows.append(
                        {
                            **current,
                            "previous_total_operating_revenue": previous["total_operating_revenue"],
                        }
                    )

        return {
            "balance_equation": balance_rows,
            "revenue_yoy_recalculation": yoy_rows,
        }

    def _validated_projection_rows(self, table: str, metrics: tuple[str, ...]) -> list[dict[str, Any]]:
        if table not in self.schema:
            raise ValueError(f"unknown financial statement table: {table!r}")
        table_fields = {field.field_name for field in self.schema[table]}
        unknown_metrics = sorted(set(metrics) - table_fields)
        if unknown_metrics:
            raise ValueError(f"unknown metrics for {table}: {unknown_metrics}")

        marker = self.db.parameter_marker
        rows = self.db.query(
            f"SELECT projection.stock_code, projection.report_period, projection.report_year, "
            "fact.metric, fact.normalized_value FROM "
            f"{table} AS projection INNER JOIN financial_fact AS fact "
            "ON fact.stock_code = projection.stock_code AND fact.period = projection.report_period "
            "INNER JOIN financial_source AS source ON source.source_key = fact.source_key "
            "AND source.stock_code = fact.stock_code AND source.period = fact.period "
            f"WHERE fact.validation_status = {marker} AND fact.statement_scope = {marker} "
            f"AND fact.period_type = {marker} AND fact.table_name = {marker} "
            f"AND source.authority_status = {marker} "
            f"AND fact.metric IN ({self.db.parameter_markers(len(metrics))})",
            (
                "VALIDATED",
                "consolidated",
                expected_period_type_for_table(table).value,
                table,
                SourceAuthorityStatus.CURRENT.value,
                *metrics,
            ),
        )
        grouped: dict[tuple[str, str], dict[str, Any]] = {}
        for row in rows:
            key = (str(row["stock_code"]), str(row["report_period"]))
            target = grouped.setdefault(
                key,
                {
                    "stock_code": row["stock_code"],
                    "report_period": row["report_period"],
                    "report_year": row["report_year"],
                },
            )
            metric = str(row["metric"])
            if metric in target:
                raise RuntimeError(f"multiple VALIDATED facts conflict for {key[0]}/{key[1]}/{metric}")
            target[metric] = row["normalized_value"]
        return [grouped[key] for key in sorted(grouped)]

    def _financial_consistency_issues(
        self,
        consistency_inputs: dict[str, list[dict[str, Any]]] | None = None,
    ) -> list[dict[str, Any]]:
        issues: list[dict[str, Any]] = []
        inputs = consistency_inputs if consistency_inputs is not None else self._financial_consistency_inputs()
        for row in inputs["balance_equation"]:
            assets = float(row["asset_total_assets"])
            liabilities = float(row["liability_total_liabilities"])
            equity = float(row["equity_total_equity"])
            difference = assets - liabilities - equity
            tolerance = max(abs(assets) * 0.005, 1.0)
            if abs(difference) > tolerance:
                issues.append(
                    {
                        "code": "balance_equation_mismatch",
                        "stock_code": row["stock_code"],
                        "report_period": row["report_period"],
                        "difference": round(difference, 4),
                        "tolerance": round(tolerance, 4),
                    }
                )

        for item in inputs["revenue_yoy_recalculation"]:
            reported = float(item["operating_revenue_yoy_growth"])
            current_value = float(item["total_operating_revenue"])
            previous_value = float(item["previous_total_operating_revenue"])
            recomputed = (current_value - previous_value) / previous_value * 100
            if abs(reported - recomputed) > 0.1:
                issues.append(
                    {
                        "code": "yoy_recalculation_mismatch",
                        "stock_code": item["stock_code"],
                        "report_period": item["report_period"],
                        "reported": round(reported, 4),
                        "recomputed": round(recomputed, 4),
                    }
                )
        return issues

    def _ensure_database_ready(self) -> None:
        ready = True
        table_counts: dict[str, int] = {}
        for table in self.schema:
            rows = self.db.query(f"SELECT COUNT(*) AS cnt FROM {table}")
            count = int(rows[0]["cnt"]) if rows else 0
            table_counts[table] = count
            if count <= 0:
                ready = False
        if not ready:
            raise RuntimeError(f"database is not ready for task mode: {table_counts}")
        trusted = self.db.query(
            "SELECT COUNT(*) AS cnt FROM financial_fact ff INNER JOIN financial_source fs "
            "ON fs.source_key = ff.source_key "
            "AND fs.stock_code = ff.stock_code AND fs.period = ff.period "
            f"WHERE ff.validation_status = {self.db.parameter_marker} "
            f"AND fs.authority_status = {self.db.parameter_marker}",
            ("VALIDATED", SourceAuthorityStatus.CURRENT.value),
        )
        if not trusted or int(trusted[0]["cnt"]) <= 0:
            raise RuntimeError(
                "database has no VALIDATED financial facts; complete human review before question answering"
            )
        consistency_issues = self._financial_consistency_issues()
        if consistency_issues:
            raise RuntimeError(
                f"database has {len(consistency_issues)} financial consistency issue(s); resolve them before answering"
            )

    def _required_tables_by_period(self) -> dict[tuple[str, str], set[str]]:
        required: dict[tuple[str, str], set[str]] = {}
        state_path = self.paths.output_dir / "ingestion_state.json"
        for meta in self._load_ingestion_state(state_path).values():
            if meta.get("status") != "parsed":
                continue
            stock_code = str(meta.get("stock_code", ""))
            report_period = str(meta.get("report_period", ""))
            if not stock_code or not report_period:
                continue
            key = (stock_code, report_period)
            available = meta.get("available_tables")
            if available is None:
                required.setdefault(key, set(self.schema.keys()))
                continue
            required.setdefault(key, set()).update(
                str(table_name) for table_name in available if str(table_name) in self.schema
            )
        return required

    def _has_complete_period_record(
        self,
        stock_code: str,
        report_period: str,
        required_tables: list[str] | None = None,
    ) -> bool:
        tables = (
            list(self.schema.keys())
            if required_tables is None
            else [table_name for table_name in required_tables if table_name in self.schema]
        )
        if not tables:
            return True
        for table in tables:
            rows = self.db.query(
                f"SELECT COUNT(*) AS cnt FROM {table} "
                f"WHERE stock_code = {self.db.parameter_marker} AND report_period = {self.db.parameter_marker}",
                (stock_code, report_period),
            )
            count = int(rows[0]["cnt"]) if rows else 0
            if count <= 0:
                return False
        facts = self.db.query(
            "SELECT COUNT(*) AS cnt FROM financial_fact ff INNER JOIN financial_source fs "
            "ON fs.source_key = ff.source_key "
            "AND fs.stock_code = ff.stock_code AND fs.period = ff.period "
            f"WHERE ff.stock_code = {self.db.parameter_marker} AND ff.period = {self.db.parameter_marker} "
            f"AND ff.extractor_version = {self.db.parameter_marker} "
            f"AND fs.authority_status = {self.db.parameter_marker}",
            (stock_code, report_period, self.extractor_version, SourceAuthorityStatus.CURRENT.value),
        )
        return bool(facts and int(facts[0]["cnt"]) > 0)

    def _delete_period_row(self, table: str, stock_code: str, report_period: str) -> None:
        quoted_table = f"`{table}`" if self.db.backend == "mysql" else table
        sql = (
            f"DELETE FROM {quoted_table} WHERE stock_code={self.db.parameter_marker} "
            f"AND report_period={self.db.parameter_marker}"
        )
        self.db.execute(sql, (stock_code, report_period))

    def _append_ingestion_log(self, item: dict[str, Any]) -> None:
        logs = self.run_log.setdefault("ingestion", [])
        status = str(item.get("status") or "unknown")
        retention = "sampled" if status in SAMPLED_INGESTION_STATUSES else "full"
        stats = self.run_log.setdefault("ingestion_event_stats", {}).setdefault(
            status,
            {"retention": retention, "total": 0, "sampled": 0, "dropped": 0},
        )
        stats["total"] += 1
        should_retain = retention == "full" or stats["sampled"] < self.app_config.ingestion_log_limit
        if should_retain:
            logs.append(item)
            stats["sampled"] += 1
        else:
            stats["dropped"] += 1
            self.run_log["ingestion_overflow"] = int(self.run_log.get("ingestion_overflow", 0)) + 1

    @staticmethod
    def _load_ingestion_state(path: Path) -> dict[str, dict[str, Any]]:
        store = IngestionStateStore(path)
        if not path.exists():
            return {}
        try:
            raw = store.load()
        except InvalidIngestionStateError as exc:
            if str(exc).startswith("Ingestion state must be a JSON object:"):
                raise ValueError(f"Invalid ingestion state object: {path}") from exc
            raise
        if raw.get("version") != INGESTION_STATE_VERSION:
            return {}
        files = raw.get("files")
        if not isinstance(files, dict):
            raise ValueError(f"Invalid ingestion state files mapping: {path}")
        return files

    @staticmethod
    def _save_ingestion_state(path: Path, state: dict[str, dict[str, Any]]) -> None:
        payload = {"version": INGESTION_STATE_VERSION, "files": state}
        IngestionStateStore(path).save(payload)

    @staticmethod
    def _matches_ingestion_signature(cached: dict[str, Any], current: dict[str, Any]) -> bool:
        return all(cached.get(field) == current[field] for field in INGESTION_SIGNATURE_FIELDS)

    def _file_signature(self, path: Path) -> dict[str, int | str]:
        resolved = path.resolve(strict=True)
        cached = self._file_signature_cache.get(resolved)
        if cached is not None:
            return dict(cached)
        st = resolved.stat()
        signature: dict[str, int | str] = {
            "size": int(st.st_size),
            "mtime_ns": int(st.st_mtime_ns),
            "content_sha256": sha256_file(resolved),
        }
        if self._ingestion_context_hashes is None:
            raise RuntimeError("ingestion context hashes are not initialized")
        schema_sha256, company_master_sha256 = self._ingestion_context_hashes
        signature["context_fingerprint"] = ingestion_context_fingerprint(
            dataset_sha256=str(signature["content_sha256"]),
            schema_sha256=schema_sha256,
            company_master_sha256=company_master_sha256,
        )
        self._file_signature_cache[resolved] = signature
        return dict(signature)

    # ---------- Task 2 ----------
    def answer_task2(self) -> Path:
        df = pd.read_excel(self.paths.task2_xlsx)
        rows_out: list[dict[str, Any]] = []
        for _, row in df.iterrows():
            qid = str(row["编号"]).strip()
            turns = json.loads(str(row["问题"]).strip())
            session = SessionState()
            answers: list[dict[str, Any]] = []
            sql_history: list[str] = []
            chart_paths: list[dict[str, str]] = []
            turn_logs: list[dict[str, Any]] = []

            for idx, turn in enumerate(turns, start=1):
                question = str(turn["Q"])
                analysis = self.task2_analyzer.analyze_turn(question, session)
                session = analysis.get("session_state", session)
                if analysis.get("need_clarify"):
                    content = analysis.get("clarify_question") or "请补充关键信息。"
                    answers.append({"Q": question, "A": {"content": content}})
                    turn_logs.append({"Q": question, "intent": analysis, "clarify": True})
                    continue

                query_spec = dict(analysis.get("query_spec") or {})
                compiled, result_rows, query_audit = self._execute_planned_query(
                    {"intent": analysis.get("intent"), "query_spec": query_spec}
                )
                calculation = apply_post_compute(query_spec, result_rows)
                result_rows = list(calculation.rows)
                query_audit["calculation"] = calculation.metadata
                sql_history.append(compiled.sql)
                images = self._render_chart_for_query_spec(qid, idx, query_spec, result_rows)
                chart_paths.extend(images)
                image_paths = [item["path"] for item in images]
                content = self._compose_task2_content(question, query_spec, result_rows, images)
                fact_sources = query_audit["fact_sources"]
                answers.append(
                    {
                        "Q": question,
                        "A": {
                            "content": content,
                            "sources": fact_sources,
                            **({"image": image_paths} if image_paths else {}),
                        },
                    }
                )
                turn_logs.append(
                    {
                        "Q": question,
                        "intent": analysis.get("intent"),
                        "query_spec": query_spec,
                        "sql": compiled.sql,
                        "params": list(compiled.params),
                        "query_audit": query_audit,
                        "rows": len(result_rows),
                        "chart_path": image_paths,
                    }
                )

            rows_out.append(
                {
                    "编号": qid,
                    "问题": json.dumps(turns, ensure_ascii=False),
                    "SQL 查询语句": "\n".join(sql_history) if sql_history else "无",
                    "图形格式": self._graph_type_from_images(chart_paths),
                    "回答": json.dumps(answers, ensure_ascii=False),
                }
            )
            self.run_log["task2"].append(
                {
                    "id": qid,
                    "intent": [log.get("intent") for log in turn_logs],
                    "query_spec": [log.get("query_spec") for log in turn_logs if log.get("query_spec")],
                    "sql": [log.get("sql") for log in turn_logs if log.get("sql")],
                    "params": [log.get("params") for log in turn_logs if "params" in log],
                    "query_audit": [log.get("query_audit") for log in turn_logs if log.get("query_audit")],
                    "rows": [log.get("rows") for log in turn_logs if "rows" in log],
                    "chart_path": [log.get("chart_path") for log in turn_logs if log.get("chart_path")],
                    "turn_logs": turn_logs,
                }
            )

        result_path = self.paths.base_dir / "result_2.xlsx"
        pd.DataFrame(rows_out).to_excel(result_path, index=False)
        return result_path

    def _compose_task2_content(
        self,
        question: str,
        query_spec: dict[str, Any],
        rows: list[dict[str, Any]],
        images: list[dict[str, str]],
    ) -> str:
        if self.llm_client.enabled:
            prompt = {
                "question": question,
                "query_spec": query_spec,
                "rows": rows[:20],
                "images": images,
            }
            ans = self.llm_client.complete_json(
                '你是财报分析助手。返回JSON: {"content":"..."}，content必须严格基于查询结果，不得编造。',
                json.dumps(prompt, ensure_ascii=False),
            )
            if not ans.get("content"):
                raise RuntimeError("Invalid LLM answer response: missing content")
            return str(ans["content"])

        if not rows:
            return "未查询到符合条件的数据。"

        analysis_type = str(query_spec.get("analysis_type", "single_metric"))
        metric_key = str(query_spec.get("metric", "total_profit"))
        metric_spec = TASK2_METRICS.get(metric_key)
        metric_col = metric_spec.field if metric_spec else metric_key
        metric_label = METRIC_LABELS.get(metric_col, metric_spec.label if metric_spec else metric_key)
        unit = METRIC_UNITS.get(metric_col, metric_spec.unit if metric_spec else "")

        if analysis_type == "single_metric":
            slots = {
                "stock_abbr": query_spec.get("stock_abbr", ""),
                "report_period": query_spec.get("report_period", ""),
                "metric": metric_col,
            }
            return format_single_metric_answer(slots, rows[0])

        if analysis_type in {"trend", "comparison"}:
            return format_trend_analysis_answer(rows, metric_col=metric_col, metric_label=metric_label)

        if analysis_type == "topn_metric":
            return format_topn_analysis_answer(rows, metric_col=metric_col, metric_label=metric_label)

        if analysis_type == "intersection_topn":
            joined = "；".join(
                f"{item.get('stock_abbr', item.get('stock_code', '未知公司'))}"
                f" {metric_label}={item.get(metric_col)}{unit}"
                for item in rows
            )
            return f"同时满足排名条件的公司有：{joined}" if joined else "未查询到符合条件的数据。"

        parts: list[str] = []
        for item in rows[:10]:
            company = str(item.get("stock_abbr") or item.get("stock_code") or "未知公司")
            values = []
            for field in query_spec.get("select_fields", []):
                if field in {"stock_code", "stock_abbr", "report_period", "report_year"}:
                    continue
                if field not in item:
                    continue
                label = METRIC_LABELS.get(field, field)
                suffix = METRIC_UNITS.get(field, "")
                values.append(f"{label}={item.get(field)}{suffix}")
            parts.append(f"{company}（{'，'.join(values)}）")
        return "；".join(parts) if parts else "未查询到符合条件的数据。"

    # ---------- Task 3 ----------
    def build_knowledge_base(self) -> None:
        metadata = self._load_research_metadata()
        pdf_files = sorted(self.paths.research_dir.rglob("*.pdf"))
        max_docs = self.app_config.kb_max_documents
        max_chunks_per_paper = self.app_config.kb_max_chunks_per_paper
        failed_papers = 0
        for pdf_path in pdf_files:
            title = pdf_path.stem
            meta = metadata.get(self._research_title_key(title), {})
            try:
                pages = self.pdf_page_extractor.extract(pdf_path)
            except Exception as exc:  # pragma: no cover
                failed_papers += 1
                self.run_log.setdefault("kb_errors", []).append({"paper_path": str(pdf_path), "error": str(exc)})
                continue
            if not any(page.text.strip() for page in pages):
                failed_papers += 1
                self.run_log.setdefault("kb_errors", []).append(
                    {"paper_path": str(pdf_path), "error": "no extractable page text"}
                )
                continue
            paper_path = f"./{pdf_path.relative_to(self.paths.base_dir)}"
            company = self._research_metadata_text(meta, RESEARCH_COMPANY_METADATA_KEYS)
            industry = self._research_metadata_text(meta, RESEARCH_INDUSTRY_METADATA_KEYS)
            published_at = self._research_metadata_date(meta, RESEARCH_PUBLICATION_METADATA_KEYS)
            indexed_chunks = 0
            for page in pages:
                if indexed_chunks >= max_chunks_per_paper:
                    break
                chunks = self._chunk_text(page.text)[: max_chunks_per_paper - indexed_chunks]
                for chunk in chunks:
                    self.kb.add_document(
                        title=title,
                        paper_path=paper_path,
                        text=chunk,
                        paper_image=self._extract_paper_image_hint(page.text),
                        company=company,
                        industry=industry,
                        published_at=published_at,
                        page_no=page.page_no,
                    )
                    indexed_chunks += 1
                    if len(self.kb.documents) >= max_docs:
                        self.run_log["kb_summary"] = {
                            "papers_seen": str(pdf_path),
                            "docs_indexed": len(self.kb.documents),
                            "failed_papers": failed_papers,
                            "capped": True,
                        }
                        return
                if len(self.kb.documents) >= max_docs:
                    self.run_log["kb_summary"] = {
                        "papers_seen": str(pdf_path),
                        "docs_indexed": len(self.kb.documents),
                        "failed_papers": failed_papers,
                        "capped": True,
                    }
                    return
        self.run_log["kb_summary"] = {
            "papers_total": len(pdf_files),
            "docs_indexed": len(self.kb.documents),
            "failed_papers": failed_papers,
            "capped": False,
        }

    def answer_task3(self) -> Path:
        df = pd.read_excel(self.paths.task3_xlsx)
        rows_out: list[dict[str, Any]] = []
        for _, row in df.iterrows():
            qid = str(row["编号"]).strip()
            turns = json.loads(str(row["问题"]).strip())
            answers: list[dict[str, Any]] = []
            sql_history: list[str] = []
            log_item = {
                "id": qid,
                "task_graph": [],
                "references": [],
                "sql": [],
                "params": [],
                "query_audit": [],
                "sql_check": [],
                "rows": [],
                "chart_path": [],
                "intent": [],
                "slots": [],
            }
            session = SessionState()

            for idx, turn in enumerate(turns, start=1):
                question = str(turn["Q"])
                task_graph = self.planner.plan_subtasks(question, session.to_dict())
                exec_state: dict[str, Any] = {}
                refs_collected: list[dict[str, Any]] = []
                fact_sources_collected: list[dict[str, Any]] = []
                reason_text = ""
                image_paths: list[str] = []
                for task in task_graph.get("subtasks", []):
                    kind = str(task.get("kind", "sql"))
                    goal = str(task.get("goal", question))
                    if kind == "sql":
                        intent_payload = self.planner.parse_intent(goal, session.to_dict())
                        session = session.update(slots=intent_payload.get("slots", {}))
                        intent_payload["slots"] = dict(session.slots)
                        log_item["intent"].append(intent_payload.get("intent"))
                        log_item["slots"].append(dict(session.slots))
                        compiled, rows, query_audit = self._execute_planned_query(intent_payload)
                        sql_history.append(compiled.sql)
                        log_item["sql_check"].append(True)
                        log_item["rows"].append(len(rows))
                        exec_state[task["id"]] = {
                            "sql": compiled.sql,
                            "params": list(compiled.params),
                            "rows": rows,
                            "sql_check": True,
                            "query_audit": query_audit,
                        }
                        log_item["sql"].append(compiled.sql)
                        log_item["params"].append(list(compiled.params))
                        log_item["query_audit"].append(query_audit)
                        fact_sources_collected.extend(query_audit["fact_sources"])
                        images = self._render_chart_for_intent(qid, idx, intent_payload, rows)
                        image_paths.extend(images)
                        if images:
                            log_item["chart_path"].append(images)
                    elif kind == "retrieval":
                        refs = self._kb_references(goal, top_k=3)
                        exec_state[task["id"]] = {"references": refs}
                        refs_collected.extend(refs)
                    else:  # reason
                        reason_text = self._reason_from_state(question, exec_state, refs_collected)
                        exec_state[task["id"]] = {"reason": reason_text}

                if not reason_text:
                    reason_text = self._reason_from_state(question, exec_state, refs_collected)

                answer_payload: dict[str, Any] = {"content": reason_text}
                if image_paths:
                    answer_payload["image"] = image_paths
                if refs_collected:
                    answer_payload["references"] = refs_collected
                if fact_sources_collected:
                    answer_payload["fact_sources"] = fact_sources_collected
                answers.append({"Q": question, "A": answer_payload})
                log_item["task_graph"].append(task_graph)
                log_item["references"].append(refs_collected)

            rows_out.append(
                {
                    "编号": qid,
                    "问题": json.dumps(turns, ensure_ascii=False),
                    "SQL 查询语法": "\n".join(sql_history) if sql_history else "无",
                    "回答": json.dumps(answers, ensure_ascii=False),
                }
            )
            self.run_log["task3"].append(log_item)

        result_path = self.paths.base_dir / "result_3.xlsx"
        pd.DataFrame(rows_out).to_excel(result_path, index=False)
        return result_path

    def _execute_planned_query(
        self, intent_payload: dict[str, Any]
    ) -> tuple[CompiledQuery, list[dict[str, Any]], dict[str, Any]]:
        try:
            compiled = self.sql_planner.compile(intent_payload, backend=self.db.backend)
        except QueryValidationError as exc:
            self.run_log.setdefault("query_audit", []).append(
                {
                    "timestamp": datetime.now(timezone.utc).isoformat(),
                    "dialect": self.db.backend,
                    "status": "rejected",
                    "stage": "planning",
                    "error": str(exc),
                }
            )
            raise
        try:
            rows = self.safe_query_executor.execute(compiled)
        except Exception:
            audit = self._query_audit_dict(self.safe_query_executor.audit_log[-1])
            audit["stage"] = "execution"
            self.run_log.setdefault("query_audit", []).append(audit)
            raise
        audit = self.safe_query_executor.audit_log[-1]
        audit_payload = self._query_audit_dict(audit)
        audit_payload["stage"] = "execution"
        try:
            audit_payload["fact_sources"] = self._validated_fact_sources(intent_payload, rows)
        except UntrustedFactError as exc:
            audit_payload["trust_status"] = "rejected"
            audit_payload["trust_error"] = str(exc)
            self.run_log.setdefault("query_audit", []).append(audit_payload)
            raise
        audit_payload["trust_status"] = "validated"
        self.run_log.setdefault("query_audit", []).append(audit_payload)
        return compiled, rows, audit_payload

    @staticmethod
    def _query_audit_dict(audit: QueryAuditRecord) -> dict[str, Any]:
        return {
            "query_id": audit.query_id,
            "timestamp": audit.timestamp,
            "dialect": audit.dialect,
            "status": audit.status,
            "parameter_count": audit.parameter_count,
            "row_count": audit.row_count,
            "error": audit.error,
        }

    def _validated_fact_sources(
        self,
        intent_payload: dict[str, Any],
        rows: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        if not rows:
            return []
        query_spec = intent_payload.get("query_spec")
        if isinstance(query_spec, dict):
            metric_keys = [str(item) for item in query_spec.get("metrics", [])]
            if not metric_keys and query_spec.get("metric"):
                metric_keys = [str(query_spec["metric"])]
            default_period = str(query_spec.get("report_period") or "")
            default_code = str(query_spec.get("stock_code") or "")
        else:
            slots = intent_payload.get("slots", {})
            metric_keys = [str(slots.get("metric") or "")]
            default_period = str(slots.get("report_period") or "")
            default_code = str(slots.get("stock_code") or "")

        metric_fields = [self.sql_planner.registry.resolve_metric(key).field for key in metric_keys if key]
        sources: list[dict[str, Any]] = []
        for row in rows:
            stock_code = str(row.get("stock_code") or default_code)
            if not stock_code:
                resolved = self.company_resolver.resolve(str(row.get("stock_abbr") or ""))
                stock_code = str(resolved.get("stock_code") or "") if resolved else ""
            period = str(row.get("report_period") or default_period)
            if not stock_code or not period:
                raise UntrustedFactError("Query result is missing stock_code or report_period provenance keys")
            for metric in metric_fields:
                if metric not in row or row[metric] is None:
                    raise UntrustedFactError(f"Query result has no reviewable value for {stock_code}/{period}/{metric}")
                source = self._validated_fact_source(stock_code, period, metric)
                sources.append(source)
        return sources

    def _validated_fact_source(
        self,
        stock_code: str,
        period: str,
        metric: str,
    ) -> dict[str, Any]:
        marker = self.db.parameter_marker
        target = self.sql_planner.registry.resolve_metric(metric)
        period_type = expected_period_type_for_table(target.table).value
        facts = self.db.query(
            "SELECT ff.fact_key, ff.normalized_value, ff.target_unit, ff.source_file, ff.page_no, ff.table_name, "
            "ff.row_label, ff.column_label, ff.validation_status FROM financial_fact ff "
            "INNER JOIN financial_source fs ON fs.source_key = ff.source_key "
            "AND fs.stock_code = ff.stock_code AND fs.period = ff.period "
            f"WHERE ff.stock_code = {marker} AND ff.period = {marker} AND ff.metric = {marker} "
            f"AND ff.validation_status = {marker} AND ff.statement_scope = {marker} AND ff.period_type = {marker} "
            f"AND fs.authority_status = {marker}",
            (
                stock_code,
                period,
                metric,
                "VALIDATED",
                "consolidated",
                period_type,
                SourceAuthorityStatus.CURRENT.value,
            ),
            use_cache=False,
        )
        if len(facts) != 1:
            raise UntrustedFactError(
                f"Expected exactly one consolidated VALIDATED fact for {stock_code}/{period}/{metric}, got {len(facts)}"
            )
        fact = facts[0]
        return {
            "fact_key": fact["fact_key"],
            "stock_code": stock_code,
            "period": period,
            "metric": metric,
            "normalized_value": fact["normalized_value"],
            "unit": fact["target_unit"],
            "source_file": fact["source_file"],
            "page_no": fact["page_no"],
            "table_name": fact["table_name"],
            "row_label": fact["row_label"],
            "column_label": fact["column_label"],
        }

    # ---------- Helpers ----------
    def _graph_type_from_images(self, images: list[Any]) -> str:
        if not images:
            return "无"
        first = images[0]
        if isinstance(first, dict):
            chart_type = str(first.get("type", ""))
            return {
                "line": "折线图",
                "bar": "柱状图",
                "horizontal_bar": "水平柱状图",
                "pie": "饼图",
                "combo": "组合图",
            }.get(chart_type, "无")
        if len(images) == 1:
            return "折线图"
        return "组合图"

    def _render_chart_for_intent(
        self,
        qid: str,
        seq: int,
        intent_payload: dict[str, Any],
        rows: list[dict[str, Any]],
    ) -> list[str]:
        if not rows:
            return []
        intent = intent_payload.get("intent")
        if intent == "single_metric":
            return []
        if intent in {"trend"}:
            x_vals: list[str] = []
            y_vals: list[float] = []
            metric_col = None
            for row in rows:
                if metric_col is None:
                    metric_col = next(
                        (k for k, v in row.items() if k != "report_period" and isinstance(v, (int, float))), None
                    )
                if not metric_col:
                    continue
                value = row.get(metric_col)
                if value is None:
                    continue
                try:
                    y = float(value)
                except (TypeError, ValueError):
                    continue
                x_vals.append(str(row.get("report_period", "")))
                y_vals.append(y)
            if x_vals and y_vals:
                return [
                    self._plot_line_chart(qid, seq, x_vals, y_vals, title=f"{metric_col}趋势", y_label=str(metric_col))
                ]
            return []
        # topn/comparison -> bar
        metric_col = next(
            (
                k
                for k, v in rows[0].items()
                if k not in {"stock_abbr", "stock_code", "report_period"} and isinstance(v, (int, float))
            ),
            None,
        )
        if not metric_col:
            return []
        names: list[str] = []
        values: list[float] = []
        for r in rows[:20]:
            value = r.get(metric_col)
            if value is None:
                continue
            try:
                fv = float(value)
            except (TypeError, ValueError):
                continue
            names.append(str(r.get("stock_abbr", r.get("stock_code", ""))))
            values.append(fv)
        if not names or not values:
            return []
        return [self._plot_bar_chart(qid, seq, names, values, title=f"{metric_col}对比", y_label=str(metric_col))]

    def _render_chart_for_query_spec(
        self,
        qid: str,
        seq: int,
        query_spec: dict[str, Any],
        rows: list[dict[str, Any]],
    ) -> list[dict[str, str]]:
        if not rows:
            return []
        chart_request = str(query_spec.get("chart_request") or "")
        analysis_type = str(query_spec.get("analysis_type") or "")
        metric_key = str(query_spec.get("metric") or "")
        metric_spec = TASK2_METRICS.get(metric_key)
        metric_col = metric_spec.field if metric_spec else metric_key
        title = f"{METRIC_LABELS.get(metric_col, metric_key)}分析"
        y_label = METRIC_LABELS.get(metric_col, metric_col)

        if chart_request == "line" or analysis_type in {"trend", "comparison"}:
            ordered = sorted(rows, key=lambda item: _period_sort_key(str(item.get("report_period", ""))))
            x_vals = [str(item.get("report_period", "")) for item in ordered]
            y_vals = [float(item.get(metric_col, 0) or 0) for item in ordered if item.get(metric_col) is not None]
            if x_vals and y_vals and len(x_vals) == len(y_vals):
                path = self._plot_line_chart(qid, seq, x_vals, y_vals, title=title, y_label=y_label)
                return [{"path": path, "type": "line"}]
            return []

        if chart_request == "horizontal_bar":
            names = [
                str(item.get("stock_abbr", item.get("stock_code", "")))
                for item in rows[:20]
                if item.get(metric_col) is not None
            ]
            values = [float(item.get(metric_col, 0) or 0) for item in rows[:20] if item.get(metric_col) is not None]
            if names and values:
                path = self._plot_bar_chart(qid, seq, names, values, title=title, y_label=y_label, horizontal=True)
                return [{"path": path, "type": "horizontal_bar"}]
            return []

        if chart_request == "pie":
            names = [
                str(item.get("stock_abbr", item.get("stock_code", "")))
                for item in rows[:8]
                if item.get(metric_col) is not None
            ]
            values = [float(item.get(metric_col, 0) or 0) for item in rows[:8] if item.get(metric_col) is not None]
            if names and values:
                path = self._plot_pie_chart(qid, seq, names, values, title=title)
                return [{"path": path, "type": "pie"}]
            return []

        if chart_request == "bar" or analysis_type in {"topn_metric", "filter", "intersection_topn"}:
            names = [
                str(item.get("stock_abbr", item.get("stock_code", "")))
                for item in rows[:20]
                if item.get(metric_col) is not None
            ]
            values = [float(item.get(metric_col, 0) or 0) for item in rows[:20] if item.get(metric_col) is not None]
            if names and values:
                path = self._plot_bar_chart(qid, seq, names, values, title=title, y_label=y_label)
                return [{"path": path, "type": "bar"}]
        return []

    def _kb_references(self, query: str, top_k: int = 3) -> list[dict[str, Any]]:
        result = self.kb.search_evidence(query, top_k=top_k, filters=self._kb_evidence_filter(query))
        if result.status is EvidenceStatus.INSUFFICIENT_EVIDENCE:
            return []
        return [citation.as_dict() for citation in result.citations]

    def _reason_from_state(self, question: str, exec_state: dict[str, Any], refs: list[dict[str, Any]]) -> str:
        retrieval_requested = any(isinstance(item, dict) and "references" in item for item in exec_state.values())
        if retrieval_requested and not refs:
            return "证据不足：未找到可定位到研报页码的原文，无法给出可靠归因。"
        if self.llm_client.enabled:
            payload = {"question": question, "exec_state": exec_state, "references": refs}
            ans = self.llm_client.complete_json(
                '你是财务分析归因助手。返回JSON: {"content":"..."}，要求包含结论、依据和可靠性说明。',
                json.dumps(payload, ensure_ascii=False),
            )
            if not ans.get("content"):
                raise RuntimeError("Invalid LLM reasoning response: missing content")
            return str(ans["content"])
        sql_rows = 0
        topn_rows: list[dict[str, Any]] = []
        for item in exec_state.values():
            if isinstance(item, dict) and isinstance(item.get("rows"), list):
                sql_rows += len(item["rows"])
                topn_rows.extend(item.get("rows", []))
        if topn_rows:
            metric_col = "total_profit" if any("total_profit" in r for r in topn_rows) else None
            if metric_col:
                return format_topn_analysis_answer(
                    topn_rows, metric_col=metric_col, metric_label=METRIC_LABELS.get(metric_col, metric_col)
                )
            # Trend-like rows in task3 SQL path.
            if all("report_period" in r for r in topn_rows):
                metric_col = next(
                    (
                        k
                        for k in topn_rows[0].keys()
                        if k not in {"stock_abbr", "stock_code", "report_period", "report_year"}
                        and isinstance(topn_rows[0].get(k), (int, float))
                    ),
                    None,
                )
                if metric_col:
                    return format_trend_analysis_answer(
                        topn_rows, metric_col=metric_col, metric_label=METRIC_LABELS.get(metric_col, metric_col)
                    )
        if refs:
            return format_reason_with_references(question, sql_rows_count=sql_rows, references=refs)
        return f"基于结构化数据分析：共检索到 {sql_rows} 条记录。"

    def _kb_evidence_filter(self, query: str) -> EvidenceFilter:
        resolved = self.company_resolver.resolve(query)
        company = str(resolved["stock_abbr"]) if resolved and resolved.get("stock_code") else None
        industries = sorted(
            {document.industry for document in self.kb.documents if document.industry},
            key=len,
            reverse=True,
        )
        industry = next((item for item in industries if item in query), None)
        return EvidenceFilter(company=company, industry=industry)

    def _load_research_metadata(self) -> dict[str, dict[str, Any]]:
        metadata: dict[str, dict[str, Any]] = {}
        conflicting_titles: set[str] = set()
        xlsx_files = [
            self.paths.research_dir / "个股_研报信息.xlsx",
            self.paths.research_dir / "行业_研报信息.xlsx",
        ]
        for xlsx_path in xlsx_files:
            if not xlsx_path.exists():
                continue
            wb = load_workbook(xlsx_path, read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            rows = ws.iter_rows(min_row=1, values_only=True)
            header = [str(item) for item in next(rows)]
            for row in rows:
                values = {header[i]: row[i] for i in range(len(header))}
                title = str(values.get("title") or "").strip()
                if not title:
                    continue
                title_key = self._research_title_key(title)
                if title_key in conflicting_titles:
                    continue
                existing = metadata.get(title_key)
                if existing is not None and existing != values:
                    metadata.pop(title_key)
                    conflicting_titles.add(title_key)
                    continue
                metadata[title_key] = values
        if conflicting_titles:
            self.run_log["kb_metadata_conflicts"] = [
                {
                    "title_key": title_key,
                    "error": "multiple research metadata rows resolve to the same PDF title",
                }
                for title_key in sorted(conflicting_titles)
            ]
        return metadata

    @staticmethod
    def _research_title_key(title: str) -> str:
        return re.sub(r"\s+", " ", title.strip()).replace("/", "_").replace("\\", "_")

    @staticmethod
    def _research_metadata_text(meta: dict[str, Any], keys: tuple[str, ...]) -> str | None:
        for key in keys:
            value = meta.get(key)
            if value is not None and str(value).strip():
                return str(value).strip()
        return None

    @staticmethod
    def _research_metadata_date(meta: dict[str, Any], keys: tuple[str, ...]) -> date | None:
        for key in keys:
            value = meta.get(key)
            if value is None or not str(value).strip():
                continue
            if isinstance(value, datetime):
                return value.date()
            if isinstance(value, date):
                return value
            text = str(value).strip()
            if re.fullmatch(r"\d{4}-\d{2}-\d{2}(?: 00:00:00)?", text):
                return date.fromisoformat(text[:10])
            raise ValueError(f"Invalid research publication date for {key}: {value!r}")
        return None

    @staticmethod
    def _chunk_text(text: str, max_chunk_len: int = 280) -> list[str]:
        cleaned = re.sub(r"\s+", " ", text)
        segments = re.split(r"[。！？；]\s*", cleaned)
        chunks: list[str] = []
        buff = ""
        for segment in segments:
            segment = segment.strip()
            if not segment:
                continue
            if len(buff) + len(segment) + 1 <= max_chunk_len:
                buff = f"{buff}。{segment}" if buff else segment
            else:
                if buff:
                    chunks.append(buff)
                buff = segment
        if buff:
            chunks.append(buff)
        return chunks

    @staticmethod
    def _extract_paper_image_hint(text: str) -> str:
        match = re.search(r"(图表\s*\d+[：:][^\n。]{0,40})", text)
        return match.group(1) if match else ""

    def _plot_line_chart(
        self,
        qid: str,
        seq: int,
        x_values: list[str],
        y_values: list[float],
        *,
        title: str,
        y_label: str,
    ) -> str:
        file_name = f"{qid}_{seq}.jpg"
        out_path = self.paths.result_dir / file_name
        plt.figure(figsize=(8, 4.8))
        plt.plot(x_values, y_values, marker="o")
        plt.title(title)
        plt.xlabel("报告期")
        plt.ylabel(y_label)
        plt.grid(alpha=0.3)
        plt.tight_layout()
        plt.savefig(out_path, dpi=160)
        plt.close()
        return f"./result/{file_name}"

    def _plot_bar_chart(
        self,
        qid: str,
        seq: int,
        x_values: list[str],
        y_values: list[float],
        *,
        title: str,
        y_label: str,
        horizontal: bool = False,
    ) -> str:
        file_name = f"{qid}_{seq}.jpg"
        out_path = self.paths.result_dir / file_name
        plt.figure(figsize=(8, 4.8))
        if horizontal:
            plt.barh(x_values, y_values, color="#2a7fff")
            plt.xlabel(y_label)
        else:
            plt.bar(x_values, y_values, color="#2a7fff")
            plt.ylabel(y_label)
        plt.title(title)
        if not horizontal:
            plt.xticks(rotation=20)
        plt.tight_layout()
        plt.savefig(out_path, dpi=160)
        plt.close()
        return f"./result/{file_name}"

    def _plot_pie_chart(
        self,
        qid: str,
        seq: int,
        labels: list[str],
        values: list[float],
        *,
        title: str,
    ) -> str:
        file_name = f"{qid}_{seq}.jpg"
        out_path = self.paths.result_dir / file_name
        plt.figure(figsize=(7.2, 5.2))
        plt.pie(values, labels=labels, autopct="%1.1f%%")
        plt.title(title)
        plt.tight_layout()
        plt.savefig(out_path, dpi=160)
        plt.close()
        return f"./result/{file_name}"
