from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .core import report_period_order_value
from .safe_query import SessionState


@dataclass(slots=True)
class MetricSpec:
    key: str
    label: str
    table: str
    field: str
    unit: str
    aliases: tuple[str, ...]


TASK2_METRICS: dict[str, MetricSpec] = {
    "total_profit": MetricSpec(
        key="total_profit",
        label="利润总额",
        table="income_sheet",
        field="total_profit",
        unit="万元",
        aliases=("利润总额", "核心利润指标", "亏钱"),
    ),
    "net_profit": MetricSpec(
        key="net_profit",
        label="净利润",
        table="income_sheet",
        field="net_profit",
        unit="万元",
        aliases=("净利润",),
    ),
    "total_operating_revenue": MetricSpec(
        key="total_operating_revenue",
        label="营业总收入",
        table="income_sheet",
        field="total_operating_revenue",
        unit="万元",
        aliases=("营业总收入", "主营业务收入", "收入", "营收", "销售额"),
    ),
    "operating_revenue_yoy_growth": MetricSpec(
        key="operating_revenue_yoy_growth",
        label="营业总收入同比增长率",
        table="income_sheet",
        field="operating_revenue_yoy_growth",
        unit="%",
        aliases=("营业总收入同比增长率", "营业总收入同比", "主营业务收入同比增长率", "销售额年同比"),
    ),
    "operating_revenue_qoq_growth": MetricSpec(
        key="operating_revenue_qoq_growth",
        label="营业总收入环比增长率",
        table="core_performance_indicators_sheet",
        field="operating_revenue_qoq_growth",
        unit="%",
        aliases=("营业总收入环比增长率", "主营业务收入环比增长率", "环比增长率"),
    ),
    "asset_liability_ratio": MetricSpec(
        key="asset_liability_ratio",
        label="资产负债率",
        table="balance_sheet",
        field="asset_liability_ratio",
        unit="%",
        aliases=("资产负债率", "负债总额/资产总额"),
    ),
    "operating_cf_net_amount": MetricSpec(
        key="operating_cf_net_amount",
        label="经营性现金流量净额",
        table="cash_flow_sheet",
        field="operating_cf_net_amount",
        unit="万元",
        aliases=("经营性现金流量净额", "经营性现金流净额", "经营活动产生的现金流量净额"),
    ),
    "operating_expense_rnd_expenses": MetricSpec(
        key="operating_expense_rnd_expenses",
        label="研发费用",
        table="income_sheet",
        field="operating_expense_rnd_expenses",
        unit="万元",
        aliases=("研发费用", "营业总支出-研发费用"),
    ),
    "operating_expense_selling_expenses": MetricSpec(
        key="operating_expense_selling_expenses",
        label="销售费用",
        table="income_sheet",
        field="operating_expense_selling_expenses",
        unit="万元",
        aliases=("销售费用", "营业总支出-销售费用"),
    ),
    "gross_profit_margin": MetricSpec(
        key="gross_profit_margin",
        label="销售毛利率",
        table="core_performance_indicators_sheet",
        field="gross_profit_margin",
        unit="%",
        aliases=("销售毛利率",),
    ),
    "net_profit_margin": MetricSpec(
        key="net_profit_margin",
        label="销售净利率",
        table="core_performance_indicators_sheet",
        field="net_profit_margin",
        unit="%",
        aliases=("销售净利率",),
    ),
    "asset_inventory": MetricSpec(
        key="asset_inventory",
        label="存货",
        table="balance_sheet",
        field="asset_inventory",
        unit="万元",
        aliases=("存货", "资产-存货"),
    ),
    "asset_accounts_receivable": MetricSpec(
        key="asset_accounts_receivable",
        label="应收账款",
        table="balance_sheet",
        field="asset_accounts_receivable",
        unit="万元",
        aliases=("应收账款", "资产-应收账款"),
    ),
    "equity_unappropriated_profit": MetricSpec(
        key="equity_unappropriated_profit",
        label="未分配利润",
        table="balance_sheet",
        field="equity_unappropriated_profit",
        unit="万元",
        aliases=("未分配利润", "股东权益-未分配利润"),
    ),
    "roe_weighted_excl_non_recurring": MetricSpec(
        key="roe_weighted_excl_non_recurring",
        label="加权平均净资产收益率（扣非）",
        table="core_performance_indicators_sheet",
        field="roe_weighted_excl_non_recurring",
        unit="%",
        aliases=("加权平均净资产收益率（扣非）",),
    ),
    "roe": MetricSpec(
        key="roe",
        label="净资产收益率",
        table="core_performance_indicators_sheet",
        field="roe",
        unit="%",
        aliases=("收益率", "净资产收益率"),
    ),
}


_ALIAS_TO_METRIC: list[tuple[str, MetricSpec]] = []
for _spec in TASK2_METRICS.values():
    for _alias in _spec.aliases:
        _ALIAS_TO_METRIC.append((_alias, _spec))
_ALIAS_TO_METRIC.sort(key=lambda item: len(item[0]), reverse=True)


CHINESE_TOP_N = {
    "前一": 1,
    "前二": 2,
    "前三": 3,
    "前四": 4,
    "前五": 5,
    "前六": 6,
    "前七": 7,
    "前八": 8,
    "前九": 9,
    "前十": 10,
}


def _normalize_number(raw: str) -> float:
    return float(raw.replace(",", "").strip())


def _parse_threshold_value(question: str) -> float | None:
    match = re.search(r"(超过|高于|大于|低于|小于)\s*([0-9]+(?:\.[0-9]+)?)\s*(亿元|万元|%|元)?", question)
    if not match:
        return None
    value = _normalize_number(match.group(2))
    unit = match.group(3) or ""
    if unit == "亿元":
        return value * 10000.0
    if unit == "元":
        return value / 10000.0
    return value


def _parse_top_n(question: str, default: int = 10) -> int:
    match = re.search(r"top\s*(\d+)", question, flags=re.IGNORECASE)
    if match:
        return max(1, min(int(match.group(1)), 100))
    for token, value in CHINESE_TOP_N.items():
        if token in question:
            return value
    return default


def _specific_period_to_code(year: int, token: str) -> str:
    token = token.strip()
    if token in {"第一季度", "一季度"}:
        return f"{year}Q1"
    if token in {"上半年", "半年度", "半年度报告"}:
        return f"{year}Q2"
    if token in {"前三季度", "第三季度", "三季度"}:
        return f"{year}Q3"
    return f"{year}FY"


def resolve_periods(question: str, context: dict[str, Any] | None = None) -> dict[str, Any]:
    context = context or {}
    result: dict[str, Any] = {}
    range_match = re.search(
        r"(\d{4})\s*年?\s*(?:至|到|-)\s*(\d{4})\s*年?\s*"
        r"(第一季度|一季度|上半年|半年度|前三季度|第三季度|三季度|年度|年报)?",
        question,
    )
    if range_match:
        start_year = int(range_match.group(1))
        end_year = int(range_match.group(2))
        end_token = range_match.group(3) or "年度"
        result["start_period"] = f"{start_year}Q1"
        result["end_period"] = _specific_period_to_code(end_year, end_token)
        return result

    explicit_periods: list[str] = []
    for year, token in re.findall(
        r"(\d{4})\s*年\s*(第一季度|一季度|上半年|半年度|前三季度|第三季度|三季度|年度|年报)",
        question,
    ):
        explicit_periods.append(_specific_period_to_code(int(year), token))
    if len(explicit_periods) >= 2:
        explicit_periods = sorted(set(explicit_periods), key=report_period_order_value)
        result["periods"] = explicit_periods
        result["start_period"] = explicit_periods[0]
        result["end_period"] = explicit_periods[-1]
        return result
    if explicit_periods:
        result["report_period"] = explicit_periods[-1]
        return result

    year_only = re.search(r"(\d{4})\s*年", question)
    if year_only:
        result["report_period"] = f"{int(year_only.group(1))}FY"
        return result

    if "去年" in question:
        anchor = str(context.get("report_period") or context.get("end_period") or "")
        if re.fullmatch(r"\d{4}(?:Q[1-3]|FY)", anchor):
            result["report_period"] = f"{int(anchor[:4]) - 1}FY"
        else:
            result["needs_period_anchor"] = True
        return result
    if "近三年" in question or "近几年" in question:
        end_period = str(context.get("report_period") or context.get("end_period") or "")
        if re.fullmatch(r"\d{4}(?:Q[1-3]|FY)", end_period):
            result["start_period"] = f"{int(end_period[:4]) - 2}Q1"
            result["end_period"] = end_period
        else:
            result["needs_period_anchor"] = True
        return result
    if "前三季度" in question:
        anchor = str(context.get("report_period") or context.get("end_period") or "")
        if re.fullmatch(r"\d{4}(?:Q[1-3]|FY)", anchor):
            result["report_period"] = f"{anchor[:4]}Q3"
        else:
            result["needs_period_anchor"] = True
    elif "上半年" in question:
        anchor = str(context.get("report_period") or context.get("end_period") or "")
        if re.fullmatch(r"\d{4}(?:Q[1-3]|FY)", anchor):
            result["report_period"] = f"{anchor[:4]}Q2"
        else:
            result["needs_period_anchor"] = True
    return result


@dataclass(slots=True)
class CompanyRecord:
    stock_code: str
    stock_abbr: str
    company_name: str


class CompanyResolver:
    def __init__(self, records: list[CompanyRecord]) -> None:
        self.records = records
        self.by_code = {item.stock_code: item for item in records}

    @classmethod
    def from_xlsx(cls, path: Path) -> "CompanyResolver":
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        records: list[CompanyRecord] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[1] is None or not row[2]:
                continue
            code = str(int(row[1])).zfill(6)
            abbr = str(row[2]).strip()
            company_name = str(row[3] or "").strip()
            records.append(CompanyRecord(stock_code=code, stock_abbr=abbr, company_name=company_name))
        return cls(records)

    def resolve(self, text: str) -> dict[str, str] | None:
        if not text:
            return None

        for match in re.findall(r"\b(\d{6})\b", text):
            if match in self.by_code:
                return self._to_dict(self.by_code[match])

        for match in re.findall(r"\b(\d{3})\b", text):
            code = match.zfill(6)
            if code in self.by_code:
                return self._to_dict(self.by_code[code])

        for item in self.records:
            if item.stock_abbr and item.stock_abbr in text:
                return self._to_dict(item)
        for item in self.records:
            if item.company_name and item.company_name in text:
                return self._to_dict(item)

        if "企业名称：" in text or "企业名称:" in text:
            candidate = re.split(r"企业名称[:：]", text, maxsplit=1)[-1].split("，", 1)[0].split(",", 1)[0].strip()
            resolved = self._resolve_partial(candidate)
            if resolved:
                return resolved

        chinese_tokens = re.findall(r"[\u4e00-\u9fff]{2,}", text)
        for token in sorted(chinese_tokens, key=len, reverse=True):
            resolved = self._resolve_partial(token)
            if resolved:
                return resolved
        return None

    def _resolve_partial(self, candidate: str) -> dict[str, str] | None:
        if not candidate:
            return None
        hits = [item for item in self.records if candidate in item.stock_abbr or candidate in item.company_name]
        if len(hits) == 1:
            return self._to_dict(hits[0])
        return None

    @staticmethod
    def _to_dict(item: CompanyRecord) -> dict[str, str]:
        return {
            "stock_code": item.stock_code,
            "stock_abbr": item.stock_abbr,
            "company_name": item.company_name,
        }


class QuestionAnalyzer:
    def __init__(self, company_resolver: CompanyResolver) -> None:
        self.company_resolver = company_resolver

    def analyze_turn(
        self,
        question: str,
        context: dict[str, Any] | SessionState | None = None,
    ) -> dict[str, Any]:
        session_state = context if isinstance(context, SessionState) else SessionState(slots=context or {})
        context = dict(session_state.slots)
        question = str(question).strip()

        def with_session_state(result: dict[str, Any]) -> dict[str, Any]:
            result["session_state"] = SessionState(
                slots=result.get("context", {}),
                context={
                    "last_intent": str(result.get("intent", "")),
                    "needs_clarification": bool(result.get("need_clarify")),
                },
            )
            return result

        if self._is_chart_followup(question) and context.get("last_query_spec"):
            query_spec = dict(context["last_query_spec"])
            query_spec["chart_request"] = self._detect_chart_request(question) or query_spec.get("chart_request")
            context["analysis_type"] = query_spec.get("analysis_type")
            context["chart_request"] = query_spec.get("chart_request")
            context["last_query_spec"] = query_spec
            return with_session_state(
                {
                    "intent": query_spec.get("analysis_type", "single_metric"),
                    "need_clarify": False,
                    "clarify_question": "",
                    "context": context,
                    "query_spec": query_spec,
                }
            )

        company = self.company_resolver.resolve(question)
        if company:
            previous_code = context.get("stock_code")
            if previous_code and company.get("stock_code") and company["stock_code"] != previous_code:
                for key in (
                    "metric",
                    "metrics",
                    "report_period",
                    "start_period",
                    "end_period",
                    "periods",
                    "chart_request",
                    "analysis_type",
                    "last_query_spec",
                    "needs_period_anchor",
                ):
                    context.pop(key, None)
            context.update(company)

        metrics = self._detect_metrics(question)
        if metrics:
            context["metric"] = metrics[0].key
            context["metrics"] = [item.key for item in metrics]
        else:
            metrics = [TASK2_METRICS[key] for key in context.get("metrics", []) if key in TASK2_METRICS]
            if not metrics and context.get("metric") in TASK2_METRICS:
                metrics = [TASK2_METRICS[context["metric"]]]

        period_info = resolve_periods(question, context)
        if "report_period" in period_info:
            for key in ("start_period", "end_period", "periods", "needs_period_anchor"):
                context.pop(key, None)
        elif "periods" in period_info:
            for key in ("report_period", "needs_period_anchor"):
                context.pop(key, None)
        elif "start_period" in period_info or "end_period" in period_info:
            for key in ("report_period", "periods", "needs_period_anchor"):
                context.pop(key, None)
        context.update(period_info)

        analysis_type = self._detect_analysis_type(question, metrics, context)
        chart_request = self._detect_chart_request(question) or context.get("chart_request")
        if chart_request:
            context["chart_request"] = chart_request
        context["analysis_type"] = analysis_type

        clarify_question = self._clarify_question(context)
        if clarify_question:
            return with_session_state(
                {
                    "intent": analysis_type,
                    "need_clarify": True,
                    "clarify_question": clarify_question,
                    "context": context,
                    "query_spec": None,
                }
            )

        query_spec = self._build_query_spec(question, metrics, context)
        context["last_query_spec"] = query_spec
        return with_session_state(
            {
                "intent": analysis_type,
                "need_clarify": False,
                "clarify_question": "",
                "context": context,
                "query_spec": query_spec,
            }
        )

    def _detect_metrics(self, question: str) -> list[MetricSpec]:
        hits: dict[str, tuple[int, int, MetricSpec]] = {}
        for alias, spec in _ALIAS_TO_METRIC:
            pos = question.find(alias)
            if pos < 0:
                continue
            current = hits.get(spec.key)
            candidate = (pos, -len(alias), spec)
            if current is None or candidate < current:
                hits[spec.key] = candidate
        ordered = sorted(hits.values(), key=lambda item: (item[0], item[1]))
        return [item[2] for item in ordered]

    @staticmethod
    def _is_chart_followup(question: str) -> bool:
        return any(token in question for token in ("绘制", "柱状图", "折线图", "饼图")) and not any(
            token in question for token in ("多少", "哪些", "是多少", "同比", "环比")
        )

    @staticmethod
    def _detect_chart_request(question: str) -> str | None:
        if "水平柱状图" in question:
            return "horizontal_bar"
        if "柱状图" in question:
            return "bar"
        if "饼图" in question:
            return "pie"
        if "折线图" in question or "趋势图" in question:
            return "line"
        return None

    @staticmethod
    def _detect_analysis_type(question: str, metrics: list[MetricSpec], context: dict[str, Any]) -> str:
        if "均排名前" in question:
            return "intersection_topn"
        if any(token in question for token in ("趋势", "变化", "近三年", "近几年", "绘图", "折线图", "可视化")):
            return "trend"
        if any(token in question for token in ("相比", "对比", "比较")) or len(context.get("periods", [])) >= 2:
            return "comparison"
        if any(token in question.lower() for token in ("top",)) or any(
            token in question for token in ("排名", "最高", "前五", "前三", "前十")
        ):
            return "topn_metric"
        if any(token in question for token in ("超过", "高于", "低于", "为负", "为正", "有哪些", "哪些")) and metrics:
            return "filter"
        return "single_metric"

    @staticmethod
    def _clarify_question(context: dict[str, Any]) -> str:
        analysis_type = context.get("analysis_type", "single_metric")
        if context.get("needs_period_anchor"):
            return "请补充用于相对时间计算的年份或报告期。"
        if not context.get("stock_abbr") and analysis_type in {"single_metric", "trend", "comparison"}:
            return "请补充公司名称或股票代码。"
        if not context.get("metric"):
            return "请补充你要查询的财务指标。"
        if analysis_type in {"single_metric", "filter", "topn_metric", "comparison"} and not (
            context.get("report_period") or context.get("start_period")
        ):
            return "请补充明确的年份和报告期。"
        return ""

    def _build_query_spec(
        self,
        question: str,
        metrics: list[MetricSpec],
        context: dict[str, Any],
    ) -> dict[str, Any]:
        analysis_type = str(context.get("analysis_type", "single_metric"))
        if not metrics and context.get("metric") in TASK2_METRICS:
            metrics = [TASK2_METRICS[context["metric"]]]
        metric = metrics[0]
        query_spec: dict[str, Any] = {
            "analysis_type": analysis_type,
            "metric": metric.key,
            "metrics": [item.key for item in metrics],
            "table": metric.table,
            "stock_code": context.get("stock_code", ""),
            "stock_abbr": context.get("stock_abbr", ""),
            "report_period": context.get("report_period", ""),
            "start_period": context.get("start_period", ""),
            "end_period": context.get("end_period", context.get("report_period", "")),
            "chart_request": context.get("chart_request"),
            "filters": [],
            "select_fields": [],
            "top_n": _parse_top_n(question),
        }
        ratio_metrics = self._detect_ratio_operands(question, metrics)
        if ratio_metrics is not None:
            numerator, denominator = ratio_metrics
            query_spec["post_compute"] = "ratio"
            query_spec["calculation"] = {
                "numerator": numerator.field,
                "denominator": denominator.field,
                "output_field": f"{numerator.field}_to_{denominator.field}_ratio",
            }

        if analysis_type == "single_metric":
            query_spec["select_fields"] = ["stock_abbr", "report_period", metric.field]
            return query_spec

        if analysis_type == "trend":
            query_spec["select_fields"] = ["report_period", "report_year", "stock_abbr", metric.field]
            query_spec["chart_request"] = query_spec.get("chart_request") or "line"
            return query_spec

        if analysis_type == "topn_metric":
            query_spec["select_fields"] = ["stock_code", "stock_abbr"] + [
                TASK2_METRICS[item.key].field for item in metrics
            ]
            return query_spec

        if analysis_type == "intersection_topn":
            query_spec["select_fields"] = ["stock_code", "stock_abbr"] + [
                TASK2_METRICS[item.key].field for item in metrics
            ]
            query_spec["post_compute"] = "intersection_topn"
            query_spec["chart_request"] = None
            return query_spec

        if analysis_type == "comparison":
            periods = list(context.get("periods", []))
            query_spec["periods"] = periods
            query_spec["select_fields"] = ["report_period", "report_year", "stock_abbr", metric.field]
            if query_spec.get("chart_request") is None and any(token in question for token in ("趋势", "折线图")):
                query_spec["chart_request"] = "line"
            query_spec["post_compute"] = "period_comparison" if len(periods) >= 2 else None
            return query_spec

        # filter
        query_spec["select_fields"] = ["stock_code", "stock_abbr"] + [TASK2_METRICS[item.key].field for item in metrics]
        threshold = _parse_threshold_value(question)
        if threshold is not None:
            op = ">"
            if "低于" in question or "小于" in question:
                op = "<"
            query_spec["filters"].append({"field": metric.field, "op": op, "value": threshold})
        if "为负" in question:
            query_spec["filters"].append({"field": metric.field, "op": "<", "value": 0})
        if "为正" in question:
            query_spec["filters"].append({"field": metric.field, "op": ">", "value": 0})
        if len(metrics) > 1 and "均高于行业均值" in question:
            query_spec["post_compute"] = "above_industry_mean"
        if len(metrics) > 1 and "且" in question and not query_spec["filters"]:
            for item in metrics:
                if item.field == "operating_cf_net_amount":
                    query_spec["filters"].append({"field": item.field, "op": "<", "value": 0})
                elif item.field == "net_profit":
                    query_spec["filters"].append({"field": item.field, "op": ">", "value": 0})
        return query_spec

    @staticmethod
    def _detect_ratio_operands(question: str, metrics: list[MetricSpec]) -> tuple[MetricSpec, MetricSpec] | None:
        ratio_position = question.rfind("比值")
        if ratio_position < 0 or len(metrics) < 2:
            return None
        prefix = question[:ratio_position]
        positioned: list[tuple[int, MetricSpec]] = []
        for metric in metrics:
            last_position = max((prefix.rfind(alias) for alias in metric.aliases), default=-1)
            if last_position >= 0:
                positioned.append((last_position, metric))
        if len(positioned) < 2:
            return None
        positioned.sort(key=lambda item: item[0])
        return positioned[-2][1], positioned[-1][1]
