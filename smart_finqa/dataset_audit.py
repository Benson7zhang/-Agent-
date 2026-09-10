"""Read-only inventory and coverage audit for a formal Smart FinQA dataset."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from collections.abc import Callable, Sequence
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .ingestion import CompanyIndex
from .schema import SHEET_TO_TABLE, load_schema_from_xlsx

AUDIT_SCHEMA_VERSION = "1.0"
MINIMUM_HEAD_TEXT_CHARACTERS = 80
TEMPORARY_SUFFIXES = {".part", ".temp", ".tmp"}
STOCK_CODE_PATTERN = re.compile(r"(?<!\d)(\d{6})(?!\d)")

ReportInspector = Callable[[Path, CompanyIndex], Any]
ReportSelector = Callable[[Sequence[Any]], Sequence[Any]]


def audit_dataset(
    dataset_dir: str | Path,
    *,
    include_sha256: bool = False,
    report_inspector: ReportInspector | None = None,
    report_selector: ReportSelector | None = None,
) -> dict[str, Any]:
    """Audit source files without writing to or mutating the dataset directory."""
    root = Path(dataset_dir).expanduser().resolve()
    if not root.is_dir():
        raise ValueError(f"dataset_dir is not a directory: {root}")

    component_errors: list[dict[str, str]] = []
    financial_dir = _find_attachment_directory(root, "附件2", "财务报告", component_errors)
    research_dir = _find_attachment_directory(root, "附件5", "研报", component_errors)
    files = sorted((path for path in root.rglob("*") if path.is_file()), key=lambda item: item.as_posix())
    inventory = _build_inventory(root, files, financial_dir, research_dir, include_sha256)

    company_path = _find_attachment_workbook(root, "附件1", component_errors)
    company_master, company_index = _audit_company_master(root, company_path)
    schema_path = _find_attachment_workbook(root, "附件3", component_errors)
    database_schema = _audit_database_schema(root, schema_path)
    question_sets = {
        "attachment_4": _audit_question_set(root, _find_attachment_workbook(root, "附件4", component_errors)),
        "attachment_6": _audit_question_set(root, _find_attachment_workbook(root, "附件6", component_errors)),
    }

    financial_files = _pdf_files(financial_dir)
    financial_reports = _audit_financial_reports(
        root,
        financial_files,
        company_index,
        report_inspector=report_inspector,
        report_selector=report_selector,
    )
    research_reports = _audit_research_reports(root, research_dir)

    return {
        "schema_version": AUDIT_SCHEMA_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "dataset_dir": str(root),
        "read_only": True,
        "hashing": {"included": include_sha256, "algorithm": "sha256" if include_sha256 else None},
        "inventory": inventory,
        "company_master": company_master,
        "database_schema": database_schema,
        "question_sets": question_sets,
        "financial_reports": financial_reports,
        "research_reports": research_reports,
        "gold_readiness": _audit_gold_readiness(root),
        "accuracy": {
            "status": "unavailable",
            "value": None,
            "reason": "Dataset audit does not evaluate predictions against a verified real golden dataset.",
        },
        "errors": component_errors,
    }


def write_audit_report(report: dict[str, Any], output: str | Path, *, dataset_dir: str | Path) -> Path:
    """Write one JSON artifact, rejecting any destination inside the source dataset."""
    root = Path(dataset_dir).expanduser().resolve()
    output_path = Path(output).expanduser().resolve()
    if output_path == root or output_path.is_relative_to(root):
        raise ValueError("output must be outside dataset_dir to preserve the read-only source boundary")
    if output_path.exists() and output_path.is_dir():
        raise ValueError(f"output must be a file path: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return output_path


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="只读审计 Smart FinQA 正式数据集，并输出 JSON 报告。")
    parser.add_argument("--dataset-dir", required=True, help="正式数据目录")
    parser.add_argument("--output", required=True, help="审计 JSON 输出路径，必须位于正式数据目录之外")
    parser.add_argument(
        "--include-sha256",
        action="store_true",
        help="为所有文件计算 SHA-256；默认关闭以避免全量读取大文件",
    )
    args = parser.parse_args(argv)
    try:
        report = audit_dataset(args.dataset_dir, include_sha256=args.include_sha256)
        output_path = write_audit_report(report, args.output, dataset_dir=args.dataset_dir)
    except (OSError, ValueError) as exc:
        print(f"数据审计失败：{exc}", file=sys.stderr)
        return 2
    print(f"审计报告已写入：{output_path}")
    return 0


def _build_inventory(
    root: Path,
    files: list[Path],
    financial_dir: Path | None,
    research_dir: Path | None,
    include_sha256: bool,
) -> dict[str, Any]:
    entries: list[dict[str, Any]] = []
    category_counts: Counter[str] = Counter()
    extension_counts: Counter[str] = Counter()
    temporary_count = 0
    usable_bytes = 0
    for path in files:
        temporary = _is_temporary_file(path)
        category = _file_category(path, financial_dir, research_dir, temporary)
        size_bytes = path.stat().st_size
        entry: dict[str, Any] = {
            "path": _relative_path(root, path),
            "size_bytes": size_bytes,
            "category": category,
            "temporary": temporary,
        }
        if include_sha256:
            entry["sha256"] = _sha256_file(path)
        entries.append(entry)
        extension_counts[path.suffix.lower() or "<none>"] += 1
        if temporary:
            temporary_count += 1
        else:
            category_counts[category] += 1
            usable_bytes += size_bytes
    categories = {
        name: category_counts.get(name, 0)
        for name in ("financial_report_pdf", "research_report_pdf", "excel_workbook", "other")
    }
    return {
        "total_files": len(files),
        "usable_files": len(files) - temporary_count,
        "temporary_files": temporary_count,
        "total_bytes": sum(item["size_bytes"] for item in entries),
        "usable_bytes": usable_bytes,
        "categories": categories,
        "extension_counts": dict(sorted(extension_counts.items())),
        "files": entries,
    }


def _audit_company_master(root: Path, path: Path | None) -> tuple[dict[str, Any], CompanyIndex]:
    if path is None:
        return (
            {"status": "missing", "path": None, "company_count": 0, "error": "attachment 1 not found"},
            _empty_company_index(),
        )
    try:
        index = CompanyIndex.from_xlsx(path)
    except Exception as exc:
        return (
            {
                "status": "error",
                "path": _relative_path(root, path),
                "company_count": 0,
                "error": _error_message(exc),
            },
            _empty_company_index(),
        )
    return (
        {
            "status": "ok",
            "path": _relative_path(root, path),
            "company_count": len(index.code_to_abbr),
        },
        index,
    )


def _audit_database_schema(root: Path, path: Path | None) -> dict[str, Any]:
    empty_tables = {table: {"present": False, "field_count": 0} for table in SHEET_TO_TABLE.values()}
    if path is None:
        return {
            "status": "missing",
            "path": None,
            "complete": False,
            "tables": empty_tables,
            "error": "attachment 3 not found",
        }
    try:
        schema = load_schema_from_xlsx(path)
    except Exception as exc:
        return {
            "status": "error",
            "path": _relative_path(root, path),
            "complete": False,
            "tables": empty_tables,
            "error": _error_message(exc),
        }
    tables = {
        table: {"present": table in schema, "field_count": len(schema.get(table, []))}
        for table in SHEET_TO_TABLE.values()
    }
    complete = all(item["present"] and item["field_count"] > 0 for item in tables.values())
    return {
        "status": "ok" if complete else "incomplete",
        "path": _relative_path(root, path),
        "complete": complete,
        "tables": tables,
        "error": None,
    }


def _audit_question_set(root: Path, path: Path | None) -> dict[str, Any]:
    if path is None:
        return {
            "status": "missing",
            "path": None,
            "case_count": 0,
            "question_count": 0,
            "errors": [{"error_type": "MissingInput", "message": "workbook not found"}],
        }
    case_count = 0
    question_count = 0
    errors: list[dict[str, Any]] = []
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            rows = sheet.iter_rows(values_only=True)
            header_row = next(rows, None)
            if not header_row:
                raise ValueError("question workbook has no header")
            headers = [str(value or "").strip() for value in header_row]
            try:
                question_column = headers.index("问题")
            except ValueError as exc:
                raise ValueError("question workbook is missing the 问题 column") from exc
            for row_number, row in enumerate(rows, start=2):
                if not any(value is not None for value in row):
                    continue
                case_count += 1
                raw_questions = row[question_column] if question_column < len(row) else None
                try:
                    questions = json.loads(raw_questions) if isinstance(raw_questions, str) else raw_questions
                except json.JSONDecodeError as exc:
                    errors.append(
                        {
                            "row": row_number,
                            "error_type": "JSONDecodeError",
                            "message": str(exc),
                        }
                    )
                    continue
                if not isinstance(questions, list):
                    errors.append(
                        {
                            "row": row_number,
                            "error_type": "InvalidQuestionList",
                            "message": "问题 must be a JSON array",
                        }
                    )
                    continue
                for item_index, item in enumerate(questions, start=1):
                    if not isinstance(item, dict) or not isinstance(item.get("Q"), str) or not item["Q"].strip():
                        errors.append(
                            {
                                "row": row_number,
                                "item": item_index,
                                "error_type": "InvalidQuestion",
                                "message": "each question item must contain a non-empty string Q",
                            }
                        )
                        continue
                    question_count += 1
        finally:
            workbook.close()
    except Exception as exc:
        return {
            "status": "error",
            "path": _relative_path(root, path),
            "case_count": case_count,
            "question_count": 0,
            "errors": [{"error_type": type(exc).__name__, "message": _error_message(exc)}],
        }
    return {
        "status": "error" if errors else "ok",
        "path": _relative_path(root, path),
        "case_count": case_count,
        "question_count": question_count,
        "errors": errors,
    }


def _audit_financial_reports(
    root: Path,
    pdf_files: list[Path],
    company_index: CompanyIndex,
    *,
    report_inspector: ReportInspector | None,
    report_selector: ReportSelector | None,
) -> dict[str, Any]:
    inspector = report_inspector or _default_report_inspector
    selector = report_selector or _default_report_selector
    records: list[Any] = []
    errors: list[dict[str, str]] = []
    page_counts: list[int] = []
    head_available = 0
    head_unavailable: list[str] = []
    summary_count = 0
    english_count = 0
    revision_counts: Counter[str] = Counter()
    outside_paths: defaultdict[str, set[str]] = defaultdict(set)
    outside_name_codes: defaultdict[str, set[str]] = defaultdict(set)
    failed_paths: list[tuple[Path, str]] = []

    known_codes = set(company_index.code_to_abbr)
    for path in pdf_files:
        relative_path = _relative_path(root, path)
        try:
            metadata = inspector(path, company_index)
        except Exception as exc:
            errors.append(_file_error(root, path, "inspect_report_file_metadata", exc))
            failed_paths.append((path, relative_path))
            filename_code = _stock_code_from_filename(path.name)
            if known_codes and filename_code and filename_code not in known_codes:
                outside_paths[filename_code].add(relative_path)
            continue
        records.append(metadata)
        stock_code = str(getattr(metadata, "stock_code", "") or "")
        company_in_master = getattr(metadata, "company_in_master", stock_code in known_codes)
        if known_codes and stock_code and not company_in_master:
            outside_paths[stock_code].add(relative_path)
            stock_abbr = str(getattr(metadata, "stock_abbr", "") or "").strip()
            if stock_abbr:
                outside_name_codes[_normalize_company_name(stock_abbr)].add(stock_code)
        read_error = getattr(metadata, "read_error", None)
        if read_error:
            errors.append(
                {
                    "path": relative_path,
                    "stage": "pdf_head_extraction",
                    "error_type": "PartialReadError",
                    "message": str(read_error),
                }
            )
        page_count = getattr(metadata, "page_count", None)
        if isinstance(page_count, int) and not isinstance(page_count, bool) and page_count > 0:
            page_counts.append(page_count)
        head_text_length = getattr(metadata, "head_text_length", 0)
        if isinstance(head_text_length, int) and head_text_length >= MINIMUM_HEAD_TEXT_CHARACTERS:
            head_available += 1
        else:
            head_unavailable.append(relative_path)
        summary_count += bool(getattr(metadata, "is_summary", False))
        english_count += bool(getattr(metadata, "is_english", False))
        revision_kind = getattr(metadata, "revision_kind", None)
        if revision_kind is not None:
            revision_counts[_enum_value(revision_kind)] += 1

    for path, relative_path in failed_paths:
        company_name = _company_name_from_filename(path.name)
        codes = outside_name_codes.get(company_name, set())
        if len(codes) == 1:
            outside_paths[next(iter(codes))].add(relative_path)

    decisions: list[dict[str, str]] = []
    selection_statuses: Counter[str] = Counter()
    if records:
        try:
            selected = selector(records)
        except Exception as exc:
            errors.append(
                {
                    "path": "<financial-report-set>",
                    "stage": "select_authoritative_reports",
                    "error_type": type(exc).__name__,
                    "message": _error_message(exc),
                }
            )
        else:
            for decision in selected:
                status = str(getattr(decision, "status", "unknown"))
                decision_path = Path(getattr(decision, "path"))
                selection_statuses[status] += 1
                decisions.append(
                    {
                        "path": _relative_path(root, decision_path),
                        "status": status,
                        "reason": str(getattr(decision, "reason", "")),
                    }
                )

    selected_outside_files = {item["path"] for item in decisions if item["reason"] == "outside_company_master"}
    outside_files = sorted(selected_outside_files | {path for paths in outside_paths.values() for path in paths})
    return {
        "discovered": len(pdf_files),
        "inspected": len(records),
        "metadata_coverage": _coverage(len(records), len(pdf_files)),
        "page_count": {
            "known_files": len(page_counts),
            "unknown_files": len(records) - len(page_counts),
            "total": sum(page_counts),
            "minimum": min(page_counts) if page_counts else None,
            "maximum": max(page_counts) if page_counts else None,
        },
        "head_text": {
            "minimum_characters": MINIMUM_HEAD_TEXT_CHARACTERS,
            "available": head_available,
            "unavailable": len(records) - head_available,
            "unavailable_files": sorted(head_unavailable),
        },
        "classifications": {
            "summary": summary_count,
            "english": english_count,
            "revision_kinds": dict(sorted(revision_counts.items())),
        },
        "selection_statuses": dict(sorted(selection_statuses.items())),
        "selection_decisions": sorted(decisions, key=lambda item: item["path"]),
        "conflict_count": selection_statuses.get("needs_review", 0),
        "outside_company_master": {
            "status": "available" if known_codes else "unavailable",
            "stock_codes": sorted(outside_paths),
            "file_count": len(outside_files),
            "files": outside_files,
        },
        "errors": errors,
    }


def _audit_research_reports(root: Path, research_dir: Path | None) -> dict[str, Any]:
    pdf_files = _pdf_files(research_dir)
    metadata_files = (
        []
        if research_dir is None
        else sorted(
            path
            for path in research_dir.glob("*.xlsx")
            if not _is_temporary_file(path) and path.name in {"个股_研报信息.xlsx", "行业_研报信息.xlsx"}
        )
    )
    rows = 0
    title_rows = 0
    contextual_rows = 0
    publication_rows = 0
    errors: list[dict[str, str]] = []
    title_records: defaultdict[str, list[dict[str, Any]]] = defaultdict(list)
    for workbook_path in metadata_files:
        try:
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            try:
                sheet = workbook[workbook.sheetnames[0]]
                iterator = sheet.iter_rows(values_only=True)
                header_row = next(iterator, None)
                if not header_row:
                    raise ValueError("metadata workbook has no header")
                headers = [str(value or "").strip() for value in header_row]
                for row_number, row in enumerate(iterator, start=2):
                    if not any(value is not None for value in row):
                        continue
                    rows += 1
                    values = {headers[index]: row[index] for index in range(min(len(headers), len(row)))}
                    title = str(values.get("title") or "").strip()
                    if not title:
                        errors.append(
                            {
                                "path": _relative_path(root, workbook_path),
                                "stage": "research_metadata",
                                "error_type": "MissingTitle",
                                "message": f"row {row_number} has no title",
                            }
                        )
                        continue
                    title_rows += 1
                    if any(str(values.get(key) or "").strip() for key in ("stockName", "indvInduName", "industryName")):
                        contextual_rows += 1
                    if str(values.get("publishDate") or "").strip():
                        publication_rows += 1
                    title_records[_research_title_key(title)].append(
                        {
                            "path": _relative_path(root, workbook_path),
                            "row": row_number,
                            "values": values,
                        }
                    )
            finally:
                workbook.close()
        except Exception as exc:
            errors.append(_file_error(root, workbook_path, "research_metadata", exc))

    conflicts: list[dict[str, Any]] = []
    usable_metadata: dict[str, dict[str, Any]] = {}
    for title_key, records in sorted(title_records.items()):
        signatures = {_metadata_signature(record["values"]) for record in records}
        if len(signatures) > 1:
            conflicts.append(
                {
                    "title_key": title_key,
                    "locations": [f"{record['path']}:{record['row']}" for record in records],
                    "reason": "multiple metadata rows resolve to the same PDF title",
                }
            )
        else:
            usable_metadata[title_key] = records[0]

    pdf_by_key: defaultdict[str, list[Path]] = defaultdict(list)
    for path in pdf_files:
        pdf_by_key[_research_title_key(path.stem)].append(path)
    matched_paths = [path for key, paths in pdf_by_key.items() if key in usable_metadata for path in paths]
    missing_paths = [path for key, paths in pdf_by_key.items() if key not in usable_metadata for path in paths]
    orphan_titles = sorted(key for key in usable_metadata if key not in pdf_by_key)
    duplicate_pdf_titles = [
        {"title_key": key, "paths": sorted(_relative_path(root, path) for path in paths)}
        for key, paths in sorted(pdf_by_key.items())
        if len(paths) > 1
    ]
    return {
        "discovered": len(pdf_files),
        "categories": {
            "stock": sum("个股研报" in path.parts for path in pdf_files),
            "industry": sum("行业研报" in path.parts for path in pdf_files),
        },
        "metadata": {
            "workbooks": [_relative_path(root, path) for path in metadata_files],
            "rows": rows,
            "title_rows": title_rows,
            "context_rows": contextual_rows,
            "publication_date_rows": publication_rows,
            "conflict_count": len(conflicts),
            "conflicts": conflicts,
            "matched_pdf_count": len(matched_paths),
            "coverage": _coverage(len(matched_paths), len(pdf_files)),
            "missing_pdf_metadata": sorted(_relative_path(root, path) for path in missing_paths),
            "orphan_metadata_titles": orphan_titles,
            "duplicate_pdf_titles": duplicate_pdf_titles,
        },
        "errors": errors,
    }


def _audit_gold_readiness(root: Path) -> dict[str, Any]:
    manifests = sorted(path for path in root.rglob("manifest.json") if not _is_temporary_file(path))
    if not manifests:
        return {
            "status": "not_ready",
            "manifest": None,
            "verification_status": "not_run",
            "blockers": ["missing_manifest"],
            "note": "Formal source files are not a verified real golden dataset without a valid manifest and annotations.",
        }
    if len(manifests) > 1:
        return {
            "status": "not_ready",
            "manifest": None,
            "verification_status": "not_run",
            "blockers": ["multiple_manifests"],
            "candidates": [_relative_path(root, path) for path in manifests],
            "note": "Select one versioned golden manifest before evaluation.",
        }
    manifest = manifests[0]
    try:
        payload = json.loads(manifest.read_text(encoding="utf-8"))
        if not isinstance(payload, dict):
            raise ValueError("manifest root must be an object")
    except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as exc:
        return {
            "status": "not_ready",
            "manifest": _relative_path(root, manifest),
            "verification_status": "failed",
            "blockers": ["invalid_manifest"],
            "error": _error_message(exc),
        }
    annotation_files = payload.get("annotation_files")
    missing_annotations = not isinstance(annotation_files, list) or not annotation_files
    blockers = ["missing_annotation_files"] if missing_annotations else []
    blockers.append("manifest_requires_evaluation_loader_verification")
    return {
        "status": "not_ready",
        "manifest": _relative_path(root, manifest),
        "verification_status": "not_run",
        "blockers": blockers,
        "declared_dataset_kind": payload.get("dataset_kind"),
        "declared_availability": payload.get("availability"),
        "note": "Run the versioned evaluation loader and a prediction set before reporting any accuracy metric.",
    }


def _find_attachment_directory(
    root: Path, prefix: str, required_text: str, errors: list[dict[str, str]]
) -> Path | None:
    candidates = sorted(
        path for path in root.iterdir() if path.is_dir() and path.name.startswith(prefix) and required_text in path.name
    )
    return _choose_single_path(root, candidates, f"{prefix} directory", errors)


def _find_attachment_workbook(root: Path, prefix: str, errors: list[dict[str, str]]) -> Path | None:
    candidates = sorted(
        path
        for path in root.iterdir()
        if path.is_file()
        and path.suffix.lower() == ".xlsx"
        and path.name.startswith(prefix)
        and not _is_temporary_file(path)
    )
    return _choose_single_path(root, candidates, f"{prefix} workbook", errors)


def _choose_single_path(root: Path, candidates: list[Path], label: str, errors: list[dict[str, str]]) -> Path | None:
    if not candidates:
        errors.append(
            {"path": "<dataset>", "stage": "discovery", "error_type": "MissingInput", "message": f"{label} not found"}
        )
        return None
    if len(candidates) > 1:
        errors.append(
            {
                "path": "<dataset>",
                "stage": "discovery",
                "error_type": "AmbiguousInput",
                "message": f"multiple {label} candidates: {', '.join(_relative_path(root, item) for item in candidates)}",
            }
        )
        return None
    return candidates[0]


def _file_category(path: Path, financial_dir: Path | None, research_dir: Path | None, temporary: bool) -> str:
    if temporary:
        return "temporary"
    if path.suffix.lower() == ".pdf" and financial_dir is not None and path.is_relative_to(financial_dir):
        return "financial_report_pdf"
    if path.suffix.lower() == ".pdf" and research_dir is not None and path.is_relative_to(research_dir):
        return "research_report_pdf"
    if path.suffix.lower() in {".xls", ".xlsx", ".xlsm"}:
        return "excel_workbook"
    return "other"


def _pdf_files(directory: Path | None) -> list[Path]:
    if directory is None:
        return []
    return sorted(path for path in directory.rglob("*") if path.is_file() and path.suffix.lower() == ".pdf")


def _is_temporary_file(path: Path) -> bool:
    lower_name = path.name.lower()
    return (
        path.name.startswith("~$")
        or lower_name.startswith(".")
        or lower_name.endswith("~")
        or path.suffix.lower() in TEMPORARY_SUFFIXES
    )


def _stock_code_from_filename(file_name: str) -> str | None:
    match = STOCK_CODE_PATTERN.search(file_name)
    return match.group(1) if match else None


def _company_name_from_filename(file_name: str) -> str:
    return _normalize_company_name(re.split(r"[:：]", Path(file_name).stem, maxsplit=1)[0])


def _normalize_company_name(value: str) -> str:
    return re.sub(r"\s+", "", value.strip())


def _research_title_key(title: str) -> str:
    return re.sub(r"\s+", " ", title.strip()).replace("/", "_").replace("\\", "_")


def _metadata_signature(values: dict[str, Any]) -> str:
    normalized = {key: _json_scalar(value) for key, value in sorted(values.items())}
    return json.dumps(normalized, ensure_ascii=False, sort_keys=True)


def _json_scalar(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _default_report_inspector(path: Path, company_index: CompanyIndex) -> Any:
    from .ingestion import inspect_report_file_metadata

    return inspect_report_file_metadata(path, company_index)


def _default_report_selector(records: Sequence[Any]) -> Sequence[Any]:
    from .ingestion import select_authoritative_reports

    return select_authoritative_reports(list(records))


def _empty_company_index() -> CompanyIndex:
    return CompanyIndex(code_to_abbr={}, abbr_to_code={})


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _relative_path(root: Path, path: Path) -> str:
    resolved = path.resolve()
    try:
        return resolved.relative_to(root).as_posix()
    except ValueError:
        return str(resolved)


def _file_error(root: Path, path: Path, stage: str, exc: Exception) -> dict[str, str]:
    return {
        "path": _relative_path(root, path),
        "stage": stage,
        "error_type": type(exc).__name__,
        "message": _error_message(exc),
    }


def _error_message(exc: Exception) -> str:
    return str(exc).strip() or type(exc).__name__


def _coverage(numerator: int, denominator: int) -> float:
    return numerator / denominator if denominator else 0.0


def _enum_value(value: Any) -> str:
    enum_value = getattr(value, "value", value)
    return str(enum_value)


if __name__ == "__main__":
    raise SystemExit(main())
