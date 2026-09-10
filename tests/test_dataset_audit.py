from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from smart_finqa.dataset_audit import audit_dataset, main, write_audit_report
from tests.helpers import write_company_workbook, write_schema_workbook


@dataclass(frozen=True)
class FakeReportMetadata:
    path: Path
    stock_code: str
    stock_abbr: str
    page_count: int
    head_text_length: int
    is_summary: bool = False
    is_english: bool = False
    company_in_master: bool = True


def _write_question_workbook(path: Path, count: int) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["编号", "问题类型", "问题"])
    for index in range(count):
        questions = [{"Q": f"问题 {index + 1}"}]
        if index == 0:
            questions.append({"Q": "追问"})
        sheet.append([f"Q{index + 1}", "测试", json.dumps(questions, ensure_ascii=False)])
    workbook.save(path)
    return path


def _write_research_metadata(path: Path, rows: list[tuple[object, ...]], headers: tuple[str, ...]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


def _build_dataset(tmp_path: Path) -> Path:
    dataset_dir = tmp_path / "正式数据"
    annual_dir = dataset_dir / "附件2：财务报告" / "reports"
    stock_research_dir = dataset_dir / "附件5：研报数据" / "个股研报"
    industry_research_dir = dataset_dir / "附件5：研报数据" / "行业研报"
    annual_dir.mkdir(parents=True)
    stock_research_dir.mkdir(parents=True)
    industry_research_dir.mkdir(parents=True)

    for name in ("600001_a.pdf", "600099_b.pdf", "600001_summary.pdf", "外部药业：2024年年度报告.pdf"):
        (annual_dir / name).write_bytes(f"annual:{name}".encode())
    (stock_research_dir / "匹配研报.pdf").write_bytes(b"stock")
    (stock_research_dir / "重复_标题.pdf").write_bytes(b"conflict")
    (industry_research_dir / "行业展望.pdf").write_bytes(b"industry")

    write_company_workbook(dataset_dir / "附件1：公司基本信息.xlsx")
    write_schema_workbook(dataset_dir / "附件3：数据库-表名及字段说明.xlsx")
    _write_question_workbook(dataset_dir / "附件4：问题汇总.xlsx", 2)
    _write_question_workbook(dataset_dir / "附件6：问题汇总.xlsx", 1)
    (dataset_dir / "~$附件4：问题汇总.xlsx").write_bytes(b"temporary")
    (dataset_dir / "说明.txt").write_text("read me", encoding="utf-8")

    research_dir = dataset_dir / "附件5：研报数据"
    _write_research_metadata(
        research_dir / "个股_研报信息.xlsx",
        [
            ("匹配研报", "金花股份", "2025-01-01"),
            ("重复/标题", "金花股份", "2025-01-02"),
            ("重复_标题", "金花股份", "2025-01-03"),
        ],
        ("title", "stockName", "publishDate"),
    )
    _write_research_metadata(
        research_dir / "行业_研报信息.xlsx",
        [("行业展望", "中药", "2025-01-04")],
        ("title", "industryName", "publishDate"),
    )
    return dataset_dir


def _inspect_report(path: Path, _company_index: object) -> FakeReportMetadata:
    if path.name == "外部药业：2024年年度报告.pdf":
        raise ValueError("damaged xref")
    outside = "600099" in path.name
    return FakeReportMetadata(
        path=path,
        stock_code="600099" if outside else "600080",
        stock_abbr="外部药业" if outside else "金花股份",
        page_count=5 if "summary" in path.name else (10 if "600099" in path.name else 100),
        head_text_length=20 if path.name == "600001_a.pdf" else 100,
        is_summary="summary" in path.name,
        company_in_master=not outside,
    )


def _select_reports(records: list[FakeReportMetadata]) -> list[SimpleNamespace]:
    return [
        SimpleNamespace(
            path=record.path,
            status="skipped_summary" if record.is_summary else "selected",
            reason="summary_only" if record.is_summary else "authoritative_version",
        )
        for record in records
    ]


def test_audit_dataset_reports_formal_data_coverage_without_claiming_accuracy(tmp_path: Path) -> None:
    dataset_dir = _build_dataset(tmp_path)

    report = audit_dataset(dataset_dir, report_inspector=_inspect_report, report_selector=_select_reports)

    assert report["read_only"] is True
    assert report["inventory"]["total_files"] == 15
    assert report["inventory"]["usable_files"] == 14
    assert report["inventory"]["temporary_files"] == 1
    assert report["inventory"]["categories"] == {
        "financial_report_pdf": 4,
        "research_report_pdf": 3,
        "excel_workbook": 6,
        "other": 1,
    }
    assert report["company_master"]["company_count"] == 5
    assert report["database_schema"]["complete"] is True
    assert set(report["database_schema"]["tables"]) == {
        "core_performance_indicators_sheet",
        "balance_sheet",
        "cash_flow_sheet",
        "income_sheet",
    }
    assert report["question_sets"]["attachment_4"]["case_count"] == 2
    assert report["question_sets"]["attachment_4"]["question_count"] == 3
    assert report["question_sets"]["attachment_6"]["case_count"] == 1
    assert report["question_sets"]["attachment_6"]["question_count"] == 2

    annual = report["financial_reports"]
    assert annual["discovered"] == 4
    assert annual["inspected"] == 3
    assert annual["metadata_coverage"] == pytest.approx(0.75)
    assert annual["page_count"]["total"] == 115
    assert annual["head_text"]["minimum_characters"] == 80
    assert annual["head_text"]["available"] == 2
    assert annual["head_text"]["unavailable"] == 1
    assert annual["head_text"]["unavailable_files"] == ["附件2：财务报告/reports/600001_a.pdf"]
    assert annual["classifications"]["summary"] == 1
    assert annual["selection_statuses"] == {"selected": 2, "skipped_summary": 1}
    assert annual["outside_company_master"]["stock_codes"] == ["600099"]
    assert annual["outside_company_master"]["file_count"] == 2
    assert annual["errors"][0]["path"].endswith("外部药业：2024年年度报告.pdf")
    assert annual["errors"][0]["error_type"] == "ValueError"

    research = report["research_reports"]
    assert research["discovered"] == 3
    assert research["metadata"]["rows"] == 4
    assert research["metadata"]["conflict_count"] == 1
    assert research["metadata"]["matched_pdf_count"] == 2
    assert research["metadata"]["coverage"] == pytest.approx(2 / 3)
    assert research["metadata"]["missing_pdf_metadata"] == ["附件5：研报数据/个股研报/重复_标题.pdf"]

    assert report["gold_readiness"]["status"] == "not_ready"
    assert "missing_manifest" in report["gold_readiness"]["blockers"]
    assert report["accuracy"] == {
        "status": "unavailable",
        "value": None,
        "reason": "Dataset audit does not evaluate predictions against a verified real golden dataset.",
    }
    assert report["hashing"] == {"included": False, "algorithm": None}
    assert all("sha256" not in item for item in report["inventory"]["files"])


def test_question_audit_reports_invalid_json_and_missing_q(tmp_path: Path) -> None:
    dataset_dir = _build_dataset(tmp_path)
    workbook_path = dataset_dir / "附件6：问题汇总.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["编号", "问题类型", "问题"])
    sheet.append(["Q1", "测试", '[{"Q":"有效问题"},{"answer":"缺少问题"}]'])
    sheet.append(["Q2", "测试", "not-json"])
    workbook.save(workbook_path)

    report = audit_dataset(dataset_dir, report_inspector=_inspect_report, report_selector=_select_reports)
    questions = report["question_sets"]["attachment_6"]

    assert questions["status"] == "error"
    assert questions["case_count"] == 2
    assert questions["question_count"] == 1
    assert [item["error_type"] for item in questions["errors"]] == ["InvalidQuestion", "JSONDecodeError"]


def test_outside_company_file_without_stock_code_uses_authoritative_decision(tmp_path: Path) -> None:
    dataset_dir = _build_dataset(tmp_path)

    def inspect(path: Path, company_index: object) -> FakeReportMetadata:
        if path.name == "外部药业：2024年年度报告.pdf":
            return FakeReportMetadata(
                path=path,
                stock_code="",
                stock_abbr="外部药业",
                page_count=100,
                head_text_length=100,
                company_in_master=False,
            )
        return _inspect_report(path, company_index)

    def select(records: list[FakeReportMetadata]) -> list[SimpleNamespace]:
        return [
            SimpleNamespace(
                path=record.path,
                status="needs_review" if not record.company_in_master else "selected",
                reason="outside_company_master" if not record.company_in_master else "authoritative_version",
            )
            for record in records
        ]

    report = audit_dataset(dataset_dir, report_inspector=inspect, report_selector=select)

    outside = report["financial_reports"]["outside_company_master"]
    assert outside["file_count"] == 2
    assert "附件2：财务报告/reports/外部药业：2024年年度报告.pdf" in outside["files"]


def test_include_sha256_hashes_usable_and_temporary_files(tmp_path: Path) -> None:
    dataset_dir = _build_dataset(tmp_path)

    report = audit_dataset(
        dataset_dir,
        include_sha256=True,
        report_inspector=_inspect_report,
        report_selector=_select_reports,
    )

    files = report["inventory"]["files"]
    assert report["hashing"] == {"included": True, "algorithm": "sha256"}
    assert all(len(item["sha256"]) == 64 for item in files)
    info_file = next(item for item in files if item["path"] == "说明.txt")
    assert info_file["sha256"] == hashlib.sha256("read me".encode()).hexdigest()


def test_write_audit_report_rejects_output_inside_dataset(tmp_path: Path) -> None:
    dataset_dir = _build_dataset(tmp_path)
    report = {"status": "test"}

    with pytest.raises(ValueError, match="outside dataset_dir"):
        write_audit_report(report, dataset_dir / "audit.json", dataset_dir=dataset_dir)

    assert not (dataset_dir / "audit.json").exists()


def test_cli_writes_only_requested_output(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    dataset_dir = tmp_path / "input"
    dataset_dir.mkdir()
    output_path = tmp_path / "isolated-output" / "audit.json"
    monkeypatch.setattr(
        "smart_finqa.dataset_audit.audit_dataset",
        lambda dataset_dir, include_sha256=False: {
            "dataset_dir": str(dataset_dir),
            "include_sha256": include_sha256,
        },
    )

    exit_code = main(
        [
            "--dataset-dir",
            str(dataset_dir),
            "--output",
            str(output_path),
            "--include-sha256",
        ]
    )

    assert exit_code == 0
    assert json.loads(output_path.read_text(encoding="utf-8"))["include_sha256"] is True
    assert list(dataset_dir.iterdir()) == []
