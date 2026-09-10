from __future__ import annotations

import gc
import json
import threading
import weakref
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from smart_finqa.config import AppConfig, DatabaseConfig, LLMConfig
from smart_finqa.ingestion import ReportFileMetadata, ReportRecord, ReportRevisionKind
from smart_finqa.ingestion_state import IngestionStateWriteError, dataset_source_uri
from smart_finqa.pipeline import INGESTION_STATE_VERSION, PipelinePaths, SmartFinancePipeline


def _write_schema_workbook(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name in ["核心业绩指标表", "资产负债表", "现金流量表", "利润表"]:
        ws = wb.create_sheet(sheet_name)
        ws.append(["字段名称", "中文名称", "字段类型", "字段说明"])
        ws.append(["stock_code", "股票代码", "varchar", "股票代码"])
        ws.append(["report_period", "报告期", "varchar", "报告期"])
        if sheet_name == "利润表":
            ws.append(["total_profit", "利润总额", "decimal", "利润总额"])
    wb.save(path)


def _write_company_workbook(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "基本信息表"
    ws.append(["序号", "股票代码", "A股简称", "公司名称"])
    ws.append([1, 600000, "测试公司", "测试公司股份有限公司"])
    wb.save(path)


def _report_metadata(path: Path, year: int = 2024) -> ReportFileMetadata:
    return ReportFileMetadata(
        path=path,
        stock_code="600000",
        stock_abbr="测试公司",
        report_period=f"{year}FY",
        report_year=year,
        is_summary=False,
        is_english=False,
        revision_kind=ReportRevisionKind.ORIGINAL,
        published_date=None,
        copy_index=0,
        page_count=100,
        head_text_length=1000,
    )


def test_run_mode_ingest_only(tmp_path: Path) -> None:
    test_data = tmp_path / "数据" / "测试数据"
    _write_schema_workbook(test_data / "附件3：数据库-表名及字段说明.xlsx")
    _write_company_workbook(test_data / "附件1：上市公司基本信息.xlsx")

    paths = PipelinePaths.from_base_dir(tmp_path, full_data=True)
    cfg = AppConfig(
        mode="ingest",
        full_data=True,
        db=DatabaseConfig(backend="sqlite", sqlite_path=str(tmp_path / "finance.db")),
        llm=LLMConfig(),
    )
    pipeline = SmartFinancePipeline(paths, app_config=cfg)
    pipeline.ingest_reports = lambda: None
    pipeline.validate_database = lambda: {"ok": True}
    pipeline._ensure_database_ready = lambda: None

    outputs = pipeline.run(mode="ingest")

    assert "db_path" in outputs
    assert outputs["result_2"] == ""
    assert outputs["result_3"] == ""
    run_log = json.loads(Path(outputs["run_log"]).read_text(encoding="utf-8"))
    validation_report = json.loads(Path(outputs["validation_report"]).read_text(encoding="utf-8"))
    assert run_log["schema_version"] == 2
    assert run_log["status"] == "SUCCEEDED"
    assert run_log["started_at"]
    assert run_log["finished_at"]
    assert validation_report["schema_version"] == 2
    assert validation_report["run_id"] == run_log["run_id"]
    assert run_log["validation"] == validation_report


def test_run_resets_log_identity_and_records_failure_status(tmp_path: Path) -> None:
    test_data = tmp_path / "数据" / "测试数据"
    _write_schema_workbook(test_data / "附件3：数据库-表名及字段说明.xlsx")
    _write_company_workbook(test_data / "附件1：上市公司基本信息.xlsx")
    paths = PipelinePaths.from_base_dir(tmp_path, full_data=True)
    config = AppConfig(
        mode="ingest",
        full_data=True,
        db=DatabaseConfig(backend="sqlite", sqlite_path=str(tmp_path / "finance.db")),
        llm=LLMConfig(),
    )
    pipeline = SmartFinancePipeline(paths, app_config=config)
    pipeline.ingest_reports = lambda: pipeline._append_ingestion_log({"status": "parsed", "pdf": "ok.pdf"})
    pipeline.validate_database = lambda: {"status": "UNAVAILABLE"}

    first_outputs = pipeline.run(mode="ingest")
    first_log = json.loads(Path(first_outputs["run_log"]).read_text(encoding="utf-8"))
    second_outputs = pipeline.run(mode="ingest")
    second_log = json.loads(Path(second_outputs["run_log"]).read_text(encoding="utf-8"))

    assert first_log["run_id"] != second_log["run_id"]
    assert len(first_log["ingestion"]) == 1
    assert len(second_log["ingestion"]) == 1

    def fail_ingestion() -> None:
        raise RuntimeError("synthetic ingestion failure")

    pipeline.ingest_reports = fail_ingestion
    with pytest.raises(RuntimeError, match="synthetic ingestion failure"):
        pipeline.run(mode="ingest")

    failed_log = json.loads((paths.output_dir / "run_log.json").read_text(encoding="utf-8"))
    assert failed_log["status"] == "FAILED"
    assert failed_log["finished_at"]
    assert failed_log["failure"] == {"type": "RuntimeError", "message": "synthetic ingestion failure"}


def test_ingestion_log_samples_benign_events_but_preserves_attention_events(tmp_path: Path) -> None:
    test_data = tmp_path / "数据" / "测试数据"
    _write_schema_workbook(test_data / "附件3：数据库-表名及字段说明.xlsx")
    _write_company_workbook(test_data / "附件1：上市公司基本信息.xlsx")
    config = AppConfig(
        full_data=True,
        ingestion_log_limit=1,
        db=DatabaseConfig(backend="sqlite", sqlite_path=str(tmp_path / "finance.db")),
        llm=LLMConfig(),
    )
    pipeline = SmartFinancePipeline(PipelinePaths.from_base_dir(tmp_path, full_data=True), app_config=config)
    pipeline.run_log = {"ingestion": []}

    for index in range(2):
        pipeline._append_ingestion_log({"status": "parsed", "pdf": f"parsed-{index}.pdf"})
        pipeline._append_ingestion_log({"status": "skipped_summary", "pdf": f"summary-{index}.pdf"})
    for status in ("needs_review", "needs_review", "skip_invalid", "error"):
        pipeline._append_ingestion_log({"status": status, "pdf": f"{status}.pdf"})

    retained_statuses = [event["status"] for event in pipeline.run_log["ingestion"]]
    assert retained_statuses.count("parsed") == 1
    assert retained_statuses.count("skipped_summary") == 2
    assert retained_statuses.count("needs_review") == 2
    assert retained_statuses.count("skip_invalid") == 1
    assert retained_statuses.count("error") == 1
    assert pipeline.run_log["ingestion_event_stats"] == {
        "parsed": {"retention": "sampled", "total": 2, "sampled": 1, "dropped": 1},
        "skipped_summary": {"retention": "full", "total": 2, "sampled": 2, "dropped": 0},
        "needs_review": {"retention": "full", "total": 2, "sampled": 2, "dropped": 0},
        "skip_invalid": {"retention": "full", "total": 1, "sampled": 1, "dropped": 0},
        "error": {"retention": "full", "total": 1, "sampled": 1, "dropped": 0},
    }
    assert pipeline.run_log["ingestion_overflow"] == 1


def test_run_rejects_unknown_mode(tmp_path: Path) -> None:
    test_data = tmp_path / "数据" / "测试数据"
    _write_schema_workbook(test_data / "附件3：数据库-表名及字段说明.xlsx")
    _write_company_workbook(test_data / "附件1：上市公司基本信息.xlsx")
    config = AppConfig(
        db=DatabaseConfig(backend="sqlite", sqlite_path=str(tmp_path / "finance.db")),
        llm=LLMConfig(),
    )
    pipeline = SmartFinancePipeline(PipelinePaths.from_base_dir(tmp_path, full_data=True), app_config=config)

    with pytest.raises(ValueError, match="Unsupported pipeline mode"):
        pipeline.run(mode="unknown")


def test_invalid_ingestion_state_fails_explicitly(tmp_path: Path) -> None:
    path = tmp_path / "ingestion_state.json"
    path.write_text("{not-json", encoding="utf-8")

    with pytest.raises(ValueError, match="Invalid ingestion state JSON"):
        SmartFinancePipeline._load_ingestion_state(path)

    path.write_text("[]", encoding="utf-8")
    with pytest.raises(ValueError, match="Invalid ingestion state object"):
        SmartFinancePipeline._load_ingestion_state(path)

    path.write_text(json.dumps({"version": INGESTION_STATE_VERSION, "files": []}), encoding="utf-8")
    with pytest.raises(ValueError, match="Invalid ingestion state files mapping"):
        SmartFinancePipeline._load_ingestion_state(path)


def test_pipeline_ingestion_state_save_preserves_previous_file_when_replace_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    state_path = tmp_path / "ingestion_state.json"
    original_files = {"report.pdf": {"status": "parsed"}}
    SmartFinancePipeline._save_ingestion_state(state_path, original_files)

    def fail_replace(_source: str | Path, _destination: str | Path) -> None:
        raise PermissionError("synthetic replace failure")

    monkeypatch.setattr("smart_finqa.ingestion_state.os.replace", fail_replace)
    with pytest.raises(IngestionStateWriteError, match="Unable to replace ingestion state"):
        SmartFinancePipeline._save_ingestion_state(state_path, {"other.pdf": {"status": "parsed"}})

    assert SmartFinancePipeline._load_ingestion_state(state_path) == original_files
    assert json.loads(state_path.read_text(encoding="utf-8")) == {
        "version": INGESTION_STATE_VERSION,
        "files": original_files,
    }
    assert list(tmp_path.glob(f".{state_path.name}.*.tmp")) == []


def test_formal_research_metadata_columns_are_supported() -> None:
    metadata = {
        "stockName": "华润三九",
        "indvInduName": "中药",
        "industryName": "医药生物",
        "publishDate": datetime(2025, 12, 31),
    }

    assert SmartFinancePipeline._research_metadata_text(metadata, ("stockName",)) == "华润三九"
    assert SmartFinancePipeline._research_metadata_text(metadata, ("indvInduName",)) == "中药"
    assert SmartFinancePipeline._research_metadata_text(metadata, ("industryName",)) == "医药生物"
    assert SmartFinancePipeline._research_metadata_date(metadata, ("publishDate",)) == date(2025, 12, 31)


def test_research_metadata_matches_windows_filename_and_rejects_ambiguous_rows(tmp_path: Path) -> None:
    research_dir = tmp_path / "附件5：研报数据"
    research_dir.mkdir()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["title", "publishDate"])
    sheet.append(["PD-1/VEGF研究", datetime(2025, 1, 1)])
    sheet.append(["重复标题", datetime(2025, 2, 1)])
    sheet.append(["重复标题", datetime(2025, 3, 1)])
    workbook.save(research_dir / "行业_研报信息.xlsx")

    pipeline = object.__new__(SmartFinancePipeline)
    pipeline.paths = SimpleNamespace(research_dir=research_dir)
    pipeline.run_log = {}
    metadata = pipeline._load_research_metadata()

    assert SmartFinancePipeline._research_title_key("PD-1_VEGF研究") in metadata
    assert SmartFinancePipeline._research_title_key("重复标题") not in metadata
    assert pipeline.run_log["kb_metadata_conflicts"] == [
        {
            "title_key": "重复标题",
            "error": "multiple research metadata rows resolve to the same PDF title",
        }
    ]


def test_ingestion_only_parses_selected_authoritative_report(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    data_dir = tmp_path / "正式数据"
    _write_schema_workbook(data_dir / "附件3：数据库-表名及字段说明.xlsx")
    _write_company_workbook(data_dir / "附件1：上市公司基本信息.xlsx")
    reports_dir = data_dir / "附件2：财务报告"
    reports_dir.mkdir(parents=True)
    full_path = reports_dir / "测试公司：2024年年度报告.pdf"
    summary_path = reports_dir / "测试公司：2024年年度报告摘要.pdf"
    outside_path = reports_dir / "主数据外公司：2024年年度报告.pdf"
    for path in (full_path, summary_path, outside_path):
        path.write_bytes(b"not-read-by-test")

    def metadata(path: Path, **overrides: object) -> ReportFileMetadata:
        values: dict[str, object] = {
            "path": path,
            "stock_code": "600000",
            "stock_abbr": "测试公司",
            "report_period": "2024FY",
            "report_year": 2024,
            "is_summary": False,
            "is_english": False,
            "revision_kind": ReportRevisionKind.ORIGINAL,
            "published_date": None,
            "copy_index": 0,
            "page_count": 100,
            "head_text_length": 1000,
            "company_in_master": True,
        }
        values.update(overrides)
        return ReportFileMetadata(**values)

    metadata_by_name = {
        full_path.name: metadata(full_path),
        summary_path.name: metadata(summary_path, is_summary=True),
        outside_path.name: metadata(
            outside_path,
            stock_code="123456",
            stock_abbr="主数据外公司",
            company_in_master=False,
        ),
    }
    inspected_paths: list[Path] = []

    def inspect(path: Path, _company_index: object, **_kwargs: object) -> ReportFileMetadata:
        inspected_paths.append(path)
        return metadata_by_name[path.name]

    monkeypatch.setattr("smart_finqa.pipeline.inspect_report_file_metadata", inspect)
    parsed_paths: list[Path] = []

    def fake_extract(path: Path, _company_index: object, report: ReportFileMetadata, **_kwargs: object) -> ReportRecord:
        parsed_paths.append(path)
        return ReportRecord(
            source_path=path,
            stock_code=report.stock_code,
            stock_abbr=report.stock_abbr,
            report_period=report.report_period,
            report_year=report.report_year,
            text="有效财报文本" * 20,
            pages=(),
            snapshot={},
        )

    monkeypatch.setattr("smart_finqa.pipeline.extract_report_record", fake_extract)
    config = AppConfig(
        db=DatabaseConfig(backend="sqlite", sqlite_path=str(tmp_path / "finance.db")),
        llm=LLMConfig(),
        incremental_ingest=True,
    )
    pipeline = SmartFinancePipeline(PipelinePaths.from_base_dir(tmp_path, full_data=True), app_config=config)
    monkeypatch.setattr(pipeline, "_build_rows_from_record", lambda _record: {})
    restored_sources: list[dict[str, object]] = []
    try:
        pipeline.ingest_reports()
        assert set(inspected_paths) == {full_path, summary_path, outside_path}
        inspected_paths.clear()
        parsed_paths.clear()
        pipeline.db.execute("DELETE FROM financial_source")
        pipeline.ingest_reports()
        restored_sources = pipeline.db.query(
            "SELECT source_file, authority_status FROM financial_source",
            use_cache=False,
        )
    finally:
        pipeline.db.close()

    assert inspected_paths == []
    assert parsed_paths == []
    assert restored_sources == [
        {
            "source_file": dataset_source_uri(data_dir, full_path),
            "authority_status": "CURRENT",
        }
    ]
    state = SmartFinancePipeline._load_ingestion_state(tmp_path / "outputs" / "ingestion_state.json")
    assert state[str(summary_path)]["status"] == "skipped_summary"
    assert state[str(outside_path)]["status"] == "needs_review"
    assert state[str(outside_path)]["error_code"] == "outside_company_master"
    assert pipeline.run_log["ingestion_summary"]["selection"] == {
        "selected": 1,
        "needs_review": 1,
        "skipped_summary": 1,
    }
    assert pipeline.run_log["ingestion_summary"]["reused_metadata"] == 3


def test_incremental_ingestion_parses_newly_selected_revision(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data_dir = tmp_path / "正式数据"
    _write_schema_workbook(data_dir / "附件3：数据库-表名及字段说明.xlsx")
    _write_company_workbook(data_dir / "附件1：上市公司基本信息.xlsx")
    reports_dir = data_dir / "附件2：财务报告"
    reports_dir.mkdir(parents=True)
    original_path = reports_dir / "测试公司：2024年年度报告.pdf"
    revised_path = reports_dir / "测试公司：2024年年度报告（修订版）.pdf"
    original_path.write_bytes(b"original-report-content")
    metadata_by_path = {
        original_path: _report_metadata(original_path),
        revised_path: ReportFileMetadata(
            path=revised_path,
            stock_code="600000",
            stock_abbr="测试公司",
            report_period="2024FY",
            report_year=2024,
            is_summary=False,
            is_english=False,
            revision_kind=ReportRevisionKind.REVISED,
            published_date=None,
            copy_index=0,
            page_count=100,
            head_text_length=1000,
        ),
    }
    extracted_paths: list[Path] = []

    monkeypatch.setattr(
        "smart_finqa.pipeline.inspect_report_file_metadata",
        lambda path, *_, **__: metadata_by_path[path],
    )

    def extract(path: Path, _company_index: object, report: ReportFileMetadata, **_kwargs: object) -> ReportRecord:
        extracted_paths.append(path)
        return ReportRecord(
            source_path=path,
            stock_code=report.stock_code,
            stock_abbr=report.stock_abbr,
            report_period=report.report_period,
            report_year=report.report_year,
            text="有效财报文本" * 20,
            pages=(),
            snapshot={},
        )

    monkeypatch.setattr("smart_finqa.pipeline.extract_report_record", extract)
    config = AppConfig(
        db=DatabaseConfig(backend="sqlite", sqlite_path=str(tmp_path / "finance.db")),
        llm=LLMConfig(),
        incremental_ingest=True,
    )
    pipeline = SmartFinancePipeline(PipelinePaths.from_base_dir(tmp_path, full_data=True), app_config=config)
    monkeypatch.setattr(pipeline, "_build_rows_from_record", lambda _record: {})
    try:
        pipeline.ingest_reports()
        assert extracted_paths == [original_path]
        extracted_paths.clear()

        revised_path.write_bytes(b"revised-report-content")
        pipeline.ingest_reports()
        authority_rows = pipeline.db.query(
            "SELECT source_file, authority_status FROM financial_source ORDER BY source_file",
            use_cache=False,
        )
    finally:
        pipeline.db.close()

    assert extracted_paths == [revised_path]
    assert pipeline.run_log["ingestion_summary"]["skipped_unchanged"] == 0
    assert {row["source_file"]: row["authority_status"] for row in authority_rows} == {
        dataset_source_uri(data_dir, original_path): "SUPERSEDED",
        dataset_source_uri(data_dir, revised_path): "CURRENT",
    }


@pytest.mark.parametrize("changed_context", ["schema", "company_master"])
def test_incremental_ingestion_reuses_only_when_schema_and_company_master_are_unchanged(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    changed_context: str,
) -> None:
    data_dir = tmp_path / "正式数据"
    schema_path = data_dir / "附件3：数据库-表名及字段说明.xlsx"
    company_path = data_dir / "附件1：上市公司基本信息.xlsx"
    _write_schema_workbook(schema_path)
    _write_company_workbook(company_path)
    report_path = (data_dir / "附件2：财务报告" / "测试公司：2024年年度报告.pdf").resolve()
    report_path.parent.mkdir(parents=True)
    report_path.write_bytes(b"unchanged-report-content")
    report = _report_metadata(report_path)
    inspected_paths: list[Path] = []
    extracted_paths: list[Path] = []

    def inspect(path: Path, _company_index: object, **_kwargs: object) -> ReportFileMetadata:
        inspected_paths.append(path)
        return report

    def extract(path: Path, _company_index: object, _report: ReportFileMetadata, **_kwargs: object) -> ReportRecord:
        extracted_paths.append(path)
        return ReportRecord(
            source_path=path,
            stock_code=report.stock_code,
            stock_abbr=report.stock_abbr,
            report_period=report.report_period,
            report_year=report.report_year,
            text="有效财报文本" * 20,
            pages=(),
            snapshot={},
        )

    monkeypatch.setattr("smart_finqa.pipeline.inspect_report_file_metadata", inspect)
    monkeypatch.setattr("smart_finqa.pipeline.extract_report_record", extract)
    config = AppConfig(
        db=DatabaseConfig(backend="sqlite", sqlite_path=str(tmp_path / "finance.db")),
        llm=LLMConfig(),
        incremental_ingest=True,
    )
    paths = PipelinePaths.from_base_dir(tmp_path, full_data=True)

    def run_ingestion() -> SmartFinancePipeline:
        pipeline = SmartFinancePipeline(paths, app_config=config)
        monkeypatch.setattr(pipeline, "_build_rows_from_record", lambda _record: {})
        try:
            pipeline.ingest_reports()
        finally:
            pipeline.db.close()
        return pipeline

    run_ingestion()
    state_path = tmp_path / "outputs" / "ingestion_state.json"
    first_state = SmartFinancePipeline._load_ingestion_state(state_path)
    first_context = first_state[str(report_path)]["context_fingerprint"]

    unchanged_run = run_ingestion()
    assert inspected_paths == [report_path]
    assert extracted_paths == [report_path]
    assert unchanged_run.run_log["ingestion_summary"]["reused_metadata"] == 1
    assert unchanged_run.run_log["ingestion_summary"]["skipped_unchanged"] == 1

    if changed_context == "schema":
        workbook = load_workbook(schema_path)
        workbook["利润表"]["D2"] = "更新后的股票代码说明"
    else:
        workbook = load_workbook(company_path)
        workbook["基本信息表"].append([2, 600001, "新增公司", "新增公司股份有限公司"])
    workbook.save(schema_path if changed_context == "schema" else company_path)

    changed_run = run_ingestion()
    changed_state = SmartFinancePipeline._load_ingestion_state(state_path)
    assert inspected_paths == [report_path, report_path]
    assert extracted_paths == [report_path, report_path]
    assert changed_run.run_log["ingestion_summary"]["reused_metadata"] == 0
    assert changed_run.run_log["ingestion_summary"]["skipped_unchanged"] == 0
    assert changed_state[str(report_path)]["context_fingerprint"] != first_context


def test_batch_ingestion_uses_one_dataset_uri_for_authoritative_source_and_facts(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data_dir = tmp_path / "正式数据"
    _write_schema_workbook(data_dir / "附件3：数据库-表名及字段说明.xlsx")
    _write_company_workbook(data_dir / "附件1：上市公司基本信息.xlsx")
    report_path = (data_dir / "附件2：财务报告" / "测试 公司：2024年年度报告.pdf").resolve()
    report_path.parent.mkdir(parents=True)
    report_path.write_bytes(b"stable-report-content")
    report = _report_metadata(report_path)

    monkeypatch.setattr("smart_finqa.pipeline.inspect_report_file_metadata", lambda *_, **__: report)
    monkeypatch.setattr(
        "smart_finqa.pipeline.extract_report_record",
        lambda *_, **__: ReportRecord(
            source_path=tmp_path / "legacy-path-alias.pdf",
            stock_code=report.stock_code,
            stock_abbr=report.stock_abbr,
            report_period=report.report_period,
            report_year=report.report_year,
            text="利润总额 100" * 20,
            pages=(),
            snapshot={},
        ),
    )
    paths = PipelinePaths.from_base_dir(tmp_path.resolve(), full_data=True)
    config = AppConfig(
        db=DatabaseConfig(backend="sqlite", sqlite_path=str(tmp_path / "finance.db")),
        llm=LLMConfig(),
        incremental_ingest=False,
    )
    pipeline = SmartFinancePipeline(paths, app_config=config)
    monkeypatch.setattr(
        pipeline,
        "_build_rows_from_record",
        lambda _record: {
            "income_sheet": {
                "stock_code": report.stock_code,
                "stock_abbr": report.stock_abbr,
                "report_period": report.report_period,
                "report_year": report.report_year,
                "total_profit": 100.0,
            }
        },
    )
    try:
        pipeline.ingest_reports()
        joined = pipeline.db.query(
            "SELECT fs.source_file, ff.source_file AS fact_source_file, "
            "fs.source_key, ff.source_key AS fact_source_key "
            "FROM financial_source fs INNER JOIN financial_fact ff ON ff.source_key = fs.source_key",
            use_cache=False,
        )
    finally:
        pipeline.db.close()

    expected_uri = dataset_source_uri(data_dir, report_path)
    assert joined == [
        {
            "source_file": expected_uri,
            "fact_source_file": expected_uri,
            "source_key": joined[0]["source_key"],
            "fact_source_key": joined[0]["source_key"],
        }
    ]
    state = SmartFinancePipeline._load_ingestion_state(tmp_path / "outputs" / "ingestion_state.json")
    assert str(report_path) in state
    assert expected_uri not in state


def test_ingestion_marks_empty_full_document_text_for_review(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    data_dir = tmp_path / "正式数据"
    _write_schema_workbook(data_dir / "附件3：数据库-表名及字段说明.xlsx")
    _write_company_workbook(data_dir / "附件1：上市公司基本信息.xlsx")
    reports_dir = data_dir / "附件2：财务报告"
    reports_dir.mkdir(parents=True)
    report_path = reports_dir / "测试公司：2024年年度报告.pdf"
    report_path.write_bytes(b"not-read-by-test")
    report = ReportFileMetadata(
        path=report_path,
        stock_code="600000",
        stock_abbr="测试公司",
        report_period="2024FY",
        report_year=2024,
        is_summary=False,
        is_english=False,
        revision_kind=ReportRevisionKind.ORIGINAL,
        published_date=None,
        copy_index=0,
        page_count=100,
        head_text_length=0,
    )
    monkeypatch.setattr("smart_finqa.pipeline.inspect_report_file_metadata", lambda *_, **__: report)
    monkeypatch.setattr(
        "smart_finqa.pipeline.extract_report_record",
        lambda *_, **__: ReportRecord(
            source_path=report_path,
            stock_code="600000",
            stock_abbr="测试公司",
            report_period="2024FY",
            report_year=2024,
            text="",
            pages=(),
            snapshot={},
        ),
    )
    config = AppConfig(
        db=DatabaseConfig(backend="sqlite", sqlite_path=str(tmp_path / "finance.db")),
        llm=LLMConfig(),
        incremental_ingest=False,
    )
    pipeline = SmartFinancePipeline(PipelinePaths.from_base_dir(tmp_path, full_data=True), app_config=config)
    try:
        pipeline.ingest_reports()
    finally:
        pipeline.db.close()

    state = SmartFinancePipeline._load_ingestion_state(tmp_path / "outputs" / "ingestion_state.json")
    assert state[str(report_path)]["status"] == "needs_review"
    assert state[str(report_path)]["error_code"] == "document_text_unavailable"
    assert pipeline.run_log["ingestion_summary"]["parsed_files"] == 0
    summary = pipeline.run_log["ingestion_summary"]
    assert summary["candidate_facts_upserted"] == 0
    assert summary["projection_seed_rows_upserted"] == 0
    assert summary["projection_rows_with_validated_metrics"] == 0
    assert summary["validated_metric_cells_projected"] == 0
    assert "written_rows" not in summary
    assert "written_facts" not in summary


def test_ingestion_consumes_and_releases_records_before_parsing_next_batch(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data_dir = tmp_path / "正式数据"
    _write_schema_workbook(data_dir / "附件3：数据库-表名及字段说明.xlsx")
    _write_company_workbook(data_dir / "附件1：上市公司基本信息.xlsx")
    reports_dir = data_dir / "附件2：财务报告"
    reports_dir.mkdir(parents=True)
    parse_batch_size = 4
    report_count = parse_batch_size * 2
    report_paths = [
        reports_dir / f"测试公司：{2000 + index}年年度报告-{index:02d}.pdf" for index in range(report_count)
    ]
    for path in report_paths:
        path.write_bytes(b"not-read-by-test")

    def inspect(path: Path, _company_index: object, **_kwargs: object) -> ReportFileMetadata:
        index = report_paths.index(path)
        year = 2000 + index
        return ReportFileMetadata(
            path=path,
            stock_code="600000",
            stock_abbr="测试公司",
            report_period=f"{year}FY",
            report_year=year,
            is_summary=False,
            is_english=False,
            revision_kind=ReportRevisionKind.ORIGINAL,
            published_date=None,
            copy_index=0,
            page_count=100,
            head_text_length=1000,
        )

    class ProbeRecord:
        def __init__(self, path: Path, report: ReportFileMetadata, call_index: int) -> None:
            self.source_path = path
            self.stock_code = report.stock_code
            self.stock_abbr = report.stock_abbr
            self.report_period = report.report_period
            self.report_year = report.report_year
            self.text = "有效财报文本" * 20
            self.pages = ()
            self.snapshot = {}
            self.call_index = call_index

    events: list[tuple[str, int]] = []
    first_batch_refs: list[weakref.ReferenceType[ProbeRecord]] = []
    release_checks: list[bool] = []
    extract_calls = 0
    event_lock = threading.Lock()

    def extract(path: Path, _company_index: object, report: ReportFileMetadata, **_kwargs: object) -> ProbeRecord:
        nonlocal extract_calls
        with event_lock:
            call_index = extract_calls
            extract_calls += 1
        record = ProbeRecord(path, report, call_index)
        with event_lock:
            if call_index < parse_batch_size:
                first_batch_refs.append(weakref.ref(record))
            elif call_index == parse_batch_size:
                gc.collect()
                release_checks.append(all(reference() is None for reference in first_batch_refs))
            events.append(("extract", call_index))
        return record

    def build(record: ProbeRecord) -> dict[str, dict[str, object]]:
        with event_lock:
            events.append(("build", record.call_index))
        return {}

    monkeypatch.setattr("smart_finqa.pipeline.inspect_report_file_metadata", inspect)
    monkeypatch.setattr("smart_finqa.pipeline.extract_report_record", extract)
    config = AppConfig(
        db=DatabaseConfig(backend="sqlite", sqlite_path=str(tmp_path / "finance.db")),
        llm=LLMConfig(),
        incremental_ingest=False,
        ingest_workers=2,
    )
    pipeline = SmartFinancePipeline(PipelinePaths.from_base_dir(tmp_path, full_data=True), app_config=config)
    monkeypatch.setattr(pipeline, "_build_rows_from_record", build)
    try:
        pipeline.ingest_reports()
    finally:
        pipeline.db.close()

    second_batch_extract = next(index for index, event in enumerate(events) if event == ("extract", parse_batch_size))
    builds_before_second_batch = [event for event in events[:second_batch_extract] if event[0] == "build"]
    assert len(builds_before_second_batch) == parse_batch_size
    assert release_checks == [True]
    assert pipeline.run_log["ingestion_summary"]["parse_batch_size"] == parse_batch_size
    assert pipeline.run_log["ingestion_summary"]["parsed_files"] == report_count


def test_ingestion_handles_batch_when_every_extraction_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data_dir = tmp_path / "正式数据"
    _write_schema_workbook(data_dir / "附件3：数据库-表名及字段说明.xlsx")
    _write_company_workbook(data_dir / "附件1：上市公司基本信息.xlsx")
    reports_dir = data_dir / "附件2：财务报告"
    reports_dir.mkdir(parents=True)
    report_path = reports_dir / "测试公司：2024年年度报告.pdf"
    report_path.write_bytes(b"not-read-by-test")
    report = _report_metadata(report_path)

    monkeypatch.setattr("smart_finqa.pipeline.inspect_report_file_metadata", lambda *_, **__: report)

    def fail_extract(*_args: object, **_kwargs: object) -> ReportRecord:
        raise RuntimeError("synthetic extraction failure")

    monkeypatch.setattr("smart_finqa.pipeline.extract_report_record", fail_extract)
    config = AppConfig(
        db=DatabaseConfig(backend="sqlite", sqlite_path=str(tmp_path / "finance.db")),
        llm=LLMConfig(),
        incremental_ingest=False,
    )
    pipeline = SmartFinancePipeline(PipelinePaths.from_base_dir(tmp_path, full_data=True), app_config=config)
    try:
        pipeline.ingest_reports()
    finally:
        pipeline.db.close()

    state = SmartFinancePipeline._load_ingestion_state(tmp_path / "outputs" / "ingestion_state.json")
    assert state[str(report_path)]["status"] == "error"
    assert state[str(report_path)]["error"] == "synthetic extraction failure"
    assert pipeline.run_log["ingestion_summary"]["parsed_files"] == 0


def test_ingestion_keeps_batch_pending_when_authoritative_write_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data_dir = tmp_path / "正式数据"
    _write_schema_workbook(data_dir / "附件3：数据库-表名及字段说明.xlsx")
    _write_company_workbook(data_dir / "附件1：上市公司基本信息.xlsx")
    reports_dir = data_dir / "附件2：财务报告"
    reports_dir.mkdir(parents=True)
    report_path = reports_dir / "测试公司：2024年年度报告.pdf"
    report_path.write_bytes(b"not-read-by-test")
    report = _report_metadata(report_path)

    monkeypatch.setattr("smart_finqa.pipeline.inspect_report_file_metadata", lambda *_, **__: report)
    monkeypatch.setattr(
        "smart_finqa.pipeline.extract_report_record",
        lambda *_, **__: ReportRecord(
            source_path=report_path,
            stock_code=report.stock_code,
            stock_abbr=report.stock_abbr,
            report_period=report.report_period,
            report_year=report.report_year,
            text="有效财报文本" * 20,
            pages=(),
            snapshot={},
        ),
    )
    config = AppConfig(
        db=DatabaseConfig(backend="sqlite", sqlite_path=str(tmp_path / "finance.db")),
        llm=LLMConfig(),
        incremental_ingest=False,
    )
    pipeline = SmartFinancePipeline(PipelinePaths.from_base_dir(tmp_path, full_data=True), app_config=config)
    monkeypatch.setattr(pipeline, "_build_rows_from_record", lambda _record: {})

    def fail_write(_facts: object) -> None:
        raise RuntimeError("synthetic authoritative write failure")

    monkeypatch.setattr(pipeline.db, "upsert_financial_facts", fail_write)
    try:
        with pytest.raises(RuntimeError, match="synthetic authoritative write failure"):
            pipeline.ingest_reports()
    finally:
        pipeline.db.close()

    state = SmartFinancePipeline._load_ingestion_state(tmp_path / "outputs" / "ingestion_state.json")
    assert state[str(report_path)]["status"] == "parsed_pending_commit"


def test_ingestion_rejects_source_change_before_authoritative_commit(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data_dir = tmp_path / "正式数据"
    _write_schema_workbook(data_dir / "附件3：数据库-表名及字段说明.xlsx")
    _write_company_workbook(data_dir / "附件1：上市公司基本信息.xlsx")
    reports_dir = data_dir / "附件2：财务报告"
    reports_dir.mkdir(parents=True)
    report_path = reports_dir / "测试公司：2024年年度报告.pdf"
    report_path.write_bytes(b"original-report-content")
    report = _report_metadata(report_path)

    monkeypatch.setattr("smart_finqa.pipeline.inspect_report_file_metadata", lambda *_, **__: report)
    monkeypatch.setattr(
        "smart_finqa.pipeline.extract_report_record",
        lambda *_, **__: ReportRecord(
            source_path=report_path,
            stock_code=report.stock_code,
            stock_abbr=report.stock_abbr,
            report_period=report.report_period,
            report_year=report.report_year,
            text="利润总额 100" * 20,
            pages=(),
            snapshot={},
        ),
    )
    config = AppConfig(
        db=DatabaseConfig(backend="sqlite", sqlite_path=str(tmp_path / "finance.db")),
        llm=LLMConfig(),
        incremental_ingest=False,
    )
    pipeline = SmartFinancePipeline(PipelinePaths.from_base_dir(tmp_path, full_data=True), app_config=config)
    monkeypatch.setattr(
        pipeline,
        "_build_rows_from_record",
        lambda _record: {
            "income_sheet": {
                "stock_code": report.stock_code,
                "stock_abbr": report.stock_abbr,
                "report_period": report.report_period,
                "report_year": report.report_year,
                "total_profit": 100.0,
            }
        },
    )
    original_consume = pipeline._consume_parsed_record
    source_changed = False

    def consume_then_change_source(*args: object, **kwargs: object) -> bool:
        nonlocal source_changed
        consumed = original_consume(*args, **kwargs)
        if consumed and not source_changed:
            report_path.write_bytes(b"changed-after-consume")
            source_changed = True
        return consumed

    monkeypatch.setattr(pipeline, "_consume_parsed_record", consume_then_change_source)
    try:
        with pytest.raises(RuntimeError, match="source file changed before authoritative commit"):
            pipeline.ingest_reports()
        assert pipeline.db.query("SELECT COUNT(*) AS count FROM financial_fact", use_cache=False) == [{"count": 0}]
        assert pipeline.db.query("SELECT COUNT(*) AS count FROM income_sheet", use_cache=False) == [{"count": 0}]
        assert pipeline.db.query("SELECT COUNT(*) AS count FROM financial_source", use_cache=False) == [{"count": 0}]
    finally:
        pipeline.db.close()

    assert source_changed is True
    state = SmartFinancePipeline._load_ingestion_state(tmp_path / "outputs" / "ingestion_state.json")
    assert state[str(report_path)]["status"] == "parsed_pending_commit"


@pytest.mark.parametrize(
    ("workbook_name", "sheet_name", "changed_cell", "expected_error"),
    [
        (
            "附件3：数据库-表名及字段说明.xlsx",
            "利润表",
            "D2",
            "schema workbook changed before authoritative commit",
        ),
        (
            "附件1：上市公司基本信息.xlsx",
            "基本信息表",
            "D2",
            "company workbook changed before authoritative commit",
        ),
    ],
)
def test_ingestion_rejects_context_workbook_change_before_authoritative_commit(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    workbook_name: str,
    sheet_name: str,
    changed_cell: str,
    expected_error: str,
) -> None:
    data_dir = tmp_path / "正式数据"
    _write_schema_workbook(data_dir / "附件3：数据库-表名及字段说明.xlsx")
    _write_company_workbook(data_dir / "附件1：上市公司基本信息.xlsx")
    reports_dir = data_dir / "附件2：财务报告"
    reports_dir.mkdir(parents=True)
    report_path = reports_dir / "测试公司：2024年年度报告.pdf"
    report_path.write_bytes(b"stable-report-content")
    report = _report_metadata(report_path)

    monkeypatch.setattr("smart_finqa.pipeline.inspect_report_file_metadata", lambda *_, **__: report)
    monkeypatch.setattr(
        "smart_finqa.pipeline.extract_report_record",
        lambda *_, **__: ReportRecord(
            source_path=report_path,
            stock_code=report.stock_code,
            stock_abbr=report.stock_abbr,
            report_period=report.report_period,
            report_year=report.report_year,
            text="利润总额 100" * 20,
            pages=(),
            snapshot={},
        ),
    )
    config = AppConfig(
        db=DatabaseConfig(backend="sqlite", sqlite_path=str(tmp_path / "finance.db")),
        llm=LLMConfig(),
        incremental_ingest=False,
    )
    pipeline = SmartFinancePipeline(PipelinePaths.from_base_dir(tmp_path, full_data=True), app_config=config)
    monkeypatch.setattr(
        pipeline,
        "_build_rows_from_record",
        lambda _record: {
            "income_sheet": {
                "stock_code": report.stock_code,
                "stock_abbr": report.stock_abbr,
                "report_period": report.report_period,
                "report_year": report.report_year,
                "total_profit": 100.0,
            }
        },
    )
    original_consume = pipeline._consume_parsed_record
    context_changed = False

    def consume_then_change_context(*args: object, **kwargs: object) -> bool:
        nonlocal context_changed
        consumed = original_consume(*args, **kwargs)
        if consumed and not context_changed:
            workbook_path = data_dir / workbook_name
            workbook = load_workbook(workbook_path)
            workbook[sheet_name][changed_cell] = "changed-after-consume"
            workbook.save(workbook_path)
            context_changed = True
        return consumed

    monkeypatch.setattr(pipeline, "_consume_parsed_record", consume_then_change_context)
    try:
        with pytest.raises(RuntimeError, match=expected_error):
            pipeline.ingest_reports()
        assert pipeline.db.query("SELECT COUNT(*) AS count FROM financial_fact", use_cache=False) == [{"count": 0}]
        assert pipeline.db.query("SELECT COUNT(*) AS count FROM income_sheet", use_cache=False) == [{"count": 0}]
        assert pipeline.db.query("SELECT COUNT(*) AS count FROM financial_source", use_cache=False) == [{"count": 0}]
    finally:
        pipeline.db.close()

    assert context_changed is True
    state = SmartFinancePipeline._load_ingestion_state(tmp_path / "outputs" / "ingestion_state.json")
    assert state[str(report_path)]["status"] == "parsed_pending_commit"


def test_ingestion_discards_record_outputs_when_later_row_processing_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data_dir = tmp_path / "正式数据"
    _write_schema_workbook(data_dir / "附件3：数据库-表名及字段说明.xlsx")
    _write_company_workbook(data_dir / "附件1：上市公司基本信息.xlsx")
    reports_dir = data_dir / "附件2：财务报告"
    reports_dir.mkdir(parents=True)
    report_path = reports_dir / "测试公司：2024年年度报告.pdf"
    report_path.write_bytes(b"not-read-by-test")
    report = _report_metadata(report_path)

    monkeypatch.setattr("smart_finqa.pipeline.inspect_report_file_metadata", lambda *_, **__: report)
    monkeypatch.setattr(
        "smart_finqa.pipeline.extract_report_record",
        lambda *_, **__: ReportRecord(
            source_path=report_path,
            stock_code=report.stock_code,
            stock_abbr=report.stock_abbr,
            report_period=report.report_period,
            report_year=report.report_year,
            text="有效财报文本" * 20,
            pages=(),
            snapshot={},
        ),
    )
    config = AppConfig(
        db=DatabaseConfig(backend="sqlite", sqlite_path=str(tmp_path / "finance.db")),
        llm=LLMConfig(),
        incremental_ingest=False,
    )
    pipeline = SmartFinancePipeline(PipelinePaths.from_base_dir(tmp_path, full_data=True), app_config=config)
    base_row = {"stock_code": report.stock_code, "report_period": report.report_period}
    monkeypatch.setattr(
        pipeline,
        "_build_rows_from_record",
        lambda _record: {
            "income_sheet": dict(base_row),
            "balance_sheet": dict(base_row),
        },
    )
    fact_marker = object()
    monkeypatch.setattr(pipeline, "_build_facts_from_rows", lambda *_: [fact_marker])

    def sanitize(table: str, row: dict[str, object]) -> dict[str, object]:
        if table == "balance_sheet":
            raise ValueError("synthetic row processing failure")
        return row

    monkeypatch.setattr(pipeline, "_sanitize_row_values", sanitize)
    monkeypatch.setattr(pipeline, "_validate_row", lambda *_: (True, []))
    written_facts: list[list[object]] = []
    refreshed_rows: list[tuple[str, dict[str, object]]] = []
    monkeypatch.setattr(pipeline.db, "upsert_financial_facts", lambda facts: written_facts.append(list(facts)))
    monkeypatch.setattr(
        pipeline.db,
        "refresh_fact_projection",
        lambda table, row: refreshed_rows.append((table, row)),
    )
    try:
        pipeline.ingest_reports()
    finally:
        pipeline.db.close()

    state = SmartFinancePipeline._load_ingestion_state(tmp_path / "outputs" / "ingestion_state.json")
    assert state[str(report_path)]["status"] == "error"
    assert state[str(report_path)]["error"] == "Row building failed: synthetic row processing failure"
    assert written_facts == [[]]
    assert refreshed_rows == []
    assert pipeline.run_log["ingestion_summary"]["parsed_files"] == 0
    summary = pipeline.run_log["ingestion_summary"]
    assert summary["candidate_facts_upserted"] == 0
    assert summary["projection_seed_rows_upserted"] == 0
    assert summary["projection_rows_with_validated_metrics"] == 0
    assert summary["validated_metric_cells_projected"] == 0
